"""CPU-only evidence-gate tests, using the official ATK 26.9.8 XLSX schema."""

import ast
import builtins
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time

import openpyxl
import pytest


ROOT = Path(__file__).resolve().parents[1]
CHECKER = ROOT / "tests/atk/common/check_atk_result.py"
SPEC = importlib.util.spec_from_file_location("atk_result_gate", CHECKER)
gate = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(gate)


def report(path, phase="accuracy", count=2):
    """Matches official bool cells, CPU comparison nodes and header-only perf."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    statistic = workbook.active
    statistic.title = "statistic"
    summary = workbook.create_sheet("summary")
    node = "cpu_cpu_golden" if phase == "accuracy" else "npu_npu_dut"
    suffix = gate.PHASES[phase][2]
    result = suffix if phase == "determinism" else f"{node}_{suffix}"
    statistic.append(["编号", "运行结果", result])
    for case in range(count):
        statistic.append([case, "SUCCESS", 8.6481 if phase == "performance" else True])
    if phase == "determinism":
        statistic.cell(1, 4, "output_info_list")
        for row in range(2, count + 2):
            statistic.cell(row, 4, "[{'dtype': 'torch.float32', 'shape': [2, 4], 'stride': [4, 1]}]")
    if phase == "performance":
        # Real ATK single-NPU performance has no comparison summary rows.
        summary.append(["名称"])
    elif phase == "mssanitizer":
        summary.append(["名称", "总用例数", "内存检测通过率", "内存检测是否达标"])
        summary.append([node, str(count), "100.0", "Pass"])
    else:
        summary.append(["名称", "总用例数", "执行成功用例个数", "执行失败用例个数",
                        "通过用例个数", "错误信息匹配用例个数", "通过率", "精度是否达标"]
                       + (["确定性计算是否达标"] if phase == "determinism" else []))
        summary.append([node, str(count), str(count), "0", str(count), "0", "100.0",
                        "-" if phase == "determinism" else "Pass"]
                       + (["Pass"] if phase == "determinism" else []))
    workbook.save(path)
    workbook.close()
    return path


def alter(path, sheet, cell, value):
    workbook = openpyxl.load_workbook(path)
    workbook[sheet][cell] = value
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize("phase", list(gate.PHASES))
def test_official_phase_reports(phase, tmp_path):
    assert gate.validate_report(report(tmp_path / "report.xlsx", phase), phase, 2) == 2


def test_inapplicable_reference_row_and_column(tmp_path):
    path = report(tmp_path / "report.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["summary"].append(["npu_npu_dut", "-", "-", "-", "-", "-", "-", "-"])
    workbook["statistic"].cell(1, 4, "npu_npu_dut_精度通过")
    workbook.save(path)
    workbook.close()
    assert gate.validate_report(path, "accuracy", 2) == 2


@pytest.mark.parametrize("outputs", [
    "[]", "None", None, "[None]", "[{'dtype': None, 'shape': None, 'stride': None}]",
    "not a descriptor", "[", "{}", "[{}]", "[42]",
    "[{'dtype': 'torch.float32', 'shape': [0, 2], 'stride': [2, 1]}]",
    "[{'dtype': 'invalid', 'shape': [2], 'stride': [1]}]",
    "[{'dtype': 'torch.float32', 'shape': [-1], 'stride': [1]}]",
    "[{'dtype': 'torch.float32', 'shape': [2.0], 'stride': [1]}]",
    "[{'dtype': 'torch.float32', 'shape': [True], 'stride': [1]}]",
    "[{'dtype': 'torch.float32', 'shape': [2], 'stride': []}]",
    "[{'dtype': 'torch.float32', 'shape': [2], 'stride': None}]",
    "[{'dtype': None, 'shape': [2], 'stride': [1]}]",
])
def test_determinism_true_without_real_output_evidence_fails(outputs, tmp_path):
    path = report(tmp_path / "report.xlsx", "determinism")
    alter(path, "statistic", "D2", outputs)
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, "determinism", 2)


@pytest.mark.parametrize("outputs", [
    "[None, {'dtype': 'torch.bfloat16', 'shape': [2], 'stride': [1]}]",
    "[{'dtype': None, 'shape': None, 'stride': None}, {'dtype': 'torch.float32', 'shape': [2], 'stride': [1]}]",
    "[{'dtype': 'torch.float16', 'shape': [0], 'stride': [1]}, {'dtype': 'torch.float32', 'shape': [2], 'stride': [1]}]",
    "[{'dtype': 'torch.float32', 'shape': [], 'stride': []}]",
])
def test_determinism_allows_optional_outputs_with_real_tensor(outputs, tmp_path):
    path = report(tmp_path / "report.xlsx", "determinism")
    alter(path, "statistic", "D2", outputs)
    assert gate.validate_report(path, "determinism", 2) == 2


def test_determinism_requires_output_descriptor_column(tmp_path):
    path = report(tmp_path / "report.xlsx", "determinism")
    alter(path, "statistic", "D1", "not_output_evidence")
    with pytest.raises(gate.InvalidReport, match="missing columns"):
        gate.validate_report(path, "determinism", 2)


@pytest.mark.parametrize("value", [True, "True", "TRUE"])
def test_official_boolean_representations(value, tmp_path):
    path = report(tmp_path / "report.xlsx")
    alter(path, "statistic", "C2", value)
    assert gate.validate_report(path, "accuracy", 2) == 2


@pytest.mark.parametrize("sheet,cell,value", [
    ("statistic", "B2", "FAILED"), ("statistic", "B2", "TIMEOUT"),
    ("statistic", "B2", "PROCESSING"), ("statistic", "B2", "unknown"),
    ("statistic", "B2", None), ("statistic", "A3", 0),
    ("statistic", "A2", -1), ("statistic", "A2", 0.5),
    ("statistic", "A2", None), ("statistic", "C2", False),
    ("statistic", "C2", "unknown"), ("statistic", "C2", None),
    ("statistic", "C2", 1), ("statistic", "C1", "wrong_result"),
    ("statistic", "B1", "wrong_status"), ("statistic", "C1", "编号"),
    ("summary", "B2", 0), ("summary", "B2", 3),
    ("summary", "C2", 1), ("summary", "D2", 1),
    ("summary", "E2", 0), ("summary", "E2", None),
    ("summary", "H2", "Failed"), ("summary", "H2", "unknown"),
    ("summary", "H2", None), ("summary", "H1", "wrong_verdict"),
    ("summary", "G2", "99.5"), ("summary", "A2", "wrong_node"),
])
def test_incomplete_or_failed_evidence_rejected(sheet, cell, value, tmp_path):
    path = report(tmp_path / "report.xlsx")
    alter(path, sheet, cell, value)
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, "accuracy", 2)


@pytest.mark.parametrize("phase", list(gate.PHASES))
@pytest.mark.parametrize("count", [0, 1, 3])
def test_expected_count_must_match_every_phase(phase, count, tmp_path):
    path = report(tmp_path / "report.xlsx", phase)
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, phase, count)


def test_case_ids_must_match_selected_json_not_just_count(tmp_path):
    path = report(tmp_path / "report.xlsx")
    alter(path, "statistic", "A2", 70)
    alter(path, "statistic", "A3", 71)
    with pytest.raises(gate.InvalidReport, match="ID mismatch"):
        gate.validate_report(path, "accuracy", 2, [0, 1])
    assert gate.validate_report(path, "accuracy", 2, [71, 70]) == 2


@pytest.mark.parametrize("phase", list(gate.PHASES))
def test_zero_statistic_cases_fail(phase, tmp_path):
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(report(tmp_path / "report.xlsx", phase, 0), phase)


@pytest.mark.parametrize("sheet", ["summary", "statistic"])
def test_missing_sheets_fail(sheet, tmp_path):
    path = report(tmp_path / "report.xlsx")
    workbook = openpyxl.load_workbook(path)
    del workbook[sheet]
    workbook.save(path)
    workbook.close()
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, "accuracy")


@pytest.mark.parametrize("value", [0, -1, "nan", "inf", "bad", True, None])
def test_performance_requires_real_positive_finite_timing(value, tmp_path):
    path = report(tmp_path / "report.xlsx", "performance")
    alter(path, "statistic", "C2", value)
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, "performance", 2)


def test_performance_cpu_timing_is_not_npu_evidence(tmp_path):
    path = report(tmp_path / "report.xlsx", "performance")
    alter(path, "statistic", "C1", "cpu_cpu_golden_Device性能（us）")
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, "performance", 2)


def test_performance_explicit_comparison_failure_not_ignored(tmp_path):
    path = report(tmp_path / "report.xlsx", "performance")
    alter(path, "statistic", "D1", "npu_npu_dut_device性能通过")
    alter(path, "statistic", "D2", False)
    with pytest.raises(gate.InvalidReport):
        gate.validate_report(path, "performance", 2)


def test_missing_stale_and_corrupt_newest_report_fail(tmp_path):
    with pytest.raises(gate.InvalidReport, match="no current"):
        gate.check_report(tmp_path, "op", "accuracy")
    path = report(tmp_path / "accuracy/atk_output/atk_op_old/report/good.xlsx")
    os.utime(path, ns=(1_000_000_000, 1_000_000_000))
    with pytest.raises(gate.InvalidReport, match="no current"):
        gate.check_report(tmp_path, "op", "accuracy", time.time_ns(), 2)
    bad = path.with_name("corrupt.xlsx")
    bad.write_text("not an xlsx", encoding="utf-8")
    with pytest.raises(gate.InvalidReport, match="cannot read"):
        gate.check_report(tmp_path, "op", "accuracy", expected_count=2)


def test_cli_missing_report_is_nonzero(tmp_path):
    result = subprocess.run([sys.executable, str(CHECKER), "--type", "all",
                             "--output-root", str(tmp_path), "--op", "missing"],
                            capture_output=True, text=True)
    assert result.returncode == 1
    assert result.stdout.count("Failed") == 4


MSS_START = "[mssanitizer] Start memcheck sanitizer on kernel Op_123_0\n"
MSS_FINISH = "[mssanitizer] Sanitizer finished on kernel Op_123_0. No error detected.\n"


@pytest.mark.parametrize("text", [
    "", "[mssanitizer] initialized successfully\n", MSS_START, MSS_FINISH,
    MSS_START + MSS_FINISH.replace("Op_123_0", "Other_0"),
    (MSS_START + MSS_FINISH).replace("memcheck", "racecheck"),
    (MSS_START + MSS_FINISH).replace("Op_123_0", "IsFinite_0"),
    MSS_START + MSS_FINISH + "ERROR memory out of bounds\n",
    MSS_START + MSS_FINISH + "kernel not instrumented\n",
    MSS_START + MSS_FINISH + "unsupported kernel skipped\n",
])
def test_mss_requires_complete_target_kernel_inspection(text, tmp_path):
    path = tmp_path / "mss.log"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(gate.InvalidReport):
        gate.validate_mss_log(path, None, "op", expected_count=1)


def test_mss_target_count_and_freshness(tmp_path):
    path = tmp_path / "mss.log"
    path.write_text(MSS_START + MSS_FINISH, encoding="utf-8")
    assert gate.validate_mss_log(path, None, "op", expected_count=1) == 1
    with pytest.raises(gate.InvalidReport, match="launches"):
        gate.validate_mss_log(path, None, "op", expected_count=2)
    os.utime(path, ns=(1_000_000_000, 1_000_000_000))
    with pytest.raises(gate.InvalidReport, match="stale"):
        gate.validate_mss_log(path, time.time_ns(), "op")
    with pytest.raises(gate.InvalidReport, match="missing"):
        gate.validate_mss_log(tmp_path / "absent.log", None, "op")


def test_prepare_executor_keyword_binds_to_stable_public_api():
    executor = ROOT / "tests/atk/chunk_kda_fwd_prepare/executor_chunk_kda_fwd_prepare.py"
    tree = ast.parse(executor.read_text(encoding="utf-8"))
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and isinstance(n.func, ast.Name) and n.func.id == "chunk_kda_fwd_prepare"]
    assert len(calls) == 1
    keys = {k.arg for k in calls[0].keywords}
    assert "A_log" in keys and "a_log" not in keys
    stable = ast.parse((ROOT / "torch_custom/fla_npu/fla_npu/ops/ascendc/_stable.py").read_text(encoding="utf-8"))
    function = next(n for n in stable.body if isinstance(n, ast.FunctionDef)
                    and n.name == "npu_chunk_kda_fwd_prepare")
    assert keys <= {a.arg for a in function.args.args + function.args.kwonlyargs}


@pytest.mark.parametrize("message,expected", [
    ("aclnnStatus=161001", True),
    ("aclnnChunk GetWorkspaceSize failed: 161001", True),
    ("GetWorkspaceSize failed: 161002", False),
    ("allocation failed 161001", False),
    ("GetWorkspaceSize failed: 161001extra", False),
])
def test_preprocess_negative_accepts_only_expected_aclnn_status(message, expected):
    path = ROOT / "tests/atk/pre_process_bwd_kernel_merged/executor_pre_process_bwd_kernel_merged.py"
    tree = ast.parse(path.read_text(encoding="utf-8"))
    function = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "run_npu")
    def call(*args, **kwargs):
        raise RuntimeError(message)
    namespace = {"re": re, "InputDataset": object, "OP_NAME": "probe",
                 "_negative_expected_code": lambda spec: 161001,
                 "build_inputs": lambda *args: {}, "_marker_device": lambda *args: "npu",
                 "_call_op": call}
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(path), "exec"), namespace)
    with pytest.raises(RuntimeError) as error:
        namespace["run_npu"]({"case_key": "negative", "expected_error_msg": "EXPECTED"}, object())
    assert (str(error.value) == "EXPECTED") is expected


def test_formal_case_error_expectations_are_explicit():
    negatives = []
    for path in (ROOT / "tests/atk").rglob("atk_*.json"):
        cases = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(cases, list):
            continue
        for case in cases:
            pattern = case.get("expected_error_msg")
            assert pattern is None or (isinstance(pattern, str) and pattern.strip()), (path, case.get("id"))
            if pattern:
                negatives.append((path.parent.name, case["id"], pattern))
    assert len(negatives) == 13
    assert all(op == "pre_process_bwd_kernel_merged" for op, _, _ in negatives)


@pytest.mark.parametrize("op", ["chunk_kda_fwd_prepare", "chunk_kda_fwd_finalize"])
def test_kda_generators_keep_positive_error_sentinel_null(op, monkeypatch):
    # Exercise generation without importing optional ATK/torch dependencies.
    original_import = builtins.__import__
    def without_atk(name, *args, **kwargs):
        if name == "atk" or name.startswith("atk."):
            raise ModuleNotFoundError("ATK is optional for this CPU test", name="atk")
        return original_import(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", without_atk)
    path = ROOT / f"tests/atk/{op}/gen_{op}.py"
    spec = importlib.util.spec_from_file_location(f"_test_gen_{op}", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    if op.endswith("prepare"):
        cases = [case for suite in module.SUITES for case in module._cases(suite)]
        assert len(cases) == 642
    else:
        cases = [case for build in module.SUITES.values() for case in build()]
        assert len(cases) == 222
    assert all(case["expected_error_msg"] is None for case in cases)


# Run the real shell orchestration with an ATK stub on Linux. The stub only
# writes official-schema evidence; no torch import, NPU or alternate golden.
@pytest.fixture
def runner(tmp_path):
    if os.name == "nt" or not shutil.which("bash"):
        pytest.skip("shell integration requires Linux bash")
    atk = tmp_path / "atk"
    (atk / "common").mkdir(parents=True)
    shutil.copyfile(CHECKER, atk / "common/check_atk_result.py")
    shutil.copyfile(ROOT / "tests/atk/run_test_cpu.sh", atk / "run_test_cpu.sh")
    op = atk / "op"
    op.mkdir()
    (op / "executor_op.py").write_text("", encoding="utf-8")
    for suffix in ("", "_perf", "_mss"):
        (op / f"atk_op{suffix}.json").write_text(json.dumps([{"id": 0}, {"id": 1}]))
    bin_dir = tmp_path / "bin"
    bin_dir.mkdir()
    # Reuse fixture report writer via an isolated import, without pytest tests.
    stub = bin_dir / "atk"
    stub.write_text(f"#!{sys.executable}\n" + '''
import importlib.util, json, os, pathlib, sys
if '--version' in sys.argv:
    print('26.9.8'); raise SystemExit(0)
if os.environ.get('ATK_STUB_ARGS'):
    pathlib.Path(os.environ['ATK_STUB_ARGS']).write_text(json.dumps(sys.argv[1:]))
mode = os.environ.get('ATK_STUB_MODE', '')
if mode == 'exit': raise SystemExit(7)
if mode == 'missing': raise SystemExit(0)
spec = importlib.util.spec_from_file_location('fixtures', os.environ['ATK_FIXTURE'])
fixtures = importlib.util.module_from_spec(spec); spec.loader.exec_module(fixtures)
phase = {'accuracy': 'accuracy', 'accuracy_dc': 'determinism',
         'performance_device': 'performance', 'run': 'mssanitizer'}[sys.argv[sys.argv.index('--task') + 1]]
root = pathlib.Path(sys.argv[sys.argv.index('--output_path') + 1])
assert root.is_absolute()
path = root / 'atk_output/atk_op_new/report/report.xlsx'
fixtures.report(path, phase, 1 if mode == 'partial' else 2)
if mode == 'failed': fixtures.alter(path, 'statistic', 'B2', 'FAILED')
if mode == 'wrong_ids': fixtures.alter(path, 'statistic', 'A2', 70)
if mode == 'stale': os.utime(path, ns=(1000000000, 1000000000))
''', encoding="utf-8")
    stub.chmod(0o755)
    sanitizer = bin_dir / "mssanitizer"
    sanitizer.write_text('''#!/usr/bin/env bash
[[ "$1" == --tool=memcheck && "$2" == --log-file && "$4" == --log-level && "$5" == info && "$6" == -- ]] || exit 9
if [[ "${ATK_STUB_MODE:-}" == empty_mss ]]; then
  : > "$3"
else
  printf '%s\\n' '[mssanitizer] Start memcheck sanitizer on kernel Op_0' '[mssanitizer] Sanitizer finished on kernel Op_0. No error detected.' '[mssanitizer] Start memcheck sanitizer on kernel Op_1' '[mssanitizer] Sanitizer finished on kernel Op_1. No error detected.' > "$3"
fi
shift 6
exec "$@"
''', encoding="utf-8")
    sanitizer.chmod(0o755)
    python = bin_dir / "python3"
    python.write_text(f'#!/usr/bin/env bash\nexec "{sys.executable}" "$@"\n', encoding="utf-8")
    python.chmod(0o755)
    env = os.environ.copy()
    env.update(PATH=str(bin_dir) + os.pathsep + env["PATH"],
               ATK_OUTPUT_ROOT=str(tmp_path / "output"), ATK_FIXTURE=str(Path(__file__).resolve()))
    for key in ("ATK_ENV", "CANN_ENV", "FLA_NPU_ENV", "FLA_NPU_OPP_ENV", "CASE_START", "CASE_END",
                "ATK_GM_INIT_MODE", "ATK_GM_INIT_MB"):
        env.pop(key, None)
    return atk, env


@pytest.mark.parametrize("mode", ["", "missing", "exit", "partial", "failed", "stale", "missing_checker", "empty_mss", "wrong_ids"])
def test_shell_checks_all_stages_and_fails_closed(runner, mode):
    atk, env = runner
    env["ATK_STUB_MODE"] = mode
    if mode == "missing_checker":
        (atk / "common/check_atk_result.py").unlink()
    result = subprocess.run(["bash", str(atk / "run_test_cpu.sh"), "-op=op", "-soc=A5", "-scope=all"],
                            env=env, capture_output=True, text=True)
    if not mode:
        assert result.returncode == 0, result.stdout + result.stderr
        assert result.stdout.count(": Pass (2 cases)") == 4
    else:
        assert result.returncode != 0, result.stdout + result.stderr


@pytest.mark.parametrize("mode,size,expected", [
    (None, None, "default"), ("on", "", "default"), ("off", "", None),
    ("on", "512", "512"), ("on", "1", "1"), ("enable", "512", "512"),
])
def test_shell_gm_initialization_cli(runner, mode, size, expected):
    atk, env = runner
    args_path = atk / "args.json"
    env["ATK_STUB_ARGS"] = str(args_path)
    if mode is not None:
        env["ATK_GM_INIT_MODE"] = mode
    if size is not None:
        env["ATK_GM_INIT_MB"] = size
    result = subprocess.run(["bash", str(atk / "run_test_cpu.sh"), "-op=op", "-soc=A5", "-scope=accuracy"],
                            env=env, capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
    args = json.loads(args_path.read_text())
    if expected is None:
        assert "--gm_init_flag" not in args
    else:
        index = args.index("--gm_init_flag")
        assert args[index + 1] == ("-to" if expected == "default" else expected)
        if expected != "default":
            assert f"ATK_GM_INIT_MB={size}" in result.stdout
            assert f"ATK CLI --gm_init_flag {size}" in result.stdout


@pytest.mark.parametrize("mode,size", [
    ("on", "0"), ("on", "-1"), ("on", "1.5"), ("on", "abc"),
    ("on", " 512"), ("on", "512 "), ("on", "+512"), ("on", "01"),
    ("on", "0x200"), ("on", "1e3"), ("on", "512\n"),
    ("off", "512"), ("false", "512"), ("0", "512"), ("disable", "512"),
])
def test_shell_gm_invalid_or_disabled_size_rejected(runner, mode, size):
    atk, env = runner
    args_path = atk / "args.json"
    env.update(ATK_GM_INIT_MODE=mode, ATK_GM_INIT_MB=size, ATK_STUB_ARGS=str(args_path))
    result = subprocess.run(["bash", str(atk / "run_test_cpu.sh"), "-op=op", "-soc=A5", "-scope=accuracy"],
                            env=env, capture_output=True, text=True)
    assert result.returncode != 0, result.stdout + result.stderr
    assert "ATK_GM_INIT_MB" in result.stderr
    assert not args_path.exists(), "invalid configuration must fail before launching an ATK task"


@pytest.mark.parametrize("pattern,valid", [(None, True), ("", False), (" ", False), ("\t\n", False),
                                         ("ACLNN_ERR_PARAM_INVALID: expected", True)])
def test_shell_error_expectation_preflight(runner, pattern, valid):
    atk, env = runner
    args_path = atk / "args.json"
    env["ATK_STUB_ARGS"] = str(args_path)
    path = atk / "op/atk_op.json"
    cases = json.loads(path.read_text())
    cases[0]["expected_error_msg"] = pattern
    path.write_text(json.dumps(cases))
    result = subprocess.run(["bash", str(atk / "run_test_cpu.sh"), "-op=op", "-soc=A5", "-scope=accuracy"],
                            env=env, capture_output=True, text=True)
    if valid:
        assert result.returncode == 0, result.stdout + result.stderr
        assert args_path.exists()
    else:
        assert result.returncode != 0, result.stdout + result.stderr
        assert "expected_error_msg" in result.stderr
        assert not args_path.exists()
