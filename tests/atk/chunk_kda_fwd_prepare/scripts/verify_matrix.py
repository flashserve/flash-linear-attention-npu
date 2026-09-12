#!/usr/bin/env python3
"""校验 ChunkKdaFwdPrepare ATK 分片及完整矩阵覆盖。"""

from __future__ import annotations

import argparse
import glob
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
from collections import Counter
from pathlib import Path


OP_NAME = "chunk_kda_fwd_prepare"
HOST_KEY_RE = re.compile(
    r"ChunkKdaFwdPrepare tiling:.*?outputMode=(\d+).*?tilingKey=(\d+)"
)
LAUNCH_KEY_RE = re.compile(
    r"OpName:\[[^\]]*ChunkKdaFwdPrepare[^\]]*\]\s+"
    r"Tiling Key:\s*(\d+)"
)
KERNEL_NAME_PATTERN = r"(?:chunk_kda_fwd_prepare|ChunkKdaFwdPrepare)"
SHARD_DIR_RE = re.compile(r"^shard_(0|[1-9][0-9]*)_(0|[1-9][0-9]*)$")
KERNEL_BIN_NAME_RE = re.compile(
    rf"^(?P<op>{KERNEL_NAME_PATTERN})_(?P<compile_hash>[0-9a-fA-F]{{32}})$"
)
SANITIZER_TOOLS = frozenset(
    {"memcheck", "racecheck", "initcheck", "synccheck"}
)
SANITIZER_COMMON_DIAGNOSTICS = (
    re.compile(
        r"^\s*(?:=+\s*)?ERROR SUMMARY:\s*[1-9][0-9]*"
        r"(?:\s+errors?)?\b",
        re.IGNORECASE | re.MULTILINE,
    ),
)
SANITIZER_TOOL_DIAGNOSTICS = {
    "memcheck": (
        re.compile(
            r"^\s*=+\s*(?:ERROR|WARNING)\s*:\s*"
            r"(?:out\s+of\s+bounds|illegal\s+(?:read|write)|"
            r"unaligned\s+access|multi-core\s+corruption|memory\s+leak|"
            r"LeakCheck:\s*detected\s+memory\s+leaks|Unused\s+memory\s+of\b|"
            r"(?:illegal|invalid|double)\s+(?:free|release))\b",
            re.IGNORECASE | re.MULTILINE,
        ),
    ),
    "racecheck": (
        re.compile(
            r"^\s*=+\s*(?:ERROR|WARNING)\s*:\s*"
            r"Potential\s+(?:RAW|WAR|WAW)\s+hazard\s+detected\b",
            re.IGNORECASE | re.MULTILINE,
        ),
    ),
    "initcheck": (
        re.compile(
            r"^\s*=+\s*(?:ERROR|WARNING)\s*:\s*"
            r"(?:uninitiali[sz]ed(?:\s+memory)?\s+read|"
            r"read\s+of\s+uninitiali[sz]ed(?:\s+memory)?)\b",
            re.IGNORECASE | re.MULTILINE,
        ),
    ),
    "synccheck": (
        re.compile(
            r"^\s*=+\s*(?:ERROR|WARNING)\s*:\s*"
            r"(?:Unpaired|Redundant)\s+(?:set_flag|wait_flag)\s+"
            r"instructions?\s+detected\b",
            re.IGNORECASE | re.MULTILINE,
        ),
        re.compile(
            r"^\s*=+\s*ERROR\s*:\s*Sync error detected\.?\s*"
            r"kernel locked up at\b",
            re.IGNORECASE | re.MULTILINE,
        ),
        re.compile(
            r"^\s*=+\s*ERROR\s*:\s*Sync error detected\.\s*"
            r"Divergent thread\(s\)\b",
            re.IGNORECASE | re.MULTILINE,
        ),
        re.compile(
            r"^\s*=+\s*SUMMARY\s*:\s*[1-9][0-9]*\s+pipe\(s\)\s+locked up\.",
            re.IGNORECASE | re.MULTILINE,
        ),
    ),
}
SANITIZER_GENERIC_DIAGNOSTIC = re.compile(
    r"^\s*=+\s*(?:ERROR|WARNING)\s*:", re.IGNORECASE | re.MULTILINE
)
# mssanitizer 有时把错误摘要单独打印成这一行，而不是附在 finish 行上。
# 该标记明确表示前面存在诊断，不能被当作 clean finish 的旁证。
SANITIZER_ERROR_MARKER = re.compile(
    r"^\s*(?:=+\s*)?(?:\[mssanitizer\]\s*)?"
    r"See\s+all\s+detected\s+errors\s+above\.\s*$",
    re.IGNORECASE | re.MULTILINE,
)
SANITIZER_FINISH_RE = re.compile(
    rf"^\s*\[mssanitizer\]\s*Sanitizer finished on kernel\s+"
    rf"(?P<kernel>{KERNEL_NAME_PATTERN}(?:_[A-Za-z0-9]+)+)\.\s+"
    rf"(?P<status>No error detected\.|See all detected errors above\.)\s*$",
    re.IGNORECASE | re.MULTILINE,
)
SANITIZER_REGISTER_RE = re.compile(
    rf"^\s*\[mssanitizer\]\s*Warning:\s*Register\s+\S+\s+was not reset "
    rf"to default\b.*?\bon kernel\s+"
    rf"(?P<kernel>{KERNEL_NAME_PATTERN}(?:_[A-Za-z0-9]+)+)\.\s+"
    rf"Expected default value is\b",
    re.IGNORECASE | re.MULTILINE,
)
SCOPE_CONTRACTS = {
    "accuracy": {
        "case_count": 200,
        "loop_nums": 1,
        "gm_init_mode": "off",
        "single_process_mode": "off",
        "tools": frozenset({""}),
        "require_tiling_log": False,
    },
    "determinism": {
        "case_count": 432,
        "loop_nums": 50,
        "gm_init_mode": "not_applicable",
        "single_process_mode": "off",
        "tools": frozenset({""}),
        "require_tiling_log": True,
    },
    "mssanitizer": {
        "case_count": 432,
        "loop_nums": 1,
        "gm_init_mode": "not_applicable",
        "single_process_mode": "off",
        "tools": SANITIZER_TOOLS,
        "require_tiling_log": True,
    },
}
SCOPE_TIMEOUT_LIMITS = {
    "accuracy": 60,
    "determinism": 60,
    "mssanitizer": 1000,
}


def _load_cases(path: Path) -> list[dict]:
    _reject_symlink_chain(path, "用例文件")
    cases = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(cases, list) or not cases:
        raise ValueError(f"用例文件必须是非空列表：{path}")
    return cases


def _case_spec(case: dict) -> dict:
    for item in case.get("inputs", ()):
        if item.get("name") == "case_spec":
            value = item.get("range_values")
            return json.loads(value) if isinstance(value, str) else value
    raise ValueError(f"case {case.get('id')} 缺少 case_spec")


def _validate_accuracy_seed_contract(cases: list[dict]) -> None:
    groups: dict[str, list[tuple[int, int, int]]] = {}
    for case in cases:
        case_id = case.get("id")
        spec = _case_spec(case)
        logical_key = spec.get("logical_case_key")
        seed = spec.get("seed")
        seed_index = spec.get("seed_index")
        if not isinstance(logical_key, str) or not logical_key:
            raise ValueError(f"accuracy case {case_id} 缺少 logical_case_key")
        if (
            not isinstance(seed, int)
            or isinstance(seed, bool)
            or not isinstance(seed_index, int)
            or isinstance(seed_index, bool)
        ):
            raise ValueError(f"accuracy case {case_id} 的 seed/seed_index 非法")
        if case.get("default_seed") != seed:
            raise ValueError(
                f"accuracy case {case_id} 的 default_seed 与 case_spec.seed 不一致"
            )
        groups.setdefault(logical_key, []).append((int(case_id), seed, seed_index))

    if len(cases) != 200 or len(groups) != 56:
        raise ValueError("accuracy 正式矩阵必须是 200 条、56 个逻辑场景")
    for logical_key, entries in groups.items():
        seeds = [seed for _, seed, _ in entries]
        seed_indices = [seed_index for _, _, seed_index in entries]
        if len(entries) not in (3, 4):
            raise ValueError(
                f"accuracy 逻辑场景 {logical_key} 必须包含 3 或 4 条用例"
            )
        if len(set(seeds)) != len(entries):
            raise ValueError(f"accuracy 逻辑场景 {logical_key} 存在重复 seed")
        if sorted(seed_indices) != list(range(len(entries))):
            raise ValueError(
                f"accuracy 逻辑场景 {logical_key} 的 seed_index 必须从 0 连续编号"
            )


def _case_hash(path: Path) -> str:
    _reject_symlink_chain(path, "用例文件")
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _file_hash(path: Path) -> str:
    _reject_symlink_chain(path, "证据文件")
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError(f"证据文件不存在或为空：{path}")
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _reject_symlink(path: Path, label: str) -> None:
    """拒绝可替换原始证据的符号链接。"""
    if path.is_symlink():
        raise ValueError(f"{label}不能是符号链接：{path}")


def _reject_symlink_chain(path: Path, label: str, stop_at: Path | None = None) -> None:
    """检查文件及其父目录，避免通过祖先符号链接替换运行时产物。"""
    current = Path(path)
    boundary = Path(stop_at) if stop_at is not None else None
    while True:
        _reject_symlink(current, label)
        if boundary is not None and current == boundary:
            return
        parent = current.parent
        if parent == current:
            return
        current = parent


def _key_digest(pairs: list[tuple[int, int]]) -> str:
    payload = "".join(f"{case_id},{key}\n" for case_id, key in pairs)
    return hashlib.sha256(payload.encode("ascii")).hexdigest()


def _integer_digest(values: set[int]) -> str:
    payload = "".join(f"{value}\n" for value in sorted(values))
    return hashlib.sha256(payload.encode("ascii")).hexdigest()


def _string_digest(values: set[str]) -> str:
    payload = "".join(f"{value}\n" for value in sorted(values))
    return hashlib.sha256(payload.encode("ascii")).hexdigest()


def _json_digest(value: object) -> str:
    payload = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


_DTYPE_ALIASES = {
    "bf16": "bf16",
    "bfloat16": "bf16",
    "fp32": "fp32",
    "float32": "fp32",
}
_SIMPLIFIED_DTYPE_TOKENS = {"0": "fp32", "27": "bf16"}
_SIMPLIFIED_FORMAT_TOKENS = {"2": "ND"}
_SIMPLIFIED_TENSOR_COUNT = 18


def _normalize_dispatch_dtype(value: object, label: str) -> str:
    normalized = _DTYPE_ALIASES.get(str(value).strip().lower())
    if normalized is None:
        raise ValueError(f"{label} dtype 非法：{value!r}")
    return normalized


def _validate_simplified_keys(
    support_info: dict, dispatch_signature: dict[str, dict], label: str
) -> list[str]:
    if support_info.get("simplifiedKeyMode") != 0:
        raise ValueError(f"kernel metadata 的 simplifiedKeyMode 不是 0：{label}")
    simplified_keys = support_info.get("simplifiedKey")
    if (
        not isinstance(simplified_keys, list)
        or not simplified_keys
        or any(not isinstance(value, str) or not value for value in simplified_keys)
    ):
        raise ValueError(f"kernel metadata 的 simplifiedKey 非法：{label}")

    seen: set[str] = set()
    for value in simplified_keys:
        if value in seen:
            raise ValueError(f"kernel metadata 的 simplifiedKey 重复：{label}")
        seen.add(value)
        fields = value.split("/")
        if (
            len(fields) != _SIMPLIFIED_TENSOR_COUNT + 2
            or fields[0] != "ChunkKdaFwdPrepare"
            or re.fullmatch(r"d=[01],p=[01]", fields[1]) is None
        ):
            raise ValueError(f"kernel metadata 的 simplifiedKey 结构非法：{label}")
        tensors = fields[2:]
        parsed: list[tuple[str, str]] = []
        for token in tensors:
            match = re.fullmatch(r"([0-9]+),([0-9]+)", token)
            if match is None:
                raise ValueError(
                    f"kernel metadata 的 simplifiedKey tensor 非法：{label}"
                )
            dtype = _SIMPLIFIED_DTYPE_TOKENS.get(match.group(1))
            tensor_format = _SIMPLIFIED_FORMAT_TOKENS.get(match.group(2))
            if dtype is None or tensor_format is None:
                raise ValueError(
                    f"kernel metadata 的 simplifiedKey token 未知：{label}"
                )
            parsed.append((dtype, tensor_format))
        for input_name, tensor_index in (("g", 3), ("beta", 4)):
            expected = dispatch_signature[input_name]
            if parsed[tensor_index] != (
                expected["dtype"],
                expected["format"],
            ):
                raise ValueError(
                    "kernel metadata 的 simplifiedKey 与 supportInfo.inputs "
                    f"中 {input_name} 不一致：{label}"
                )
    return simplified_keys


def _metadata_dispatch_contract(metadata: dict, label: str) -> dict:
    """从真实 supportInfo 中提取 g/beta 调度签名。"""
    support_info = metadata.get("supportInfo")
    if not isinstance(support_info, dict):
        raise ValueError(f"kernel metadata 缺少 supportInfo：{label}")
    inputs = support_info.get("inputs")
    if not isinstance(inputs, list) or not inputs:
        raise ValueError(f"kernel metadata 的 supportInfo.inputs 非法：{label}")

    by_name: dict[str, dict] = {}
    indices: set[int] = set()
    for item in inputs:
        if not isinstance(item, dict):
            raise ValueError(f"supportInfo.inputs 元素非法：{label}")
        name = str(item.get("name", ""))
        if not name or name in by_name:
            raise ValueError(f"supportInfo.inputs 的 name 缺失或重复：{label}")
        index = item.get("index")
        if not isinstance(index, int) or isinstance(index, bool) or index in indices:
            raise ValueError(f"supportInfo.inputs 的 index 非法或重复：{label}")
        by_name[name] = item
        indices.add(index)

    dispatch_signature: dict[str, dict] = {}
    for input_name, case_name, expected_index in (
        ("g", "gate_dtype", 3),
        ("beta", "beta_dtype", 4),
    ):
        item = by_name.get(input_name)
        if item is None:
            raise ValueError(f"supportInfo.inputs 缺少 {input_name}：{label}")
        if item.get("index") != expected_index:
            raise ValueError(
                f"supportInfo.inputs 的 {input_name} index 非法：{label}"
            )
        input_format = str(item.get("format", "")).upper()
        if input_format != "ND":
            raise ValueError(
                f"supportInfo.inputs 的 {input_name} format 非法：{label}"
            )
        param_type = str(item.get("paramType", "")).lower()
        if param_type != "required":
            raise ValueError(
                f"supportInfo.inputs 的 {input_name} paramType 非法：{label}"
            )
        dispatch_signature[input_name] = {
            "index": expected_index,
            "dtype": _normalize_dispatch_dtype(
                item.get("dtype"), f"supportInfo.inputs.{input_name}"
            ),
            "format": input_format,
            "param_type": param_type,
            "case_field": case_name,
        }

    simplified_key = _validate_simplified_keys(
        support_info, dispatch_signature, label
    )
    return {
        "dispatch_signature": dispatch_signature,
        "simplified_key_mode": 0,
        "support_info_sha256": _json_digest(support_info),
        "simplified_key_sha256": _json_digest(simplified_key),
    }


def _manifest_dispatch_signature(binary: dict) -> tuple[str, str]:
    if binary.get("simplified_key_mode") != 0:
        raise ValueError("runtime manifest 的 simplified_key_mode 不是 0")
    signature = binary.get("dispatch_signature")
    if not isinstance(signature, dict) or set(signature) != {"g", "beta"}:
        raise ValueError("runtime manifest 缺少 g/beta dispatch_signature")
    expected = {
        "g": (3, "gate_dtype"),
        "beta": (4, "beta_dtype"),
    }
    dtypes: dict[str, str] = {}
    for name, (index, case_field) in expected.items():
        item = signature.get(name)
        if not isinstance(item, dict) or set(item) != {
            "index",
            "dtype",
            "format",
            "param_type",
            "case_field",
        }:
            raise ValueError(f"runtime manifest 的 {name} 调度签名非法")
        if (
            item.get("index") != index
            or item.get("format") != "ND"
            or item.get("param_type") != "required"
            or item.get("case_field") != case_field
        ):
            raise ValueError(f"runtime manifest 的 {name} 调度签名不一致")
        dtypes[name] = _normalize_dispatch_dtype(
            item.get("dtype"), f"runtime manifest dispatch_signature.{name}"
        )
    for field in ("support_info_sha256", "simplified_key_sha256"):
        if not _is_sha256(binary.get(field)):
            raise ValueError(f"runtime manifest 的 {field} 非法")
    return dtypes["g"], dtypes["beta"]


def _validate_scope_contract(args: argparse.Namespace, case_count: int) -> None:
    contract = SCOPE_CONTRACTS[args.scope]
    timeout_limit = SCOPE_TIMEOUT_LIMITS[args.scope]
    if not 1 <= args.timeout <= timeout_limit:
        raise ValueError(
            f"{args.scope} timeout 必须在 1..{timeout_limit} 秒内，"
            f"实际为 {args.timeout}"
        )
    if case_count != contract["case_count"]:
        raise ValueError(
            f"{args.scope} 正式矩阵必须为 {contract['case_count']} 条，"
            f"实际为 {case_count} 条"
        )
    if args.loop_nums != contract["loop_nums"]:
        raise ValueError(
            f"{args.scope} loop_nums 必须为 {contract['loop_nums']}，"
            f"实际为 {args.loop_nums}"
        )
    if args.gm_init_mode != contract["gm_init_mode"]:
        raise ValueError(
            f"{args.scope} gm_init_mode 必须为 {contract['gm_init_mode']}，"
            f"实际为 {args.gm_init_mode}"
        )
    if args.single_process_mode != contract["single_process_mode"]:
        raise ValueError(
            f"{args.scope} single_process_mode 必须为 "
            f"{contract['single_process_mode']}，实际为 {args.single_process_mode}"
        )
    if args.tool not in contract["tools"]:
        raise ValueError(f"{args.scope} sanitizer tool 非法：{args.tool!r}")
    require_tiling_log = bool(getattr(args, "require_tiling_log", False))
    if require_tiling_log != contract["require_tiling_log"]:
        expected = "开启" if contract["require_tiling_log"] else "关闭"
        raise ValueError(f"{args.scope} require_tiling_log 必须{expected}")


def _manifest_kernel_maps(
    manifest: dict,
) -> tuple[dict[str, int], dict[int, set[str]]]:
    """读取运行时 kernel 身份及其 TilingKey 候选关系。

    一个 TilingKey 可能由多个 hashed 编译对象提供，运行时只会选择其中
    一个对象。因此这里约束 kernelName 全局唯一，但不把 key 反向限制成
    单一 kernelName。
    """
    if manifest.get("schema") != "kda-prepare-runtime/v2":
        raise ValueError("runtime manifest schema 不是 v2")
    binaries = manifest.get("kernel_binaries")
    if not isinstance(binaries, list) or not binaries:
        raise ValueError("runtime manifest 缺少 kernel_binaries")

    kernel_map: dict[str, int] = {}
    tiling_key_map: dict[int, set[str]] = {}
    dispatch_map: dict[tuple[int, str, str], str] = {}
    for binary in binaries:
        bin_file_name = str(binary.get("bin_file_name", ""))
        match = KERNEL_BIN_NAME_RE.fullmatch(bin_file_name)
        if match is None:
            raise ValueError(f"runtime manifest 的 binFileName 非法：{bin_file_name}")
        if binary.get("compile_hash") != match.group("compile_hash"):
            raise ValueError(f"runtime manifest 的编译 hash 不一致：{bin_file_name}")
        if binary.get("metadata_kernel_name") != bin_file_name:
            raise ValueError(f"runtime manifest 的 kernelName 不一致：{bin_file_name}")
        if binary.get("metadata_bin_sha256") != binary.get("bin_file_sha256"):
            raise ValueError(f"runtime manifest 的对象 SHA256 不一致：{bin_file_name}")
        gate_dtype, beta_dtype = _manifest_dispatch_signature(binary)

        kernels = binary.get("kernels")
        if not isinstance(kernels, list) or not kernels:
            raise ValueError(f"runtime manifest 的 kernelList 为空：{bin_file_name}")
        for item in kernels:
            tiling_key = int(item["tiling_key"])
            kernel_name = str(item.get("kernel_name", ""))
            if kernel_name != f"{bin_file_name}_{tiling_key}":
                raise ValueError(
                    f"runtime manifest 的完整 kernelName 不一致：{kernel_name}"
                )
            if kernel_name in kernel_map:
                raise ValueError(
                    f"runtime manifest 的 kernelName 重复：{kernel_name}"
                )
            kernel_map[kernel_name] = tiling_key
            tiling_key_map.setdefault(tiling_key, set()).add(kernel_name)
            dispatch_key = (tiling_key, gate_dtype, beta_dtype)
            previous = dispatch_map.get(dispatch_key)
            if previous is not None:
                raise ValueError(
                    "runtime manifest 的调度签名重复："
                    f"key={tiling_key}, gate={gate_dtype}, beta={beta_dtype}, "
                    f"kernels={previous},{kernel_name}"
                )
            dispatch_map[dispatch_key] = kernel_name

    if manifest.get("kernel_name_count") != len(kernel_map):
        raise ValueError("runtime manifest 的 kernelName 数量不一致")
    if manifest.get("kernel_name_sha256") != _string_digest(set(kernel_map)):
        raise ValueError("runtime manifest 的 kernelName 摘要不一致")
    return kernel_map, tiling_key_map


def _manifest_kernel_map(manifest: dict) -> dict[str, int]:
    """返回 kernelName 到 TilingKey 的唯一映射。"""
    return _manifest_kernel_maps(manifest)[0]


def _manifest_key_to_kernels(manifest: dict) -> dict[int, set[str]]:
    """返回 TilingKey 到可选 kernelName 集合的映射。"""
    return _manifest_kernel_maps(manifest)[1]


def _manifest_dispatch_map(manifest: dict) -> dict[tuple[int, str, str], str]:
    """返回 TilingKey 与 g/beta dtype 到唯一 kernelName 的映射。"""
    _manifest_kernel_maps(manifest)
    result: dict[tuple[int, str, str], str] = {}
    for binary in manifest["kernel_binaries"]:
        gate_dtype, beta_dtype = _manifest_dispatch_signature(binary)
        for item in binary["kernels"]:
            result[(int(item["tiling_key"]), gate_dtype, beta_dtype)] = str(
                item["kernel_name"]
            )
    return result


def _case_kernel_bindings(manifest: dict, cases: list[dict]) -> list[dict]:
    """把冻结 case 精确绑定到兼容其输入 dtype 的编译对象。"""
    dispatch_map = _manifest_dispatch_map(manifest)
    bindings: list[dict] = []
    seen_cases: set[int] = set()
    for case in cases:
        case_id = int(case["id"])
        if case_id in seen_cases:
            raise ValueError(f"冻结 case id 重复：{case_id}")
        seen_cases.add(case_id)
        spec = _case_spec(case)
        if "expected_tiling_key" not in spec:
            continue
        tiling_key = int(spec["expected_tiling_key"])
        gate_dtype = _normalize_dispatch_dtype(
            spec.get("gate_dtype"), f"case {case_id} gate_dtype"
        )
        beta_dtype = _normalize_dispatch_dtype(
            spec.get("beta_dtype"), f"case {case_id} beta_dtype"
        )
        dispatch_key = (tiling_key, gate_dtype, beta_dtype)
        kernel_name = dispatch_map.get(dispatch_key)
        if kernel_name is None:
            raise ValueError(
                "冻结 case 没有匹配输入 dtype 的编译对象："
                f"case={case_id}, key={tiling_key}, gate={gate_dtype}, "
                f"beta={beta_dtype}"
            )
        bindings.append(
            {
                "case_id": case_id,
                "tiling_key": tiling_key,
                "gate_dtype": gate_dtype,
                "beta_dtype": beta_dtype,
                "kernel_name": kernel_name,
            }
        )
    return bindings


def _binding_digest(bindings: list[dict]) -> str:
    normalized = sorted(
        bindings,
        key=lambda item: (
            int(item["case_id"]),
            int(item["tiling_key"]),
            str(item["gate_dtype"]),
            str(item["beta_dtype"]),
            str(item["kernel_name"]),
        ),
    )
    return _json_digest(normalized)


def _is_sha256(value: object) -> bool:
    return (
        isinstance(value, str)
        and re.fullmatch(r"[0-9a-f]{64}", value) is not None
    )


def _load_runtime_manifest(
    path: Path,
    case_file: Path,
    require_sanitizer: bool,
    require_complete_key_set: bool,
) -> dict:
    _reject_symlink_chain(path, "runtime manifest")
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise ValueError(f"runtime manifest 无法读取：{path}：{error}") from error
    if not isinstance(manifest, dict):
        raise ValueError(f"runtime manifest 必须是 JSON object：{path}")

    kernel_map = _manifest_kernel_map(manifest)
    binaries = manifest["kernel_binaries"]
    compiled_keys = set(kernel_map.values())
    if manifest.get("platform") not in {
        "ascend910b",
        "ascend910_93",
        "ascend950",
    }:
        raise ValueError("runtime manifest 的 platform 非法")
    if not _is_sha256(manifest.get("runtime_sha256")):
        raise ValueError("runtime manifest 的 runtime_sha256 非法")
    if not _is_sha256(manifest.get("op_api_sha256")):
        raise ValueError("runtime manifest 的 op_api_sha256 非法")
    _validate_runtime_toolchain(manifest, require_sanitizer)
    if manifest.get("object_count") != len(binaries):
        raise ValueError("runtime manifest 的对象数量不一致")
    if manifest.get("metadata_count") != len(binaries):
        raise ValueError("runtime manifest 的 metadata 数量不一致")
    if manifest.get("compiled_tiling_key_count") != len(compiled_keys):
        raise ValueError("runtime manifest 的 TilingKey 数量不一致")
    if manifest.get("compiled_tiling_key_sha256") != _integer_digest(
        compiled_keys
    ):
        raise ValueError("runtime manifest 的 TilingKey 摘要不一致")
    if manifest.get("complete_key_set_required") is not require_complete_key_set:
        raise ValueError("runtime manifest 的完整 TilingKey 集合要求不一致")
    if manifest.get("sanitizer_required") is not require_sanitizer:
        raise ValueError("runtime manifest 的 sanitizer 要求不一致")
    expected_sanitizer_objects = len(binaries) if require_sanitizer else 0
    if manifest.get("sanitizer_object_count") != expected_sanitizer_objects:
        raise ValueError("runtime manifest 的 sanitizer 对象数量不一致")

    for binary in binaries:
        for field in (
            "metadata_sha256",
            "metadata_bin_sha256",
            "bin_file_sha256",
        ):
            if not _is_sha256(binary.get(field)):
                raise ValueError(f"runtime manifest 的 {field} 非法")

    wrapper_hashes = manifest.get("python_wrapper_sha256")
    expected_wrappers = {
        "__init__.py",
        "ops/ascendc/__init__.py",
        "ops/ascendc/_aclnn_ctypes.py",
        "ops/ascendc/_kda_policy.py",
        "ops/ascendc/_runtime.py",
    }
    if (
        not isinstance(wrapper_hashes, dict)
        or set(wrapper_hashes) != expected_wrappers
    ):
        raise ValueError("runtime manifest 的 Python 包装器集合不闭合")
    if any(not _is_sha256(value) for value in wrapper_hashes.values()):
        raise ValueError("runtime manifest 的 Python 包装器摘要非法")
    artifact_hashes = manifest.get("test_artifact_sha256")
    if not isinstance(artifact_hashes, dict) or not artifact_hashes:
        raise ValueError("runtime manifest 缺少测试程序摘要")
    if any(not _is_sha256(value) for value in artifact_hashes.values()):
        raise ValueError("runtime manifest 的测试程序摘要非法")

    cases = _load_cases(case_file)
    _, expected_pairs = _expected_slice(cases, 0, len(cases))
    expected_keys = {key for _, key in expected_pairs}
    missing_keys = expected_keys - compiled_keys
    extra_keys = compiled_keys - expected_keys
    if missing_keys or (require_complete_key_set and extra_keys):
        raise ValueError(
            "runtime manifest 的编译 key 与冻结矩阵不一致："
            f"missing={sorted(missing_keys)}, extra={sorted(extra_keys)}"
        )
    bindings = _case_kernel_bindings(manifest, cases)
    if {int(item["tiling_key"]) for item in bindings} != expected_keys:
        raise ValueError("runtime manifest 的 case 调度签名覆盖不闭合")
    return manifest


def _candidate_opp_roots() -> list[Path]:
    roots: list[Path] = []
    op_api_lib = os.environ.get("FLA_NPU_OP_API_LIB", "").strip()
    if op_api_lib:
        raw_lib_path = Path(op_api_lib).expanduser()
        _reject_symlink_chain(raw_lib_path, "FLA_NPU_OP_API_LIB")
        lib_path = raw_lib_path.resolve()
        if not lib_path.is_file():
            raise ValueError(f"FLA_NPU_OP_API_LIB 不存在：{lib_path}")
        roots.append(lib_path.parents[2])
    for name in ("FLA_NPU_OPP_PATH", "ASCEND_CUSTOM_OPP_PATH"):
        for value in os.environ.get(name, "").split(os.pathsep):
            if value:
                raw_root = Path(value).expanduser()
                # 必须在规范化前检查原始路径及其祖先，避免把符号链接解析成可信根目录。
                _reject_symlink_chain(raw_root, f"{name} OPP 根目录")
                roots.append(raw_root)

    spec = importlib.util.find_spec("fla_npu")
    if spec is not None and spec.submodule_search_locations:
        package_root = Path(next(iter(spec.submodule_search_locations)))
        raw_opp_root = package_root / "opp"
        # package_root 及其 opp 子目录都属于运行时搜索根，解析前必须完成检查。
        _reject_symlink_chain(raw_opp_root, "Python 包 OPP 根目录")
        roots.append(raw_opp_root)

    unique: list[Path] = []
    seen: set[Path] = set()
    for root in roots:
        resolved = root.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(resolved)
    return unique


def _find_runtime_kernel_dir(soc: str) -> Path:
    relative = Path("op_impl/ai_core/tbe/kernel")
    for root in _candidate_opp_roots():
        vendor_roots = [root / "vendors/fla_npu_transformer", root]
        matches: set[Path] = set()
        for vendor_root in vendor_roots:
            kernel_root = vendor_root / relative
            _reject_symlink_chain(kernel_root, "kernel 根目录", root)
            if not kernel_root.is_dir():
                continue
            platforms = [kernel_root / soc] if soc != "auto" else kernel_root.iterdir()
            for platform in platforms:
                _reject_symlink_chain(platform, "kernel 平台目录", root)
                candidate = platform / OP_NAME
                if candidate.is_dir():
                    _reject_symlink_chain(candidate, "kernel 目录", root)
                    matches.add(candidate.resolve())
        if matches:
            if len(matches) != 1:
                raise ValueError(
                    f"当前 OPP 中找到多个 {OP_NAME} kernel 目录，请显式指定 SoC"
                )
            return matches.pop()
    raise ValueError(f"当前运行环境中找不到 {OP_NAME} kernel 目录")


def _runtime_op_api_lib() -> Path:
    value = os.environ.get("FLA_NPU_OP_API_LIB", "").strip()
    if value:
        raw_path = Path(value).expanduser()
        _reject_symlink_chain(raw_path, "FLA_NPU_OP_API_LIB")
        path = raw_path.resolve()
        if path.is_file():
            return path
        raise ValueError(f"FLA_NPU_OP_API_LIB 不存在：{path}")

    for root in _candidate_opp_roots():
        for candidate in (
            root / "op_api/lib/libcust_opapi.so",
            root / "vendors/fla_npu_transformer/op_api/lib/libcust_opapi.so",
        ):
            _reject_symlink_chain(candidate, "op_api 动态库", root)
            if candidate.is_file():
                return candidate.resolve()
    raise ValueError("当前运行环境中找不到 libcust_opapi.so")


def _runtime_manifest(
    case_file: Path,
    soc: str,
    require_sanitizer: bool,
    require_complete_key_set: bool = False,
    test_artifacts: tuple[Path, ...] = (),
) -> dict:
    kernel_dir = _find_runtime_kernel_dir(soc)
    op_api_lib = _runtime_op_api_lib()
    vendor_root = op_api_lib.parents[2]
    try:
        kernel_dir.relative_to(vendor_root)
    except ValueError as error:
        raise ValueError("op_api 与 kernel 不属于同一运行时 OPP") from error
    object_files = sorted(kernel_dir.glob("*.o"))
    metadata_files = sorted(kernel_dir.glob("*.json"))
    if not object_files or not metadata_files:
        raise ValueError(f"kernel 目录缺少 .o 或 metadata：{kernel_dir}")

    compiled_keys: set[int] = set()
    compiled_kernel_names: set[str] = set()
    referenced_objects: set[str] = set()
    kernel_binaries = []
    for path in metadata_files:
        _reject_symlink_chain(path, "kernel metadata", kernel_dir)
        metadata = json.loads(path.read_text(encoding="utf-8"))
        bin_file_name = str(metadata.get("binFileName", ""))
        bin_file_suffix = str(metadata.get("binFileSuffix", ".o"))
        bin_name_match = KERNEL_BIN_NAME_RE.fullmatch(bin_file_name)
        if bin_name_match is None:
            raise ValueError(f"kernel metadata 的 binFileName 非法：{path.name}")
        if path.stem != bin_file_name:
            raise ValueError(
                f"kernel metadata 文件名与 binFileName 不一致：{path.name}"
            )
        metadata_kernel_name = str(metadata.get("kernelName", ""))
        if metadata_kernel_name != bin_file_name:
            raise ValueError(
                f"kernel metadata 顶层 kernelName 与 binFileName 不一致：{path.name}"
            )
        dispatch_contract = _metadata_dispatch_contract(metadata, path.name)

        object_name = f"{bin_file_name}{bin_file_suffix}"
        object_path = kernel_dir / object_name
        _reject_symlink_chain(object_path, "kernel 对象", kernel_dir)
        referenced_objects.add(object_name)
        object_sha256 = _file_hash(object_path)
        metadata_bin_sha256 = str(metadata.get("sha256", "")).lower()
        if metadata_bin_sha256 != object_sha256:
            raise ValueError(
                f"kernel metadata 的 sha256 与对象不一致：{object_name}"
            )

        kernel_list = metadata.get("kernelList")
        if not isinstance(kernel_list, list) or not kernel_list:
            raise ValueError(f"kernel metadata 的 kernelList 为空：{path.name}")
        kernels = []
        binary_tiling_keys: set[int] = set()
        for item in kernel_list:
            tiling_key = int(item["tilingKey"])
            kernel_name = str(item.get("kernelName", ""))
            if kernel_name != f"{bin_file_name}_{tiling_key}":
                raise ValueError(
                    "kernelList 完整 kernelName 与 binFileName/TilingKey 不一致："
                    f"{kernel_name}, {bin_file_name}, {tiling_key}"
                )
            if kernel_name in compiled_kernel_names:
                raise ValueError(f"kernelName 重复：{kernel_name}")
            if tiling_key in binary_tiling_keys:
                raise ValueError(
                    "同一编译对象的 TilingKey 重复："
                    f"object={bin_file_name}, key={tiling_key}"
                )
            binary_tiling_keys.add(tiling_key)
            compiled_kernel_names.add(kernel_name)
            compiled_keys.add(tiling_key)
            kernels.append(
                {"kernel_name": kernel_name, "tiling_key": tiling_key}
            )
        kernel_binaries.append(
            {
                "metadata_file": path.name,
                "metadata_sha256": _file_hash(path),
                "metadata_kernel_name": metadata_kernel_name,
                "metadata_bin_sha256": metadata_bin_sha256,
                "bin_file_name": bin_file_name,
                "bin_file_suffix": bin_file_suffix,
                "bin_file_sha256": object_sha256,
                "compile_hash": bin_name_match.group("compile_hash"),
                **dispatch_contract,
                "kernels": kernels,
            }
        )
    if not compiled_keys:
        raise ValueError("kernel metadata 中没有 TilingKey")
    actual_objects = {path.name for path in object_files}
    if actual_objects != referenced_objects:
        raise ValueError(
            "kernel metadata 引用的对象集合不闭合："
            f"missing={sorted(referenced_objects - actual_objects)}, "
            f"extra={sorted(actual_objects - referenced_objects)}"
        )

    cases = _load_cases(case_file)
    _, expected_pairs = _expected_slice(cases, 0, len(cases))
    expected_keys = {key for _, key in expected_pairs}
    missing_keys = expected_keys - compiled_keys
    extra_keys = compiled_keys - expected_keys
    if missing_keys or (require_complete_key_set and extra_keys):
        raise ValueError(
            "安装包编译 key 与冻结矩阵不一致："
            f"missing={sorted(missing_keys)}, extra={sorted(extra_keys)}"
        )

    digest = hashlib.sha256()
    for path in object_files + metadata_files:
        digest.update(path.name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    digest.update(b"op_api/lib/libcust_opapi.so\0")
    digest.update(op_api_lib.read_bytes())
    digest.update(b"\0")

    test_artifact_hashes = {}
    for path in test_artifacts:
        _reject_symlink_chain(path, "测试程序")
        resolved = path.resolve()
        test_artifact_hashes[resolved.name] = _file_hash(resolved)
        digest.update(f"test/{resolved.name}\0".encode("utf-8"))
        digest.update(resolved.read_bytes())
        digest.update(b"\0")

    wrapper_hashes = {}
    package_spec = importlib.util.find_spec("fla_npu")
    if package_spec is None or not package_spec.submodule_search_locations:
        raise ValueError("找不到实际加载的 fla_npu Python 包")
    raw_package_root = Path(next(iter(package_spec.submodule_search_locations)))
    _reject_symlink_chain(raw_package_root, "Python 包")
    package_root = raw_package_root.resolve()
    for relative in (
        "__init__.py",
        "ops/ascendc/__init__.py",
        "ops/ascendc/_aclnn_ctypes.py",
        "ops/ascendc/_kda_policy.py",
        "ops/ascendc/_runtime.py",
    ):
        path = package_root / relative
        _reject_symlink_chain(path, "Python 包装器", package_root)
        wrapper_hashes[relative] = _file_hash(path)
        digest.update(f"python/{relative}\0".encode("utf-8"))
        digest.update(path.read_bytes())
        digest.update(b"\0")

    toolchain = _runtime_toolchain(require_sanitizer)
    digest.update(b"toolchain\0")
    digest.update(
        json.dumps(
            toolchain,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    )
    digest.update(b"\0")

    sanitizer_objects = 0
    if require_sanitizer:
        for path in object_files:
            result = subprocess.run(
                ("nm", str(path)),
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                errors="replace",
            )
            if result.returncode != 0 or not _has_sanitizer_symbol(
                result.stdout, soc
            ):
                raise ValueError(f"kernel 对象缺少 sanitizer 符号：{path.name}")
            sanitizer_objects += 1

    manifest = {
        "schema": "kda-prepare-runtime/v2",
        "platform": kernel_dir.parent.name,
        "runtime_sha256": digest.hexdigest(),
        "op_api_sha256": _file_hash(op_api_lib),
        "test_artifact_sha256": test_artifact_hashes,
        "python_wrapper_sha256": wrapper_hashes,
        "toolchain": toolchain,
        "object_count": len(object_files),
        "metadata_count": len(metadata_files),
        "compiled_tiling_key_count": len(compiled_keys),
        "compiled_tiling_key_sha256": _integer_digest(compiled_keys),
        "kernel_name_count": len(compiled_kernel_names),
        "kernel_name_sha256": _string_digest(compiled_kernel_names),
        "kernel_binaries": kernel_binaries,
        "complete_key_set_required": require_complete_key_set,
        "sanitizer_required": require_sanitizer,
        "sanitizer_object_count": sanitizer_objects,
    }
    bindings = _case_kernel_bindings(manifest, cases)
    if {int(item["tiling_key"]) for item in bindings} != expected_keys:
        raise ValueError("安装包的 case 调度签名覆盖不闭合")
    return manifest


def _has_sanitizer_symbol(nm_output: str, soc: str) -> bool:
    """识别各架构 sanitizer 编译对象中的运行时符号。"""
    lowered = nm_output.lower()
    if "sanitizer" in lowered:
        return True
    return soc == "ascend950" and "__mstx_dfx_report_stub" in lowered


def _verify_current_runtime(args: argparse.Namespace) -> dict:
    require_sanitizer = args.scope == "mssanitizer"
    require_complete_key_set = args.scope != "accuracy"
    existing = _load_runtime_manifest(
        args.runtime_manifest,
        args.case_file,
        require_sanitizer,
        require_complete_key_set,
    )
    actual = _runtime_manifest(
        args.case_file,
        args.soc,
        require_sanitizer,
        require_complete_key_set,
        tuple(args.test_artifact),
    )
    if existing != actual:
        raise ValueError(
            "runtime manifest 与当前实际 OPP、包装器或测试程序不一致"
        )
    return existing


def verify_runtime(args: argparse.Namespace) -> int:
    _reject_symlink_chain(args.output, "runtime manifest")
    manifest = _runtime_manifest(
        args.case_file,
        args.soc,
        args.require_sanitizer,
        args.require_complete_key_set,
        tuple(args.test_artifact),
    )
    if args.output.exists():
        existing = json.loads(args.output.read_text(encoding="utf-8"))
        if existing != manifest:
            raise ValueError(
                "运行时 kernel 已变化，拒绝在原矩阵目录续跑："
                f"existing={existing}, actual={manifest}"
            )
    else:
        args.output.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


def _find_report(root: Path, scope: str) -> Path:
    _reject_symlink_chain(root, "ATK 报告目录")
    patterns = {
        "accuracy": (
            f"accuracy/atk_output/atk_{OP_NAME}_*/report/*.xlsx",
            f"accuracy/atk_{OP_NAME}_*/report/*.xlsx",
        ),
        "determinism": (
            f"determinism/atk_output/atk_{OP_NAME}_*/report/*.xlsx",
            f"determinism/atk_{OP_NAME}_*/report/*.xlsx",
            f"atk_{OP_NAME}_*/report/*.xlsx",
        ),
        "mssanitizer": (
            f"mssanitizer/atk_output/atk_{OP_NAME}_*/report/*.xlsx",
            f"mssanitizer/atk_{OP_NAME}_*/report/*.xlsx",
            f"atk_{OP_NAME}_*/report/*.xlsx",
        ),
    }
    paths: list[Path] = []
    for pattern in patterns[scope]:
        paths.extend(Path(value) for value in glob.glob(str(root / pattern)))
    if not paths:
        raise ValueError(f"{root} 下没有找到 {scope} 报告")
    for path in set(paths):
        _reject_symlink_chain(path, "ATK 报告")
    return max(set(paths), key=lambda path: path.stat().st_mtime_ns)


def _report_case_ids(report: Path, scope: str) -> tuple[list[int], int]:
    _reject_symlink_chain(report, "ATK 报告")
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("读取 ATK 报告需要安装 openpyxl") from error

    workbook = openpyxl.load_workbook(report, read_only=True, data_only=True)
    try:
        if "statistic" not in workbook.sheetnames:
            raise ValueError(f"报告缺少 statistic sheet：{report.name}")
        if "summary" not in workbook.sheetnames:
            raise ValueError(f"报告缺少 summary sheet：{report.name}")

        statistic_rows = list(workbook["statistic"].iter_rows(values_only=True))
        if not statistic_rows:
            raise ValueError(f"statistic sheet 为空：{report.name}")
        statistic_header = {
            str(value).strip(): index
            for index, value in enumerate(statistic_rows[0])
            if value is not None
        }
        if "编号" not in statistic_header or "运行结果" not in statistic_header:
            raise ValueError(f"statistic sheet 缺少编号或运行结果列：{report.name}")

        case_ids = []
        memory_columns = [
            index
            for name, index in statistic_header.items()
            if name == "内存检测通过" or name.endswith("_内存检测通过")
        ]
        accuracy_columns = [
            index
            for name, index in statistic_header.items()
            if name == "精度通过" or name.endswith("_精度通过")
        ]
        determinism_columns = [
            index
            for name, index in statistic_header.items()
            if name == "确定性计算结果" or name.endswith("_确定性计算结果")
        ]
        if scope == "mssanitizer" and not memory_columns:
            raise ValueError(f"statistic sheet 缺少内存检测通过列：{report.name}")
        if scope == "accuracy" and not accuracy_columns:
            raise ValueError(f"statistic sheet 缺少精度通过列：{report.name}")
        if scope == "determinism" and not determinism_columns:
            raise ValueError(f"statistic sheet 缺少确定性计算结果列：{report.name}")

        id_column = statistic_header["编号"]
        run_column = statistic_header["运行结果"]
        for row in statistic_rows[1:]:
            case_id = row[id_column] if id_column < len(row) else None
            if case_id is None:
                continue
            case_ids.append(int(case_id))
            run_result = row[run_column] if run_column < len(row) else None
            if str(run_result).strip().upper() != "SUCCESS":
                raise ValueError(f"case {case_id} 运行结果不是 SUCCESS")
            checked_columns = memory_columns if scope == "mssanitizer" else []
            if scope == "accuracy":
                checked_columns = accuracy_columns
            elif scope == "determinism":
                checked_columns = determinism_columns
            checked_values = [
                row[column] if column < len(row) else None
                for column in checked_columns
            ]
            active_values = [
                value for value in checked_values if value not in (None, "")
            ]
            if checked_columns and (
                not active_values
                or any(
                    value is not True and str(value).strip().lower() != "true"
                    for value in active_values
                )
            ):
                raise ValueError(f"case {case_id} 的检查结果不是全部 True")

        summary_rows = list(workbook["summary"].iter_rows(values_only=True))
        if not summary_rows:
            raise ValueError(f"summary sheet 为空：{report.name}")
        summary_header = {
            str(value).strip(): index
            for index, value in enumerate(summary_rows[0])
            if value is not None
        }
        required = {"总用例数"}
        if scope in ("accuracy", "determinism"):
            required.update(
                {
                    "执行成功用例个数",
                    "执行失败用例个数",
                    "通过用例个数",
                    "精度是否达标",
                }
            )
            if scope == "determinism":
                required.add("确定性计算是否达标")
        else:
            required.update({"内存检测通过率", "内存检测是否达标"})
        missing = sorted(required - summary_header.keys())
        if missing:
            raise ValueError(f"summary sheet 缺少列 {missing}：{report.name}")

        total = 0
        passed = 0
        for row in summary_rows[1:]:
            total_value = row[summary_header["总用例数"]]
            if total_value is None:
                continue
            row_total = int(total_value)
            total += row_total
            if scope in ("accuracy", "determinism"):
                executed = int(row[summary_header["执行成功用例个数"]])
                failed = int(row[summary_header["执行失败用例个数"]])
                row_passed = int(row[summary_header["通过用例个数"]])
                if (
                    executed != row_total
                    or failed != 0
                    or row_passed != row_total
                ):
                    raise ValueError(f"报告存在执行或检查失败：{report.name}")
                accuracy_status = str(
                    row[summary_header["精度是否达标"]]
                ).strip()
                if scope == "accuracy" and accuracy_status != "Pass":
                    raise ValueError(f"精度达标列失败：{report.name}")
                # 确定性报告的精度列由 ATK 标为不适用；真正的确定性
                # 结论由下方“确定性计算是否达标”列严格校验。
                if scope == "determinism" and accuracy_status not in {
                    "Pass",
                    "-",
                }:
                    raise ValueError(f"精度达标列失败：{report.name}")
                if scope == "determinism" and str(
                    row[summary_header["确定性计算是否达标"]]
                ).strip() != "Pass":
                    raise ValueError(f"确定性达标列失败：{report.name}")
                passed += row_passed
            else:
                rate = float(
                    str(row[summary_header["内存检测通过率"]]).strip().rstrip("%")
                )
                status = str(row[summary_header["内存检测是否达标"]]).strip()
                if rate != 100.0 or status != "Pass":
                    raise ValueError(f"内存检测未全部通过：{report.name}")
                passed += row_total
        if total == 0 or passed != total:
            raise ValueError(
                f"报告通过数不匹配：total={total}, passed={passed}"
            )
        return case_ids, total
    finally:
        workbook.close()


def _logged_keys(log_path: Path) -> tuple[set[int], set[int]]:
    _reject_symlink_chain(log_path, "console 日志")
    host_keys: set[int] = set()
    launch_keys: set[int] = set()
    with log_path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            host_match = HOST_KEY_RE.search(line)
            if host_match:
                output_mode = int(host_match.group(1))
                tiling_key = int(host_match.group(2))
                encoded_output_mode = (tiling_key >> 23) & 0x3
                if output_mode != encoded_output_mode:
                    raise ValueError(
                        "Host 日志中的 outputMode 与 TilingKey 编码不一致："
                        f"mode={output_mode}, key={tiling_key}"
                    )
                host_keys.add(tiling_key)
            launch_match = LAUNCH_KEY_RE.search(line)
            if launch_match:
                launch_keys.add(int(launch_match.group(1)))
    return host_keys, launch_keys


def _read_sanitizer_logs(log_paths: tuple[Path, ...]) -> str:
    contents = []
    seen: set[Path] = set()
    for path in log_paths:
        _reject_symlink_chain(path, "sanitizer 证据文件")
        resolved = path.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        if not path.is_file() or path.stat().st_size == 0:
            raise ValueError(f"sanitizer 证据文件不存在或为空：{path}")
        contents.append(path.read_text(encoding="utf-8", errors="replace"))
    return "\n".join(contents)


def _sanitizer_diagnostic(evidence: str, tool: str) -> str | None:
    """返回第一条 sanitizer 异常；普通业务 ERROR 不作为异常。"""
    inactive_re = re.compile(
        rf"No\s+active\s+sanitizer\s+tool\s+on\s+kernel[^\r\n]*"
        rf"{KERNEL_NAME_PATTERN}",
        re.IGNORECASE,
    )
    inactive = inactive_re.search(evidence)
    if inactive is not None:
        raise ValueError(f"{tool} 日志显示目标 kernel 未加载 sanitizer 插桩")

    register_matches = list(SANITIZER_REGISTER_RE.finditer(evidence))
    if register_matches:
        return evidence[
            evidence.rfind("\n", 0, register_matches[0].start()) + 1 :
            evidence.find("\n", register_matches[0].end())
            if evidence.find("\n", register_matches[0].end()) >= 0
            else len(evidence)
        ].strip()

    for diagnostic_re in (
        *SANITIZER_COMMON_DIAGNOSTICS,
        *SANITIZER_TOOL_DIAGNOSTICS[tool],
        SANITIZER_GENERIC_DIAGNOSTIC,
        SANITIZER_ERROR_MARKER,
    ):
        diagnostic = diagnostic_re.search(evidence)
        if diagnostic is not None:
            line_start = evidence.rfind("\n", 0, diagnostic.start()) + 1
            line_end = evidence.find("\n", diagnostic.end())
            if line_end < 0:
                line_end = len(evidence)
            return evidence[line_start:line_end].strip()
    return None


def _parse_atk_version(output: str) -> str:
    versions = [
        line.strip()
        for line in output.splitlines()
        if re.fullmatch(r"[0-9]+\.[0-9]+\.[0-9]+", line.strip())
    ]
    if len(versions) != 1:
        raise ValueError("无法从 atk --version 唯一解析版本")
    return versions[0]


def _parse_mssanitizer_revisions(output: str) -> tuple[str, str]:
    sanitizer = re.findall(r"^\s*mssanitizer\s+(\S+)\s*$", output, re.MULTILINE)
    common = re.findall(r"^\s*msopscommon\s+(\S+)\s*$", output, re.MULTILINE)
    if len(sanitizer) != 1 or len(common) != 1:
        raise ValueError(
            "无法从 mssanitizer --version 唯一解析 mssanitizer/msopscommon revision"
        )
    return sanitizer[0], common[0]


def _tool_version_output(command: tuple[str, ...], label: str) -> str:
    try:
        result = subprocess.run(
            command,
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            errors="replace",
            timeout=10,
        )
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ValueError(f"无法执行 {label} 版本查询：{error}") from error
    if result.returncode != 0 or not result.stdout.strip():
        raise ValueError(f"{label} 版本查询失败，返回码 {result.returncode}")
    return result.stdout


def _tool_executable_sha256(name: str) -> str:
    executable = shutil.which(name)
    if not executable:
        raise ValueError(f"找不到测试工具：{name}")
    resolved = Path(executable).resolve()
    if not resolved.is_file():
        raise ValueError(f"测试工具不是普通文件：{name}")
    return _file_hash(resolved)


def _runtime_toolchain(require_sanitizer: bool) -> dict:
    atk_output = _tool_version_output(("atk", "--version"), "atk")
    toolchain = {
        "atk": {
            "version": _parse_atk_version(atk_output),
            "executable_sha256": _tool_executable_sha256("atk"),
        }
    }
    if require_sanitizer:
        mss_output = _tool_version_output(
            ("mssanitizer", "--version"), "mssanitizer"
        )
        mss_revision, common_revision = _parse_mssanitizer_revisions(
            mss_output
        )
        toolchain["mssanitizer"] = {
            "revision": mss_revision,
            "msopscommon_revision": common_revision,
            "executable_sha256": _tool_executable_sha256("mssanitizer"),
        }
    return toolchain


def _validate_runtime_toolchain(manifest: dict, require_sanitizer: bool) -> None:
    toolchain = manifest.get("toolchain")
    expected_names = {"atk"}
    if require_sanitizer:
        expected_names.add("mssanitizer")
    if not isinstance(toolchain, dict) or set(toolchain) != expected_names:
        raise ValueError("runtime manifest 的测试工具集合不闭合")
    atk = toolchain.get("atk")
    if (
        not isinstance(atk, dict)
        or set(atk) != {"version", "executable_sha256"}
        or re.fullmatch(r"[0-9]+\.[0-9]+\.[0-9]+", str(atk.get("version", "")))
        is None
        or not _is_sha256(atk.get("executable_sha256"))
    ):
        raise ValueError("runtime manifest 的 ATK 工具身份非法")
    if require_sanitizer:
        mss = toolchain.get("mssanitizer")
        if (
            not isinstance(mss, dict)
            or set(mss)
            != {"revision", "msopscommon_revision", "executable_sha256"}
            or not str(mss.get("revision", ""))
            or not str(mss.get("msopscommon_revision", ""))
            or not _is_sha256(mss.get("executable_sha256"))
        ):
            raise ValueError("runtime manifest 的 mssanitizer 工具身份非法")


def _sanitizer_log_summary(
    log_paths: tuple[Path, ...],
    tool: str,
    expected_finish_count: int,
    expected_kernels: set[str] | None = None,
) -> dict:
    """校验统一 runner 的原始 mssanitizer 日志，不依赖 ATK xlsx。"""
    if tool not in SANITIZER_TOOLS:
        raise ValueError(f"非法 sanitizer 工具：{tool}")
    if expected_finish_count <= 0:
        raise ValueError("expected_finish_count 必须为正数")
    evidence = _read_sanitizer_logs(log_paths)
    diagnostic = _sanitizer_diagnostic(evidence, tool)
    if diagnostic is not None:
        raise ValueError(f"{tool} 检测到 sanitizer 异常：{diagnostic}")

    start_re = re.compile(
        rf"Start\s+{re.escape(tool)}\s+sanitizer\s+on\s+kernel\s+"
        rf"(?P<kernel>{KERNEL_NAME_PATTERN}(?:_[A-Za-z0-9]+)+)",
        re.IGNORECASE,
    )
    start_counts = Counter(start_re.findall(evidence))
    started = set(start_counts)
    duplicate_starts = {
        name: count for name, count in start_counts.items() if count != 1
    }
    if duplicate_starts:
        raise ValueError(
            f"{tool} sanitizer Start 次数不为 1：{duplicate_starts}"
        )
    finish_matches = list(SANITIZER_FINISH_RE.finditer(evidence))
    if not finish_matches:
        raise ValueError(f"{tool} 日志缺少 sanitizer 完成状态")
    finish_statuses: dict[str, set[str]] = {}
    finish_counts: Counter[str] = Counter()
    for match in finish_matches:
        kernel_name = match.group("kernel")
        finish_counts[kernel_name] += 1
        finish_statuses.setdefault(kernel_name, set()).add(
            match.group("status").lower()
        )
    duplicate_finishes = {
        name: count for name, count in finish_counts.items() if count != 1
    }
    if duplicate_finishes:
        raise ValueError(
            f"{tool} sanitizer Finish 次数不为 1：{duplicate_finishes}"
        )
    failed = {
        name: sorted(statuses)
        for name, statuses in finish_statuses.items()
        if statuses != {"no error detected."}
    }
    if failed:
        raise ValueError(f"{tool} 目标 kernel 的 sanitizer 完成状态失败：{failed}")
    if len(started) != expected_finish_count:
        raise ValueError(
            f"{tool} sanitizer 启动 kernel 数量不匹配："
            f"expected={expected_finish_count}, actual={len(started)}"
        )
    finished = set(finish_statuses)
    if len(finished) != expected_finish_count:
        raise ValueError(
            f"{tool} sanitizer clean finish 数量不匹配："
            f"expected={expected_finish_count}, actual={len(finished)}"
        )
    if started != finished:
        raise ValueError(
            f"{tool} sanitizer Start/Finish kernel 集合不一致："
            f"missing_finish={sorted(started - finished)}, "
            f"extra_finish={sorted(finished - started)}"
        )
    if expected_kernels is not None:
        unexpected_finished = finished - expected_kernels
        missing_finished = expected_kernels - finished
        unexpected_started = started - expected_kernels
        missing_started = expected_kernels - started
        if (
            unexpected_finished
            or missing_finished
            or unexpected_started
            or missing_started
        ):
            raise ValueError(
                f"{tool} sanitizer kernel 集合不匹配："
                f"missing_finish={sorted(missing_finished)}, "
                f"extra_finish={sorted(unexpected_finished)}, "
                f"missing_start={sorted(missing_started)}, "
                f"extra_start={sorted(unexpected_started)}"
            )
    missing_finish = started - finished
    if missing_finish:
        raise ValueError(
            f"{tool} sanitizer 启动 kernel 缺少 clean finish："
            f"{sorted(missing_finish)}"
        )
    return {
        "tool": tool,
        "log_count": len(log_paths),
        "started_kernel_count": len(started),
        "clean_finish_count": len(finished),
        "kernel_names": sorted(finished),
        "passed": True,
    }


def _kernel_name_to_key(
    key_to_kernels: dict[int, set[str]],
) -> dict[str, int]:
    """把允许多候选的 key 映射展开为 kernelName 到 key。"""
    name_to_key: dict[str, int] = {}
    for tiling_key, names in key_to_kernels.items():
        for name in names:
            previous = name_to_key.get(name)
            if previous is not None and previous != tiling_key:
                raise ValueError(
                    "runtime manifest 的 kernelName 关联多个 TilingKey："
                    f"kernel={name}, keys={previous},{tiling_key}"
                )
            name_to_key[name] = tiling_key
    return name_to_key


def _validate_kernel_hits(
    started_kernel_names: set[str],
    expected_keys: set[int],
    key_to_kernels: dict[int, set[str]],
    label: str,
) -> tuple[set[int], dict[int, str]]:
    """校验每个期望 key 恰好命中一个编译候选。"""
    name_to_key = _kernel_name_to_key(key_to_kernels)
    unknown = started_kernel_names - set(name_to_key)
    if unknown:
        raise ValueError(
            f"{label} 命中了 runtime manifest 之外的 kernel：{sorted(unknown)}"
        )
    key_hits: dict[int, set[str]] = {}
    for name in started_kernel_names:
        key_hits.setdefault(name_to_key[name], set()).add(name)
    started_keys = set(key_hits)
    if started_keys != expected_keys:
        raise ValueError(
            f"{label} 启动 TilingKey 不闭合："
            f"missing={sorted(expected_keys - started_keys)}, "
            f"extra={sorted(started_keys - expected_keys)}"
        )
    ambiguous = {
        tiling_key: sorted(names)
        for tiling_key, names in key_hits.items()
        if len(names) != 1
    }
    if ambiguous:
        raise ValueError(
            f"{label} 同一 TilingKey 命中了多个 kernel 候选：{ambiguous}"
        )
    return started_keys, {
        tiling_key: next(iter(names)) for tiling_key, names in key_hits.items()
    }


def _sanitizer_evidence(
    console_log: Path,
    sanitizer_log: Path,
    tool: str,
    expected_keys: set[int],
    runtime_manifest: dict,
    expected_cases: list[dict] | None = None,
) -> tuple[int, set[int], set[str], str, str]:
    if tool not in SANITIZER_TOOLS:
        raise ValueError(f"非法 sanitizer 工具：{tool}")

    start_re = re.compile(
        rf"Start\s+{re.escape(tool)}\s+sanitizer\s+on\s+kernel\s+"
        rf"(?P<kernel>{KERNEL_NAME_PATTERN}(?:_[A-Za-z0-9]+)+)"
        rf"(?=$|[^A-Za-z0-9_])",
        re.IGNORECASE,
    )
    finish_re = SANITIZER_FINISH_RE
    register_re = SANITIZER_REGISTER_RE
    inactive_re = re.compile(
        rf"No\s+active\s+sanitizer\s+tool\s+on\s+kernel[^\r\n]*"
        rf"{KERNEL_NAME_PATTERN}",
        re.IGNORECASE,
    )
    console_evidence = _read_sanitizer_logs((console_log,))
    sanitizer_evidence = _read_sanitizer_logs((sanitizer_log,))
    evidence = f"{console_evidence}\n{sanitizer_evidence}"
    # -msl 是权威事件源；控制台可能只镜像其中一部分，但不能出现额外事件。
    def _event_signature(source: str) -> tuple[Counter[str], Counter[tuple[str, str]]]:
        starts = Counter(start_re.findall(source))
        finishes = Counter(
            (match.group("kernel"), match.group("status").lower())
            for match in finish_re.finditer(source)
        )
        return starts, finishes

    console_events = _event_signature(console_evidence)
    sanitizer_events = _event_signature(sanitizer_evidence)
    console_has_events = bool(console_events[0] or console_events[1])
    sanitizer_has_events = bool(sanitizer_events[0] or sanitizer_events[1])
    if sanitizer_has_events:
        if console_has_events:
            starts_are_subset = not (
                console_events[0] - sanitizer_events[0]
            )
            finishes_are_subset = not (
                console_events[1] - sanitizer_events[1]
            )
            if not starts_are_subset or not finishes_are_subset:
                raise ValueError(
                    f"{tool} console 与 sanitizer 日志的 Start/Finish 记录不一致"
                )
        event_evidence = sanitizer_evidence
    else:
        event_evidence = console_evidence
    if inactive_re.search(evidence):
        raise ValueError(f"{tool} 日志显示目标 kernel 未加载 sanitizer 插桩")
    if expected_cases is None:
        key_to_kernels = _manifest_key_to_kernels(runtime_manifest)
    else:
        bindings = _case_kernel_bindings(runtime_manifest, expected_cases)
        if {int(item["tiling_key"]) for item in bindings} != expected_keys:
            raise ValueError("sanitizer case 调度签名与期望 TilingKey 不闭合")
        key_to_kernels = {
            int(item["tiling_key"]): {str(item["kernel_name"])}
            for item in bindings
        }
    register_matches = list(register_re.finditer(evidence))
    if register_matches:
        names = {match.group("kernel") for match in register_matches}
        unexpected = names - set(_kernel_name_to_key(key_to_kernels))
        if unexpected:
            raise ValueError(
                f"{tool} 检测到非目标 kernel 的寄存器状态异常：{sorted(unexpected)}"
            )
        raise ValueError(f"{tool} 检测到目标 kernel 的寄存器状态异常")
    for diagnostic_re in (
        *SANITIZER_COMMON_DIAGNOSTICS,
        *SANITIZER_TOOL_DIAGNOSTICS[tool],
        SANITIZER_GENERIC_DIAGNOSTIC,
        SANITIZER_ERROR_MARKER,
    ):
        diagnostic = diagnostic_re.search(evidence)
        if diagnostic is not None:
            line_start = evidence.rfind("\n", 0, diagnostic.start()) + 1
            line_end = evidence.find("\n", diagnostic.end())
            if line_end < 0:
                line_end = len(evidence)
            detail = evidence[line_start:line_end].strip()
            raise ValueError(f"{tool} 检测到 sanitizer 异常：{detail}")
    matched_kernel_names = start_re.findall(event_evidence)
    start_counts = Counter(matched_kernel_names)
    duplicate_starts = {
        name: count for name, count in start_counts.items() if count != 1
    }
    if duplicate_starts:
        raise ValueError(
            f"{tool} sanitizer Start 次数不为 1：{duplicate_starts}"
        )
    started_kernel_names = set(matched_kernel_names)
    started_keys, _ = _validate_kernel_hits(
        started_kernel_names,
        expected_keys,
        key_to_kernels,
        f"{tool} sanitizer",
    )

    finish_statuses: dict[str, set[str]] = {}
    finish_counts: Counter[str] = Counter()
    for match in finish_re.finditer(event_evidence):
        kernel_name = match.group("kernel")
        finish_counts[kernel_name] += 1
        finish_statuses.setdefault(kernel_name, set()).add(
            match.group("status").lower()
        )
    duplicate_finishes = {
        name: count for name, count in finish_counts.items() if count != 1
    }
    if duplicate_finishes:
        raise ValueError(
            f"{tool} sanitizer Finish 次数不为 1：{duplicate_finishes}"
        )
    unexpected_finished = set(finish_statuses) - started_kernel_names
    if unexpected_finished:
        raise ValueError(
            f"{tool} 出现非目标 kernel 的 sanitizer 完成状态："
            f"{sorted(unexpected_finished)}"
        )
    missing_finished = started_kernel_names - set(finish_statuses)
    if missing_finished:
        raise ValueError(
            f"{tool} 缺少目标 kernel 的 sanitizer 完成状态："
            f"{sorted(missing_finished)}"
        )
    failed_finished = {
        name: sorted(statuses)
        for name, statuses in finish_statuses.items()
        if statuses != {"no error detected."}
    }
    if failed_finished:
        raise ValueError(
            f"{tool} 目标 kernel 的 sanitizer 完成状态失败：{failed_finished}"
        )
    return (
        len(matched_kernel_names),
        started_keys,
        started_kernel_names,
        _file_hash(console_log),
        _file_hash(sanitizer_log),
    )


def _expected_slice(
    cases: list[dict], start: int, end: int
) -> tuple[list[int], list[tuple[int, int]]]:
    if not (0 <= start < end <= len(cases)):
        raise ValueError(f"非法分片范围 {start}:{end}，总用例数 {len(cases)}")
    selected = cases[start:end]
    ids = [int(case["id"]) for case in selected]
    if ids != list(range(start, end)):
        raise ValueError(f"冻结用例 ID 与位置不一致：{ids[:3]}...{ids[-3:]}")
    pairs = []
    for case in selected:
        spec = _case_spec(case)
        if "expected_tiling_key" in spec:
            pairs.append((int(case["id"]), int(spec["expected_tiling_key"])))
    return ids, pairs


def _build_shard_summary(args: argparse.Namespace) -> dict:
    cases = _load_cases(args.case_file)
    _validate_scope_contract(args, len(cases))
    if args.scope == "accuracy":
        _validate_accuracy_seed_contract(cases)
    runtime_manifest = _load_runtime_manifest(
        args.runtime_manifest,
        args.case_file,
        args.scope == "mssanitizer",
        args.scope != "accuracy",
    )
    _reject_symlink_chain(args.shard_root, "分片目录")
    shard_root = args.shard_root.resolve()
    if not shard_root.is_dir():
        raise ValueError(f"分片目录不存在、不是目录或为符号链接：{args.shard_root}")
    _reject_symlink_chain(args.console_log, "console 日志")
    _reject_symlink_chain(args.summary, "分片摘要")
    if args.console_log.resolve() != shard_root / "console.log":
        raise ValueError("console 日志必须是分片目录下的 console.log")
    if args.summary.resolve() != shard_root / "summary.json":
        raise ValueError("summary 必须是分片目录下的 summary.json")

    expected_ids, expected_pairs = _expected_slice(cases, args.start, args.end)
    if args.end - args.start != 1:
        raise ValueError("正式矩阵每个分片必须只包含一条 case")
    report = _find_report(args.shard_root, args.scope)
    _reject_symlink_chain(report, "ATK 报告")
    try:
        report.resolve().relative_to(shard_root)
    except ValueError as error:
        raise ValueError(f"ATK 报告越出分片目录：{report}") from error
    report_ids, report_total = _report_case_ids(report, args.scope)
    if sorted(report_ids) != expected_ids or len(set(report_ids)) != len(report_ids):
        raise ValueError(
            f"报告 case ID 不闭合：expected={expected_ids}, actual={sorted(report_ids)}"
        )
    if report_total != len(expected_ids):
        raise ValueError(
            f"报告总数不匹配：expected={len(expected_ids)}, actual={report_total}"
        )

    expected_keys = {key for _, key in expected_pairs}
    host_keys: set[int] = set()
    launch_keys: set[int] = set()
    if args.require_tiling_log:
        host_keys, launch_keys = _logged_keys(args.console_log)
        if host_keys != expected_keys:
            raise ValueError(
                f"Host TilingKey 不匹配：missing={sorted(expected_keys - host_keys)}, "
                f"extra={sorted(host_keys - expected_keys)}"
            )
        if launch_keys != expected_keys:
            raise ValueError(
                f"Launch TilingKey 不匹配：missing={sorted(expected_keys - launch_keys)}, "
                f"extra={sorted(launch_keys - expected_keys)}"
            )

    sanitizer_start_count = 0
    sanitizer_started_keys: set[int] = set()
    sanitizer_started_kernel_names: set[str] = set()
    sanitizer_started_kernel_bindings: list[dict] = []
    sanitizer_log_sha256 = ""
    sanitizer_outer_log_sha256 = ""
    console_log_sha256 = _file_hash(args.console_log)
    if args.scope == "mssanitizer":
        sanitizer_started_kernel_bindings = _case_kernel_bindings(
            runtime_manifest, cases[args.start : args.end]
        )
        sanitizer_log = args.shard_root / f"{args.tool}.log"
        _reject_symlink_chain(sanitizer_log, "ATK sanitizer 日志")
        sanitizer_outer_log = getattr(args, "sanitizer_log", None)
        if sanitizer_outer_log is None:
            sanitizer_outer_log = sanitizer_log
        _reject_symlink_chain(sanitizer_outer_log, "外层 sanitizer 日志")
        sanitizer_outer_log = sanitizer_outer_log.resolve()
        if sanitizer_outer_log != sanitizer_log.resolve():
            raise ValueError(
                "外层 sanitizer 日志必须与 ATK -msl 日志使用同一文件"
            )
        try:
            sanitizer_outer_log.relative_to(shard_root)
        except ValueError as error:
            raise ValueError(
                f"外层 sanitizer 日志必须位于分片目录内：{sanitizer_outer_log}"
            ) from error
        key_to_kernels = _manifest_key_to_kernels(runtime_manifest)
        missing_candidates = expected_keys - set(key_to_kernels)
        if missing_candidates:
            raise ValueError(
                "外层 sanitizer 日志无法绑定到当前分片的目标 kernel："
                f"missing_keys={sorted(missing_candidates)}"
            )
        # 同一份原始日志同时承载 ATK -msl 和外层 mssanitizer 记录；即使
        # ATK 报告误报通过，也必须有实际选中候选的 clean finish 才能续跑。
        _sanitizer_log_summary(
            (sanitizer_outer_log,),
            args.tool,
            expected_finish_count=len(expected_keys),
        )
        sanitizer_outer_log_sha256 = _file_hash(sanitizer_outer_log)
        (
            sanitizer_start_count,
            sanitizer_started_keys,
            sanitizer_started_kernel_names,
            console_log_sha256,
            sanitizer_log_sha256,
        ) = _sanitizer_evidence(
            args.console_log,
            sanitizer_log,
            args.tool,
            expected_keys,
            runtime_manifest,
            cases[args.start : args.end],
        )

    return {
        "schema": "kda-prepare-atk-shard/v2",
        "scope": args.scope,
        "tool": args.tool,
        "timeout_seconds": args.timeout,
        "loop_nums": args.loop_nums,
        "gm_init_mode": args.gm_init_mode,
        "single_process_mode": args.single_process_mode,
        "require_tiling_log": args.require_tiling_log,
        "runtime_manifest_sha256": _file_hash(args.runtime_manifest),
        "start": args.start,
        "end": args.end,
        "expected_cases": len(expected_ids),
        "case_ids": expected_ids,
        "case_json_sha256": _case_hash(args.case_file),
        "expected_tiling_keys": sorted(expected_keys),
        "host_tiling_keys": sorted(host_keys),
        "launch_tiling_keys": sorted(launch_keys),
        "key_pair_sha256": _key_digest(expected_pairs),
        "report": str(report.relative_to(args.shard_root)),
        "report_sha256": _file_hash(report),
        "console_log_sha256": console_log_sha256,
        "sanitizer_log_sha256": sanitizer_log_sha256,
        **(
            {
                "sanitizer_outer_log": str(sanitizer_outer_log.relative_to(shard_root)),
                "sanitizer_outer_log_sha256": sanitizer_outer_log_sha256,
            }
            if args.scope == "mssanitizer"
            else {}
        ),
        "sanitizer_start_count": sanitizer_start_count,
        "sanitizer_started_tiling_keys": sorted(sanitizer_started_keys),
        "sanitizer_started_kernel_names": sorted(
            sanitizer_started_kernel_names
        ),
        "sanitizer_started_kernel_bindings": sorted(
            sanitizer_started_kernel_bindings,
            key=lambda item: int(item["case_id"]),
        ),
        "sanitizer_started_kernel_binding_count": len(
            sanitizer_started_kernel_bindings
        ),
        "sanitizer_started_kernel_binding_sha256": _binding_digest(
            sanitizer_started_kernel_bindings
        ),
        "passed": True,
    }


def _read_summary(path: Path) -> dict:
    _reject_symlink_chain(path, "分片摘要")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise ValueError(f"分片摘要无法读取：{path}：{error}") from error
    if not isinstance(value, dict):
        raise ValueError(f"分片摘要必须是 JSON object：{path}")
    return value


def _verify_or_write_summary(args: argparse.Namespace, summary: dict) -> None:
    _reject_symlink_chain(args.summary, "分片摘要")
    if args.summary.exists():
        if _read_summary(args.summary) != summary:
            raise ValueError(f"分片摘要与当前原始证据不一致：{args.summary}")
        return
    args.summary.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def summarize_shard(args: argparse.Namespace) -> int:
    summary = _build_shard_summary(args)
    _verify_or_write_summary(args, summary)
    print(
        f"分片 {args.start}:{args.end} 校验通过："
        f"cases={summary['expected_cases']}, "
        f"keys={len(summary['expected_tiling_keys'])}"
    )
    return 0


def _matrix_shards(matrix_root: Path, case_count: int) -> list[tuple[Path, int, int]]:
    _reject_symlink_chain(matrix_root, "矩阵目录")
    root = matrix_root.resolve()
    if not root.is_dir():
        raise ValueError(f"矩阵目录不存在、不是目录或为符号链接：{matrix_root}")
    shards: list[tuple[Path, int, int]] = []
    for path in root.iterdir():
        if not path.name.startswith("shard_"):
            continue
        match = SHARD_DIR_RE.fullmatch(path.name)
        if match is None:
            raise ValueError(f"非法分片目录：{path.name}")
        _reject_symlink_chain(path, "分片目录", root)
        if not path.is_dir():
            raise ValueError(f"非法分片目录：{path.name}")
        start, end = (int(value) for value in match.groups())
        if not 0 <= start < end <= case_count:
            raise ValueError(f"分片目录范围非法：{path.name}")
        if end - start != 1:
            raise ValueError(f"正式矩阵分片必须为单 case：{path.name}")
        shards.append((path, start, end))
    if not shards:
        raise ValueError("没有找到可聚合的分片目录")

    shards.sort(key=lambda item: (item[1], item[2]))
    cursor = 0
    for path, start, end in shards:
        if start != cursor:
            raise ValueError(
                f"分片范围不连续或重叠：期望从 {cursor} 开始，实际为 {path.name}"
            )
        cursor = end
    if cursor != case_count:
        raise ValueError(
            f"分片范围未覆盖完整矩阵：覆盖到 {cursor}，应为 {case_count}"
        )
    return shards


def _build_matrix_summary(args: argparse.Namespace) -> dict:
    cases = _load_cases(args.case_file)
    _validate_scope_contract(args, len(cases))
    if args.scope == "accuracy":
        _validate_accuracy_seed_contract(cases)
    runtime_manifest = _verify_current_runtime(args)
    summaries = []
    for shard_root, start, end in _matrix_shards(args.matrix_root, len(cases)):
        summary_path = shard_root / "summary.json"
        _reject_symlink_chain(summary_path, "分片摘要", shard_root)
        if not summary_path.is_file():
            raise ValueError(f"完整分片缺少 summary.json：{shard_root}")
        shard_args = argparse.Namespace(
            case_file=args.case_file,
            scope=args.scope,
            tool=args.tool,
            runtime_manifest=args.runtime_manifest,
            timeout=args.timeout,
            loop_nums=args.loop_nums,
            gm_init_mode=args.gm_init_mode,
            single_process_mode=args.single_process_mode,
            require_tiling_log=args.require_tiling_log,
            shard_root=shard_root,
            console_log=shard_root / "console.log",
            summary=summary_path,
            start=start,
            end=end,
            sanitizer_log=(
                shard_root / f"{args.tool}.log"
                if args.scope == "mssanitizer"
                else None
            ),
        )
        rebuilt = _build_shard_summary(shard_args)
        if _read_summary(summary_path) != rebuilt:
            raise ValueError(
                f"分片摘要与当前原始证据不一致：{summary_path}"
            )
        summaries.append(rebuilt)

    case_hash = _case_hash(args.case_file)
    case_ids: list[int] = []
    expected_keys: set[int] = set()
    host_keys: set[int] = set()
    launch_keys: set[int] = set()
    sanitizer_start_count = 0
    sanitizer_started_keys: set[int] = set()
    sanitizer_started_kernel_names: set[str] = set()
    sanitizer_started_kernel_bindings: list[dict] = []
    sanitizer_outer_log_hashes: list[str] = []
    runtime_manifest_hash = _file_hash(args.runtime_manifest)
    manifest_kernel_map: dict[str, int] = {}
    if args.scope == "mssanitizer":
        manifest_kernel_map = _manifest_kernel_map(runtime_manifest)
    for summary in summaries:
        if summary.get("schema") != "kda-prepare-atk-shard/v2":
            raise ValueError("分片摘要 schema 不是 v2")
        if (
            summary.get("scope") != args.scope
            or summary.get("tool") != args.tool
            or summary.get("timeout_seconds") != args.timeout
            or summary.get("loop_nums") != args.loop_nums
            or summary.get("gm_init_mode") != args.gm_init_mode
            or summary.get("single_process_mode") != args.single_process_mode
            or summary.get("require_tiling_log") != args.require_tiling_log
        ):
            raise ValueError(
                "分片 scope/tool/timeout/loop_nums/gm/tiling-log 不一致"
            )
        if summary.get("runtime_manifest_sha256") != runtime_manifest_hash:
            raise ValueError("分片运行时 kernel 指纹不一致")
        if summary.get("case_json_sha256") != case_hash or not summary.get("passed"):
            raise ValueError("分片用例哈希不一致或未通过")
        case_ids.extend(int(value) for value in summary["case_ids"])
        expected_keys.update(int(value) for value in summary["expected_tiling_keys"])
        host_keys.update(int(value) for value in summary["host_tiling_keys"])
        launch_keys.update(int(value) for value in summary["launch_tiling_keys"])
        if args.scope == "mssanitizer":
            outer_hash = str(summary.get("sanitizer_outer_log_sha256", ""))
            if not re.fullmatch(r"[0-9a-f]{64}", outer_hash):
                raise ValueError("内存分片缺少外层 sanitizer 原始日志哈希")
            sanitizer_outer_log_hashes.append(outer_hash)
            shard_start_count = int(summary.get("sanitizer_start_count", 0))
            if shard_start_count <= 0:
                raise ValueError("内存分片缺少 sanitizer 启动证据")
            sanitizer_start_count += shard_start_count
            shard_started_keys = {
                int(value)
                for value in summary.get("sanitizer_started_tiling_keys", ())
            }
            if shard_started_keys != {
                int(value) for value in summary["expected_tiling_keys"]
            }:
                raise ValueError("内存分片的 sanitizer 启动 TilingKey 不闭合")
            shard_started_kernel_names = {
                str(value)
                for value in summary.get(
                    "sanitizer_started_kernel_names", ()
                )
            }
            unexpected_kernel_names = (
                shard_started_kernel_names - set(manifest_kernel_map)
            )
            if unexpected_kernel_names:
                raise ValueError(
                    "内存分片命中 runtime manifest 之外的 kernel："
                    f"{sorted(unexpected_kernel_names)}"
                )
            if {
                manifest_kernel_map[name]
                for name in shard_started_kernel_names
            } != shard_started_keys:
                raise ValueError(
                    "内存分片的完整 kernelName 与 TilingKey 不一致"
                )
            if len(shard_started_kernel_names) != len(shard_started_keys):
                raise ValueError(
                    "内存分片同一 TilingKey 命中了多个 kernel 身份"
                )
            shard_bindings = summary.get("sanitizer_started_kernel_bindings")
            if not isinstance(shard_bindings, list):
                raise ValueError("内存分片缺少 kernel 调度绑定")
            if summary.get("sanitizer_started_kernel_binding_count") != len(
                shard_bindings
            ):
                raise ValueError("内存分片的 kernel 调度绑定数量不一致")
            if summary.get(
                "sanitizer_started_kernel_binding_sha256"
            ) != _binding_digest(shard_bindings):
                raise ValueError("内存分片的 kernel 调度绑定摘要不一致")
            case_bindings = _case_kernel_bindings(
                runtime_manifest,
                [cases[int(value)] for value in summary["case_ids"]],
            )
            if shard_bindings != case_bindings:
                raise ValueError("内存分片命中的 kernel 与 case dtype 不一致")
            if {
                str(item["kernel_name"]) for item in shard_bindings
            } != shard_started_kernel_names:
                raise ValueError("内存分片的 kernel 名称与调度绑定不一致")
            sanitizer_started_keys.update(shard_started_keys)
            sanitizer_started_kernel_names.update(
                shard_started_kernel_names
            )
            sanitizer_started_kernel_bindings.extend(shard_bindings)

    expected_ids = list(range(len(cases)))
    duplicate_ids = sorted(
        case_id for case_id in set(case_ids) if case_ids.count(case_id) > 1
    )
    if sorted(case_ids) != expected_ids or duplicate_ids:
        raise ValueError(
            f"矩阵 case ID 不闭合：missing={sorted(set(expected_ids) - set(case_ids))}, "
            f"extra={sorted(set(case_ids) - set(expected_ids))}, "
            f"duplicates={duplicate_ids}"
        )

    _, full_pairs = _expected_slice(cases, 0, len(cases))
    full_expected_keys = {key for _, key in full_pairs}
    if args.require_tiling_log:
        if len(full_expected_keys) != len(cases):
            raise ValueError(
                "冻结矩阵不是一条 case 对应一个唯一 TilingKey："
                f"cases={len(cases)}, keys={len(full_expected_keys)}"
            )
        if host_keys != full_expected_keys or launch_keys != full_expected_keys:
            raise ValueError("完整矩阵的 expected/host/launch TilingKey 集合不一致")
        if args.scope == "mssanitizer" and sanitizer_started_keys != full_expected_keys:
            raise ValueError("完整矩阵的 sanitizer 启动 TilingKey 集合不一致")
        if args.scope == "mssanitizer" and len(
            sanitizer_started_kernel_names
        ) != len(full_expected_keys):
            raise ValueError("完整矩阵的 sanitizer kernel 身份数量不一致")
        if args.scope == "mssanitizer":
            expected_bindings = _case_kernel_bindings(runtime_manifest, cases)
            if sanitizer_started_kernel_bindings != expected_bindings:
                raise ValueError("完整矩阵的 sanitizer kernel 调度绑定不一致")

    return {
        "schema": "kda-prepare-atk-matrix/v2",
        "scope": args.scope,
        "tool": args.tool,
        "timeout_seconds": args.timeout,
        "loop_nums": args.loop_nums,
        "gm_init_mode": args.gm_init_mode,
        "single_process_mode": args.single_process_mode,
        "require_tiling_log": args.require_tiling_log,
        "runtime_manifest_sha256": runtime_manifest_hash,
        "expected_cases": len(cases),
        "observed_cases": len(case_ids),
        "shard_count": len(summaries),
        "case_json_sha256": case_hash,
        "expected_tiling_key_count": len(full_expected_keys),
        "host_tiling_key_count": len(host_keys),
        "launch_tiling_key_count": len(launch_keys),
        "sanitizer_start_count": sanitizer_start_count,
        "sanitizer_started_tiling_key_count": len(sanitizer_started_keys),
        "sanitizer_started_kernel_name_count": len(
            sanitizer_started_kernel_names
        ),
        "sanitizer_started_kernel_name_sha256": _string_digest(
            sanitizer_started_kernel_names
        ),
        "sanitizer_started_kernel_binding_count": len(
            sanitizer_started_kernel_bindings
        ),
        "sanitizer_started_kernel_binding_sha256": _binding_digest(
            sanitizer_started_kernel_bindings
        ),
        **(
            {
                "sanitizer_outer_log_count": len(sanitizer_outer_log_hashes),
                "sanitizer_outer_log_sha256": hashlib.sha256(
                    "".join(sanitizer_outer_log_hashes).encode("ascii")
                ).hexdigest(),
            }
            if args.scope == "mssanitizer"
            else {}
        ),
        "key_pair_sha256": _key_digest(full_pairs),
        "complete": True,
        "passed": True,
    }


def _verify_or_write_matrix_summary(path: Path, summary: dict) -> None:
    _reject_symlink_chain(path, "矩阵汇总")
    if path.exists():
        if _read_summary(path) != summary:
            raise ValueError(f"矩阵汇总与当前原始证据不一致：{path}")
        return
    path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def aggregate_matrix(args: argparse.Namespace) -> int:
    result = _build_matrix_summary(args)
    _verify_or_write_matrix_summary(args.output, result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def verify_sanitizer_suite(args: argparse.Namespace) -> int:
    cases = _load_cases(args.case_file)
    _, expected_pairs = _expected_slice(cases, 0, len(cases))
    expected_keys = {key for _, key in expected_pairs}
    expected_case_count = SCOPE_CONTRACTS["mssanitizer"]["case_count"]
    if (
        len(cases) != expected_case_count
        or len(expected_keys) != expected_case_count
    ):
        raise ValueError(
            "sanitizer suite 要求 "
            f"{expected_case_count} 条 case 对应同样数量的唯一 TilingKey"
        )
    expected_tools = {"memcheck", "racecheck", "initcheck", "synccheck"}
    if len(args.aggregate) != len(expected_tools):
        raise ValueError("必须提供四种 sanitizer 的 aggregate_summary.json")

    results: dict[str, dict] = {}
    runtime_hashes: set[str] = set()
    kernel_name_hashes: set[str] = set()
    kernel_binding_hashes: set[str] = set()
    for path in args.aggregate:
        # 必须在 resolve 前检查用户提供的路径，避免把符号链接目标当作可信汇总。
        _reject_symlink_chain(path, "sanitizer 汇总")
        resolved_path = path.resolve()
        if (
            not resolved_path.is_file()
            or resolved_path.name != "aggregate_summary.json"
        ):
            raise ValueError(f"sanitizer 汇总路径非法：{path}")
        matrix_root = resolved_path.parent
        runtime_manifest = matrix_root / "runtime_manifest.json"
        existing = _read_summary(resolved_path)
        tool = str(existing.get("tool", ""))
        if tool in results:
            raise ValueError(f"sanitizer 汇总重复：{tool}")
        matrix_args = argparse.Namespace(
            case_file=args.case_file,
            scope="mssanitizer",
            tool=tool,
            runtime_manifest=runtime_manifest,
            timeout=int(existing.get("timeout_seconds", 0)),
            loop_nums=1,
            gm_init_mode="not_applicable",
            single_process_mode="off",
            require_tiling_log=True,
            matrix_root=matrix_root,
            soc=args.soc,
            test_artifact=args.test_artifact,
        )
        summary = _build_matrix_summary(matrix_args)
        if existing != summary:
            raise ValueError(
                f"sanitizer 汇总与当前原始证据不一致：{resolved_path}"
            )
        if (
            summary.get("schema") != "kda-prepare-atk-matrix/v2"
            or summary.get("scope") != "mssanitizer"
            or not summary.get("complete")
            or not summary.get("passed")
        ):
            raise ValueError(f"sanitizer 汇总未通过或 schema 错误：{path}")
        if summary.get("case_json_sha256") != _case_hash(args.case_file):
            raise ValueError(f"sanitizer 汇总的 case 哈希不一致：{path}")
        if (
            summary.get("expected_cases") != len(cases)
            or summary.get("observed_cases") != len(cases)
            or summary.get("expected_tiling_key_count") != len(expected_keys)
            or summary.get("host_tiling_key_count") != len(expected_keys)
            or summary.get("launch_tiling_key_count") != len(expected_keys)
            or summary.get("sanitizer_started_tiling_key_count")
            != len(expected_keys)
            or summary.get("sanitizer_started_kernel_name_count")
            != len(expected_keys)
            or not summary.get("sanitizer_started_kernel_name_sha256")
            or summary.get("sanitizer_started_kernel_binding_count")
            != len(expected_keys)
            or not re.fullmatch(
                r"[0-9a-f]{64}",
                str(
                    summary.get(
                        "sanitizer_started_kernel_binding_sha256", ""
                    )
                ),
            )
            or summary.get("sanitizer_outer_log_count") != len(cases)
            or not re.fullmatch(
                r"[0-9a-f]{64}",
                str(summary.get("sanitizer_outer_log_sha256", "")),
            )
            or summary.get("key_pair_sha256") != _key_digest(expected_pairs)
        ):
            raise ValueError(f"sanitizer 汇总的 case/key 覆盖不闭合：{path}")
        timeout_limit = SCOPE_TIMEOUT_LIMITS["mssanitizer"]
        if not 1 <= int(summary.get("timeout_seconds", 0)) <= timeout_limit:
            raise ValueError(
                f"sanitizer 超时配置不在 1..{timeout_limit} 秒内：{path}"
            )
        if (
            summary.get("loop_nums") != 1
            or summary.get("gm_init_mode") != "not_applicable"
            or summary.get("single_process_mode") != "off"
            or summary.get("require_tiling_log") is not True
        ):
            raise ValueError(f"sanitizer 汇总不符合正式 scope 合同：{path}")
        runtime_hashes.add(str(summary.get("runtime_manifest_sha256", "")))
        kernel_name_hashes.add(
            str(summary.get("sanitizer_started_kernel_name_sha256", ""))
        )
        kernel_binding_hashes.add(
            str(summary.get("sanitizer_started_kernel_binding_sha256", ""))
        )
        results[tool] = summary

    if set(results) != expected_tools:
        raise ValueError(
            f"sanitizer 工具不闭合：missing={sorted(expected_tools - set(results))}, "
            f"extra={sorted(set(results) - expected_tools)}"
        )
    if len(runtime_hashes) != 1 or "" in runtime_hashes:
        raise ValueError("四种 sanitizer 未使用同一份 runtime kernel")
    if len(kernel_name_hashes) != 1 or "" in kernel_name_hashes:
        raise ValueError("四种 sanitizer 命中的完整 kernel 身份不一致")
    if len(kernel_binding_hashes) != 1 or "" in kernel_binding_hashes:
        raise ValueError("四种 sanitizer 命中的 case/kernel 调度绑定不一致")

    suite = {
        "schema": "kda-prepare-atk-sanitizer-suite/v1",
        "tools": sorted(results),
        "case_json_sha256": _case_hash(args.case_file),
        "runtime_manifest_sha256": runtime_hashes.pop(),
        "sanitizer_started_kernel_name_sha256": kernel_name_hashes.pop(),
        "sanitizer_started_kernel_binding_sha256": (
            kernel_binding_hashes.pop()
        ),
        "expected_cases_per_tool": len(cases),
        "observed_cases_total": len(cases) * len(expected_tools),
        "tiling_key_count_per_tool": len(expected_keys),
        "complete": True,
        "passed": True,
    }
    _verify_or_write_matrix_summary(args.output, suite)
    print(json.dumps(suite, ensure_ascii=False, indent=2))
    return 0


def verify_sanitizer_log(args: argparse.Namespace) -> int:
    result = _sanitizer_log_summary(
        tuple(args.log), args.tool, args.expected_finish_count
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def _common_parser(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--case-file", type=Path, required=True)
    parser.add_argument(
        "--scope", choices=("accuracy", "determinism", "mssanitizer"), required=True
    )
    parser.add_argument("--tool", default="")
    parser.add_argument("--runtime-manifest", type=Path, required=True)
    parser.add_argument("--timeout", type=int, required=True)
    parser.add_argument("--loop-nums", type=int, required=True)
    parser.add_argument("--gm-init-mode", required=True)
    parser.add_argument(
        "--single-process-mode", choices=("on", "off"), required=True
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    runtime = subparsers.add_parser("runtime", help="冻结实际加载的 kernel 指纹")
    runtime.add_argument("--case-file", type=Path, required=True)
    runtime.add_argument("--soc", default="auto")
    runtime.add_argument("--require-sanitizer", action="store_true")
    runtime.add_argument("--require-complete-key-set", action="store_true")
    runtime.add_argument(
        "--test-artifact", type=Path, action="append", required=True
    )
    runtime.add_argument("--output", type=Path, required=True)
    runtime.set_defaults(func=verify_runtime)

    shard = subparsers.add_parser("shard", help="校验一个已完成的 ATK 分片")
    _common_parser(shard)
    shard.add_argument("--shard-root", type=Path, required=True)
    shard.add_argument("--console-log", type=Path, required=True)
    shard.add_argument(
        "--sanitizer-log",
        type=Path,
        help="外层 mssanitizer 原始日志；默认与分片目录下的 <tool>.log 相同",
    )
    shard.add_argument("--summary", type=Path, required=True)
    shard.add_argument("--start", type=int, required=True)
    shard.add_argument("--end", type=int, required=True)
    shard.add_argument("--require-tiling-log", action="store_true")
    shard.set_defaults(func=summarize_shard)

    aggregate = subparsers.add_parser("aggregate", help="聚合并闭环校验完整矩阵")
    _common_parser(aggregate)
    aggregate.add_argument("--matrix-root", type=Path, required=True)
    aggregate.add_argument(
        "--soc",
        choices=("ascend910b", "ascend910_93", "ascend950"),
        required=True,
    )
    aggregate.add_argument(
        "--test-artifact", type=Path, action="append", required=True
    )
    aggregate.add_argument("--output", type=Path, required=True)
    aggregate.add_argument("--require-tiling-log", action="store_true")
    aggregate.set_defaults(func=aggregate_matrix)

    suite = subparsers.add_parser(
        "sanitizer-suite", help="校验四种 sanitizer 的全量矩阵汇总"
    )
    suite.add_argument("--case-file", type=Path, required=True)
    suite.add_argument(
        "--soc",
        choices=("ascend910b", "ascend910_93", "ascend950"),
        required=True,
    )
    suite.add_argument(
        "--test-artifact", type=Path, action="append", required=True
    )
    suite.add_argument("--aggregate", type=Path, action="append", required=True)
    suite.add_argument("--output", type=Path, required=True)
    suite.set_defaults(func=verify_sanitizer_suite)

    sanitizer_log = subparsers.add_parser(
        "sanitizer-log", help="校验统一 runner 的原始 sanitizer 日志"
    )
    sanitizer_log.add_argument("--tool", choices=sorted(SANITIZER_TOOLS), required=True)
    sanitizer_log.add_argument("--log", type=Path, action="append", required=True)
    sanitizer_log.add_argument("--expected-finish-count", type=int, required=True)
    sanitizer_log.set_defaults(func=verify_sanitizer_log)

    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
