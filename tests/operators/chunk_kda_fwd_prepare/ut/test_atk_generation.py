"""检查冻结 ATK 用例与统一 JSON 中的声明完全一致。"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from unittest import mock


ROOT = Path(__file__).resolve().parents[4]
MANIFEST_PATH = ROOT / "tests/op_cases/chunk_kda_fwd_prepare.json"
GENERATOR_PATH = (
    ROOT / "tests/atk/chunk_kda_fwd_prepare/gen_chunk_kda_fwd_prepare.py"
)
ATK_DIR = GENERATOR_PATH.parent
VERIFIER_PATH = ATK_DIR / "scripts/verify_matrix.py"
TILING_MATRIX_GENERATOR_PATH = (
    ATK_DIR / "scripts/generate_tiling_key_matrix.py"
)
BUILD_WHEEL_PATH = ROOT / "scripts/build_wheel.py"
BIN_PARAM_BUILDER_PATH = ROOT / "cmake/scripts/util/ascendc_bin_param_build.py"
PREPARE_KERNEL_DIR = ROOT / "fla/ops/ascendc/kda/chunk_kda_fwd_prepare"
TILING_HOST_PATH = PREPARE_KERNEL_DIR / "op_host/chunk_kda_fwd_prepare_tiling.cpp"
TILING_KEY_PATH = PREPARE_KERNEL_DIR / "op_kernel/chunk_kda_fwd_prepare_tiling_key.h"
POLICY_PATH = PREPARE_KERNEL_DIR / "op_kernel/chunk_kda_fwd_prepare_policy.h"


def _load_generator():
    module_name = "chunk_kda_fwd_prepare_atk_generator_test"
    spec = importlib.util.spec_from_file_location(module_name, GENERATOR_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {GENERATOR_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


GENERATOR = _load_generator()


def _load_verifier():
    module_name = "chunk_kda_fwd_prepare_atk_matrix_verifier_test"
    spec = importlib.util.spec_from_file_location(module_name, VERIFIER_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {VERIFIER_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


VERIFIER = _load_verifier()


def _load_build_wheel():
    module_name = "chunk_kda_fwd_prepare_build_wheel_test"
    spec = importlib.util.spec_from_file_location(module_name, BUILD_WHEEL_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {BUILD_WHEEL_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


BUILD_WHEEL = _load_build_wheel()


def _load_bin_param_builder():
    module_name = "chunk_kda_fwd_prepare_bin_param_builder_test"
    util_dir = str(BIN_PARAM_BUILDER_PATH.parent)
    sys.path.insert(0, util_dir)
    try:
        spec = importlib.util.spec_from_file_location(
            module_name, BIN_PARAM_BUILDER_PATH
        )
        if spec is None or spec.loader is None:
            raise RuntimeError(f"cannot load {BIN_PARAM_BUILDER_PATH}")
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(util_dir)


BIN_PARAM_BUILDER = _load_bin_param_builder()


def _support_info_fixture(gate_dtype: str, beta_dtype: str) -> dict:
    metadata_dtype = {"bf16": "bfloat16", "fp32": "float32"}
    dtype_token = {"bf16": "27", "fp32": "0"}
    tensor_tokens = [
        "27,2",
        "27,2",
        "27,2",
        f"{dtype_token[gate_dtype]},2",
        f"{dtype_token[beta_dtype]},2",
        "0,2",
        *("27,2" for _ in range(9)),
        *("0,2" for _ in range(3)),
    ]
    return {
        "inputs": [
            {
                "name": "q",
                "index": 0,
                "dtype": "bfloat16",
                "format": "ND",
                "paramType": "required",
            },
            {
                "name": "k",
                "index": 1,
                "dtype": "bfloat16",
                "format": "ND",
                "paramType": "required",
            },
            {
                "name": "v",
                "index": 2,
                "dtype": "bfloat16",
                "format": "ND",
                "paramType": "required",
            },
            {
                "name": "g",
                "index": 3,
                "dtype": metadata_dtype[gate_dtype],
                "format": "ND",
                "paramType": "required",
            },
            {
                "name": "beta",
                "index": 4,
                "dtype": metadata_dtype[beta_dtype],
                "format": "ND",
                "paramType": "required",
            },
        ],
        "simplifiedKeyMode": 0,
        "simplifiedKey": [
            "/".join(
                ["ChunkKdaFwdPrepare", f"d={d},p={p}", *tensor_tokens]
            )
            for d in (0, 1)
            for p in (0, 1)
        ],
    }


def _dispatch_contract_fixture(gate_dtype: str, beta_dtype: str) -> dict:
    return VERIFIER._metadata_dispatch_contract(
        {"supportInfo": _support_info_fixture(gate_dtype, beta_dtype)},
        "fixture.json",
    )


def _toolchain_fixture(sanitizer: bool = False) -> dict:
    digest = "a" * 64
    toolchain = {
        "atk": {
            "version": "26.8.8",
            "executable_sha256": digest,
        }
    }
    if sanitizer:
        toolchain["mssanitizer"] = {
            "revision": "8.3.RC1-a1b2c3d4",
            "msopscommon_revision": "e5f6a7b8",
            "executable_sha256": digest,
        }
    return toolchain


def _runtime_manifest_fixture(
    keys: set[int],
    *,
    sanitizer: bool = False,
    complete: bool = False,
    binary_count: int = 1,
) -> dict:
    signatures = [
        ("bf16", "bf16"),
        ("bf16", "fp32"),
        ("fp32", "bf16"),
        ("fp32", "fp32"),
    ]
    if not 1 <= binary_count <= len(signatures):
        raise ValueError("fixture binary_count 必须在 1..4")
    compile_hashes = ["171e9a19df4c0d44ca00a177d8fe300c"] + [
        f"{index:032x}" for index in range(2, binary_count + 1)
    ]
    binary_names = [f"ChunkKdaFwdPrepare_{value}" for value in compile_hashes]
    kernel_names = {
        f"{binary_name}_{key}" for binary_name in binary_names for key in keys
    }
    digest = "a" * 64
    return {
        "schema": "kda-prepare-runtime/v2",
        "platform": "ascend950",
        "runtime_sha256": digest,
        "op_api_sha256": digest,
        "test_artifact_sha256": {"runner.py": digest},
        "python_wrapper_sha256": {
            "__init__.py": digest,
            "ops/ascendc/__init__.py": digest,
            "ops/ascendc/_aclnn_ctypes.py": digest,
            "ops/ascendc/_kda_policy.py": digest,
            "ops/ascendc/_runtime.py": digest,
        },
        "toolchain": _toolchain_fixture(sanitizer),
        "object_count": binary_count,
        "metadata_count": binary_count,
        "compiled_tiling_key_count": len(keys),
        "compiled_tiling_key_sha256": VERIFIER._integer_digest(keys),
        "kernel_name_count": len(kernel_names),
        "kernel_name_sha256": VERIFIER._string_digest(kernel_names),
        "kernel_binaries": [
            {
                "metadata_file": f"{binary_name}.json",
                "metadata_sha256": digest,
                "metadata_kernel_name": binary_name,
                "metadata_bin_sha256": digest,
                "bin_file_name": binary_name,
                "bin_file_suffix": ".o",
                "bin_file_sha256": digest,
                "compile_hash": compile_hash,
                **_dispatch_contract_fixture(gate_dtype, beta_dtype),
                "kernels": [
                    {
                        "kernel_name": f"{binary_name}_{key}",
                        "tiling_key": key,
                    }
                    for key in sorted(keys)
                ],
            }
            for (compile_hash, binary_name), (gate_dtype, beta_dtype) in zip(
                zip(compile_hashes, binary_names), signatures
            )
        ],
        "complete_key_set_required": complete,
        "sanitizer_required": sanitizer,
        "sanitizer_object_count": binary_count if sanitizer else 0,
    }


class ChunkKdaFwdPrepareAtkGenerationTest(unittest.TestCase):
    def test_runtime_toolchain_parsers_require_unique_identity(self):
        self.assertEqual(
            VERIFIER._parse_atk_version("ATK\n26.8.8\nready\n"),
            "26.8.8",
        )
        with self.assertRaisesRegex(ValueError, "唯一解析版本"):
            VERIFIER._parse_atk_version("ATK\n")
        with self.assertRaisesRegex(ValueError, "唯一解析版本"):
            VERIFIER._parse_atk_version("26.8.8\n26.9.0\n")

        sanitizer_revision, common_revision = (
            VERIFIER._parse_mssanitizer_revisions(
                "revision:\n"
                "mssanitizer 8.3.RC1-a1b2c3d4\n"
                "msopscommon e5f6a7b8\n"
            )
        )
        self.assertEqual(sanitizer_revision, "8.3.RC1-a1b2c3d4")
        self.assertEqual(common_revision, "e5f6a7b8")
        with self.assertRaisesRegex(ValueError, "唯一解析"):
            VERIFIER._parse_mssanitizer_revisions(
                "mssanitizer 8.3.RC1-a1b2c3d4\n"
            )

        normal = _runtime_manifest_fixture({2570})
        VERIFIER._validate_runtime_toolchain(normal, False)
        sanitizer = _runtime_manifest_fixture({2570}, sanitizer=True)
        VERIFIER._validate_runtime_toolchain(sanitizer, True)
        del sanitizer["toolchain"]["mssanitizer"]
        with self.assertRaisesRegex(ValueError, "工具集合不闭合"):
            VERIFIER._validate_runtime_toolchain(sanitizer, True)

    def test_sanitizer_build_keeps_required_debug_artifacts(self):
        args = SimpleNamespace(
            build_args=[], debug=False, sanitizer=True, oom=False
        )
        self.assertEqual(
            BUILD_WHEEL._assemble_build_args(args),
            "--bisheng_flags=ccec_g,sanitizer,dump_cce",
        )
        generator = (
            ROOT / "cmake/scripts/util/ascendc_bin_param_build.py"
        ).read_text(encoding="utf-8")
        self.assertIn('build_cmd_var += " --op_debug_level=1"', generator)
        cmake_helpers = (ROOT / "cmake/func.cmake").read_text(encoding="utf-8")
        self.assertIn('STREQUAL "ccec_g"', cmake_helpers)
        self.assertIn('list(APPEND _OPC_CONFIG "-g")', cmake_helpers)
        self.assertIn('STREQUAL "sanitizer"', cmake_helpers)
        self.assertIn('list(APPEND _OPC_CONFIG "-sanitizer")', cmake_helpers)
        self.assertIn('STREQUAL "ascend950"', cmake_helpers)
        self.assertIn("COMPUTE_UNIT ascend950pr_9599", cmake_helpers)
        self.assertIn("OPTIONS --cce-enable-sanitizer", cmake_helpers)
        custom_build = (ROOT / "cmake/custom_build.cmake").read_text(
            encoding="utf-8"
        )
        self.assertIn(
            "CONFIG ${OP_DEBUG_CONFIG} ${BISHENG_FLAGS}", custom_build
        )
        root_cmake = (ROOT / "CMakeLists.txt").read_text(encoding="utf-8")
        self.assertIn(
            "CONFIG ${OP_DEBUG_CONFIG} ${BISHENG_FLAGS}", root_cmake
        )

    def test_sanitizer_build_preserves_and_deduplicates_explicit_configs(self):
        args = SimpleNamespace(
            build_args=[
                "-O3 --bisheng_flags=ccec_g,oom "
                "--op_debug_config dump_bin,sanitizer"
            ],
            debug=False,
            sanitizer=True,
            oom=False,
        )
        values = "ccec_g,sanitizer,dump_cce,oom,dump_bin"
        self.assertEqual(
            BUILD_WHEEL._assemble_build_args(args),
            f"-O3 --bisheng_flags={values}",
        )

    def test_final_asc_opc_command_merges_debug_configs_once(self):
        builder = object.__new__(BIN_PARAM_BUILDER.BinParamBuilder)
        builder.soc = "ascend950"
        builder.op_type = "ChunkKdaFwdPrepare"
        builder.op_file = "chunk_kda_fwd_prepare"
        builder.op_intf = "chunk_kda_fwd_prepare"
        builder.tiling_keys = set()
        builder.op_debug_config = {"dump_bin", "dump_cce"}
        builder.op_super_config = []

        with tempfile.TemporaryDirectory() as tmpdir:
            builder.out_path = tmpdir
            with mock.patch.dict(os.environ, {"CI_MODE": "TRUE"}, clear=False):
                with mock.patch.object(
                    BIN_PARAM_BUILDER.BinParamBuilder,
                    "_generate_check_result",
                    return_value="",
                ):
                    builder._write_build_cmd(
                        "params.json",
                        "kernel_bin",
                        0,
                        tmpdir,
                        "ccec_g,sanitizer,dump_cce,oom",
                    )
            command = (
                Path(tmpdir) / "ChunkKdaFwdPrepare-chunk_kda_fwd_prepare-0.sh"
            ).read_text(encoding="utf-8")

        self.assertEqual(command.count("--op_debug_config="), 1)
        self.assertIn(
            "--op_debug_config=ccec_g,sanitizer,dump_cce,oom,dump_bin",
            command,
        )
        self.assertIn("--op_debug_level=1", command)

    def test_generation_config_comes_from_canonical_manifest(self):
        manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        self.assertEqual(GENERATOR.GENERATION, manifest["atk_generation"])
        source = GENERATOR_PATH.read_text(encoding="utf-8")
        self.assertNotIn("FUNCTIONAL_SPECS", source)
        self.assertNotIn("PERF_SPECS", source)

    def test_frozen_json_matches_generator(self):
        suites = (
            (
                ATK_DIR / "atk_chunk_kda_fwd_prepare.json",
                GENERATOR.build_accuracy_specs(),
                200,
            ),
            (
                ATK_DIR / "atk_chunk_kda_fwd_prepare_perf.json",
                GENERATOR.build_perf_specs(),
                10,
            ),
            (
                ATK_DIR / "atk_chunk_kda_fwd_prepare_mss.json",
                GENERATOR.build_mss_specs(),
                432,
            ),
        )
        for path, specs, expected_count in suites:
            with self.subTest(path=path.name):
                frozen = json.loads(path.read_text(encoding="utf-8"))
                self.assertEqual(len(frozen), expected_count)
                self.assertEqual(frozen, GENERATOR._payloads(specs))

    def test_output_modes_cover_accuracy_and_inference_performance(self):
        accuracy_modes = {
            spec["backward_mode"]
            for spec in GENERATOR.build_accuracy_specs()
        }
        self.assertEqual(accuracy_modes, {"none", "recompute", "save"})
        self.assertEqual(
            {
                spec["backward_mode"]
                for spec in GENERATOR.build_perf_specs()
            },
            {"none"},
        )
        self.assertEqual(
            {
                spec["backward_mode"]
                for spec in GENERATOR.build_mss_specs()
            },
            {"none", "recompute", "save"},
        )

    def test_tiling_key_matrix_is_generated_from_frozen_json(self):
        result = subprocess.run(
            [sys.executable, str(TILING_MATRIX_GENERATOR_PATH), "--check"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr or result.stdout)

    def test_performance_cases_preserve_user_model_shapes(self):
        specs = GENERATOR.build_perf_specs()
        self.assertEqual(
            [
                (spec["B"], spec["HK"], spec["HV"], spec["T"])
                for spec in specs
            ],
            [
                (2, 16, 32, 11264),
                (1, 16, 32, 11264),
                (1, 32, 32, 65536),
                (4, 96, 96, 128),
                (1, 32, 32, 160),
                (6, 6, 6, 1084),
                (1, 12, 12, 1084),
                (1, 96, 96, 8192),
                (1, 96, 96, 16384),
                (1, 8, 24, 32768),
            ],
        )
        for spec in specs:
            self.assertEqual(spec["layout"], "BNSD")
            self.assertEqual(spec["backward_mode"], "none")
            self.assertTrue(spec["use_qk_l2norm_in_kernel"])
            self.assertTrue(spec["use_gate_in_kernel"])
            self.assertTrue(spec["safe_gate"])
            self.assertTrue(spec["use_beta_sigmoid_in_kernel"])
            self.assertTrue(spec["allow_neg_eigval"])
            self.assertTrue(spec["use_exp2"])
        self.assertEqual(len(specs[2]["cu_seqlens"].split(",")), 65)
        self.assertTrue(specs[2]["explicit_chunk_indices"])
        self.assertEqual(len(specs[8]["cu_seqlens"].split(",")), 65)
        self.assertFalse(specs[8]["explicit_chunk_indices"])

    def test_readme_maps_all_frozen_case_families(self):
        source = (ATK_DIR / "README.md").read_text(encoding="utf-8")
        for declaration in GENERATOR.GENERATION["functional_cases"]:
            self.assertIn(f"`{declaration['case_key']}`", source)
        for declaration in GENERATOR.GENERATION["performance_cases"]:
            self.assertIn(f"`{declaration['case_key']}`", source)
        self.assertIn("TilingKey <-> `_mss.json` case ID", source)
        self.assertIn("17/432 个不同 TilingKey", source)
        self.assertIn(
            "f28a869f07d8e3f65e8d0c85768759bde635759a09d047cd78b4a9b7136d3b19",
            source,
        )
        self.assertIn(
            "ATK_SINGLE_PROCESS=off PERFORMANCE_TIMEOUT=60", source
        )

    def test_accuracy_has_200_fixed_seed_cases(self):
        specs = GENERATOR.build_accuracy_specs()
        logical_counts = Counter(spec["logical_case_key"] for spec in specs)
        self.assertEqual(len(specs), 200)
        self.assertEqual(len(logical_counts), 56)
        self.assertGreaterEqual(min(logical_counts.values()), 3)
        self.assertEqual(len({spec["seed"] for spec in specs}), 200)

        cases = GENERATOR._payloads(specs)
        VERIFIER._validate_accuracy_seed_contract(cases)

        duplicate_seed = json.loads(json.dumps(cases))
        first_spec = VERIFIER._case_spec(duplicate_seed[0])
        second_spec = VERIFIER._case_spec(duplicate_seed[1])
        second_spec["seed"] = first_spec["seed"]
        duplicate_seed[1]["default_seed"] = first_spec["seed"]
        for item in duplicate_seed[1]["inputs"]:
            if item.get("name") == "case_spec":
                item["range_values"] = json.dumps(second_spec)
        with self.assertRaisesRegex(ValueError, "重复 seed"):
            VERIFIER._validate_accuracy_seed_contract(duplicate_seed)

        broken_index = json.loads(json.dumps(cases))
        broken_spec = VERIFIER._case_spec(broken_index[1])
        broken_spec["seed_index"] = 0
        for item in broken_index[1]["inputs"]:
            if item.get("name") == "case_spec":
                item["range_values"] = json.dumps(broken_spec)
        with self.assertRaisesRegex(ValueError, "seed_index"):
            VERIFIER._validate_accuracy_seed_contract(broken_index)

    def test_mss_covers_every_reachable_tiling_key_and_pipeline_profile(self):
        specs = GENERATOR.build_mss_specs()
        keys = {spec["expected_tiling_key"] for spec in specs}
        self.assertEqual(len(specs), 432)
        self.assertEqual(len(keys), 432)
        self.assertEqual(min(keys), 2570)
        self.assertEqual(max(keys), 24452638)
        self.assertEqual(Counter(spec["backward_mode"] for spec in specs), {
            "none": 144,
            "recompute": 144,
            "save": 144,
        })
        self.assertEqual({spec["layout"] for spec in specs}, {
            "BNSD", "BSND", "NTD", "TND",
        })
        self.assertEqual(
            set(Counter(spec["mss_profile"] for spec in specs).values()),
            {36},
        )
        for profile in {spec["mss_profile"] for spec in specs}:
            profile_specs = [
                spec for spec in specs if spec["mss_profile"] == profile
            ]
            self.assertEqual({spec["gate_dtype"] for spec in profile_specs}, {
                "bf16", "fp32",
            })
            self.assertEqual({spec["beta_dtype"] for spec in profile_specs}, {
                "bf16", "fp32",
            })
            self.assertEqual({
                spec["use_qk_l2norm_in_kernel"] for spec in profile_specs
            }, {False, True})
            self.assertEqual({spec["use_exp2"] for spec in profile_specs}, {
                False, True,
            })
            self.assertEqual({
                (
                    spec["use_exp2"],
                    spec["dt_bias"],
                )
                for spec in profile_specs
                if spec["use_gate_in_kernel"]
            }, {
                (False, False), (False, True),
                (True, False), (True, True),
            })
        digest = hashlib.sha256(
            "".join(
                f"{spec['case_id']},{spec['expected_tiling_key']}\n"
                for spec in specs
            ).encode("ascii")
        ).hexdigest()
        self.assertEqual(
            digest,
            "f28a869f07d8e3f65e8d0c85768759bde635759a09d047cd78b4a9b7136d3b19",
        )
        tags = {tag for spec in specs for tag in spec["tags"].split(",")}
        for required in (
            "dense", "varlen", "auto_indices", "explicit_indices",
            "grid_stride", "slot_reuse", "head_split", "gva",
            "active1", "active2", "active3", "active4",
        ):
            self.assertIn(required, tags)

    def test_tiling_key_generator_matches_host_template_contract(self):
        policy = POLICY_PATH.read_text(encoding="utf-8")
        constants = {
            name: int(value)
            for name, value in re.findall(
                r"^#define\s+(CHUNK_KDA_FWD_PREPARE_[A-Z0-9_]+)\s+(\d+)\s*$",
                policy,
                re.MULTILINE,
            )
        }
        self.assertEqual(
            GENERATOR.DTYPE_TOKENS,
            {
                "bf16": constants["CHUNK_KDA_FWD_PREPARE_TPL_BF16"],
                "fp32": constants["CHUNK_KDA_FWD_PREPARE_TPL_FP32"],
            },
        )
        self.assertEqual(
            GENERATOR.BETA_MODE_TOKENS,
            {
                "raw": constants["CHUNK_KDA_FWD_PREPARE_BETA_RAW"],
                "sigmoid": constants["CHUNK_KDA_FWD_PREPARE_BETA_SIGMOID"],
                "two_sigmoid": constants[
                    "CHUNK_KDA_FWD_PREPARE_BETA_TWO_SIGMOID"
                ],
            },
        )
        self.assertEqual(
            GENERATOR.GATE_MODE_TOKENS,
            {
                "precomputed": constants[
                    "CHUNK_KDA_FWD_PREPARE_GATE_PRECOMPUTED_STEP"
                ],
                "softplus": constants["CHUNK_KDA_FWD_PREPARE_GATE_SOFTPLUS"],
                "safe": constants["CHUNK_KDA_FWD_PREPARE_GATE_SAFE_SIGMOID"],
            },
        )
        self.assertEqual(
            GENERATOR.OUTPUT_MODE_TOKENS,
            {
                "none": constants["CHUNK_KDA_FWD_PREPARE_OUTPUT_NONE"],
                "recompute": constants[
                    "CHUNK_KDA_FWD_PREPARE_OUTPUT_RECOMPUTE"
                ],
                "save": constants["CHUNK_KDA_FWD_PREPARE_OUTPUT_SAVE"],
            },
        )

        declaration = TILING_KEY_PATH.read_text(encoding="utf-8").split(
            "ASCENDC_TPL_ARGS_DECL(", 1
        )[1].split(");", 1)[0]
        axes = (
            "D_T_GATE",
            "D_T_BETA",
            "NORM_MODE",
            "BETA_MODE",
            "GATE_MODE",
            "USE_EXP2",
            "SAFE_GATE",
            "OUTPUT_MODE",
        )
        positions = [declaration.index(axis) for axis in axes]
        self.assertEqual(positions, sorted(positions))

        host = " ".join(TILING_HOST_PATH.read_text(encoding="utf-8").split())
        self.assertIn(
            "GET_TPL_TILING_KEY( gateDtypeToken, betaDtypeToken, normMode, "
            "betaMode, gateMode, static_cast<uint64_t>(*useExp2Ptr), "
            "static_cast<uint64_t>(*safeGatePtr), outputMode)",
            host,
        )

    def test_mss_profiles_cover_real_schedule_branches(self):
        profiles = {
            spec["mss_profile"]: spec
            for spec in GENERATOR.build_mss_specs()
        }

        def schedule(spec, core_count=128):
            if spec["cu_seqlens"]:
                cu = [int(value) for value in spec["cu_seqlens"].split(",")]
                chunks = sum(
                    (end - begin + 63) // 64
                    for begin, end in zip(cu, cu[1:])
                )
            else:
                chunks = int(spec["B"]) * ((int(spec["T"]) + 63) // 64)
            head_split = chunks < core_count and int(spec["HK"]) > 1
            heads_per_partition = (
                int(spec["HV"]) // int(spec["HK"])
                if head_split else int(spec["HV"])
            )
            total_work = chunks * (int(spec["HK"]) if head_split else 1)
            return chunks, heads_per_partition, total_work

        for name in (
            "varlen_ntd_grid_stride_auto",
            "varlen_tnd_grid_stride_explicit",
        ):
            chunks, _, total_work = schedule(profiles[name])
            self.assertEqual(chunks, 129)
            self.assertEqual(total_work, 129)

        chunks, heads_per_partition, total_work = schedule(
            profiles["dense_bnsd_grid_stride_gva_chunk_only"]
        )
        self.assertEqual(chunks, 130)
        self.assertEqual(heads_per_partition, 4)
        self.assertEqual(total_work, 130)

        chunks, heads_per_partition, total_work = schedule(
            profiles["dense_bsnd_gva"]
        )
        self.assertEqual(chunks, 2)
        self.assertEqual(heads_per_partition, 2)
        self.assertEqual(total_work, 4)

    def test_executor_checks_public_output_contract_without_extra_npu_kernels(self):
        source = (
            ATK_DIR / "executor_chunk_kda_fwd_prepare.py"
        ).read_text(encoding="utf-8")
        self.assertGreaterEqual(source.count("return result.contiguous()"), 2)
        self.assertIn("_validate_output_contract(", source)
        self.assertIn("torch.npu.synchronize()", source)
        self.assertIn("if not with_output:", source)
        self.assertIn("output.detach().cpu()", source)

    def test_timed_atk_scopes_have_configurable_process_scheduling(self):
        source = (ROOT / "tests/atk/run_test_cpu.sh").read_text(
            encoding="utf-8"
        )
        contract = source.split(
            "validate_chunk_kda_fwd_prepare_contract()", 1
        )[1].split("validate_case_json()", 1)[0]
        for timeout in (
            "ATK_TIMEOUT",
            "PERFORMANCE_TIMEOUT",
            "DC_TIMEOUT",
            "MSS_TIMEOUT",
        ):
            self.assertIn(
                f'validate_bounded_timeout {timeout} "${timeout}"',
                contract,
            )
        self.assertIn(
            "正式测试要求 ATK_SINGLE_PROCESS=off", contract
        )
        self.assertIn(
            "正式精度测试要求 ATK_GM_INIT_MODE=off", contract
        )
        self.assertIn(
            "正式确定性测试要求 DC_LOOP_NUMS=50", contract
        )
        self.assertIn("all 不能闭合四种 sanitizer", contract)
        accuracy = source.rsplit("if should_run accuracy; then", 1)[1].split(
            "if should_run performance; then", 1
        )[0]
        determinism = source.rsplit("if should_run determinism; then", 1)[1].split(
            "if should_run mssanitizer; then", 1
        )[0]
        self.assertIn('"${SINGLE_PROCESS_ARGS[@]}"', accuracy)
        self.assertIn('"${SINGLE_PROCESS_ARGS[@]}"', determinism)
        performance = source.rsplit("if should_run performance; then", 1)[1].split(
            "if should_run determinism; then", 1
        )[0]
        self.assertIn("record_ran_type performance", performance)

    def test_prepare_performance_enables_fluctuation_check(self):
        source = (ROOT / "tests/atk/run_test_cpu.sh").read_text(encoding="utf-8")
        performance = source.rsplit("if should_run performance; then", 1)[1].split(
            "if should_run determinism; then", 1
        )[0]
        self.assertIn("PERFORMANCE_FLUCTUATION_ARGS=(--fluctuation_check)", source)
        self.assertIn('"${PERFORMANCE_FLUCTUATION_ARGS[@]}"', performance)
        self.assertIn("--require-performance-stability", source)
        self.assertIn("PERFORMANCE_RUN_START_NS=", source)
        self.assertIn("--report-min-mtime-ns", source)

    def test_public_result_checker_keeps_legacy_all_scope(self):
        module_spec = importlib.util.spec_from_file_location(
            "chunk_kda_fwd_prepare_all_scope_checker_test",
            ROOT / "tests/atk/common/check_atk_result.py",
        )
        self.assertIsNotNone(module_spec)
        self.assertIsNotNone(module_spec.loader)
        result_checker = importlib.util.module_from_spec(module_spec)
        module_spec.loader.exec_module(result_checker)

        self.assertEqual(
            result_checker.DEFAULT_ALL_TYPES,
            ("accuracy", "determinism", "mssanitizer"),
        )
        self.assertIn("performance", result_checker.CHECKERS)

    def test_prepare_contract_is_checked_after_environment_files(self):
        source = (ROOT / "tests/atk/run_test_cpu.sh").read_text(
            encoding="utf-8"
        )
        env_end = source.index(
            'source_env_file "fla_npu_transformer环境"',
            source.index("if [[ -n \"${FLA_NPU_ENV:-"),
        )
        contract_call = source.index(
            "validate_chunk_kda_fwd_prepare_contract\n\nATK_BIN=",
            env_end,
        )
        self.assertGreater(contract_call, env_end)

    def test_prepare_accuracy_summary_uses_strict_pass_status(self):
        source = (ROOT / "tests/atk/run_test_cpu.sh").read_text(
            encoding="utf-8"
        )
        self.assertIn(
            'checker_args+=(--require-accuracy-pass)',
            source,
        )
        self.assertIn(
            'die "Prepare 必须执行结果检查，但找不到 ${RESULT_CHECK_PY}"',
            source,
        )
        checker = (ROOT / "tests/atk/common/check_atk_result.py").read_text(
            encoding="utf-8"
        )
        self.assertIn("require_accuracy_pass=False", checker)
        self.assertIn('val != "Pass"', checker)
        module_spec = importlib.util.spec_from_file_location(
            "chunk_kda_fwd_prepare_result_checker_test",
            ROOT / "tests/atk/common/check_atk_result.py",
        )
        self.assertIsNotNone(module_spec)
        self.assertIsNotNone(module_spec.loader)
        result_checker = importlib.util.module_from_spec(module_spec)
        module_spec.loader.exec_module(result_checker)
        header = [
            "名称",
            "总用例数",
            "执行失败用例个数",
            "通过用例个数",
            "精度是否达标",
        ]
        rows = [["npu_dut", 1, 0, 1, "-"]]
        self.assertTrue(result_checker._extract_summary_row(header, rows)["all_pass"])
        self.assertFalse(
            result_checker._extract_summary_row(
                header, rows, require_accuracy_pass=True
            )["all_pass"]
        )
        strict_header = [
            "名称",
            "总用例数",
            "执行成功用例个数",
            "执行失败用例个数",
            "通过用例个数",
            "精度是否达标",
        ]
        strict_rows = [["npu_dut", 2, 2, 0, 2, "Pass"]]
        self.assertTrue(
            result_checker._strict_accuracy_summary_valid(
                strict_header, strict_rows
            )
        )
        for bad_row in (
            ["npu_dut", 2, 1, 1, 2, "Pass"],
            ["npu_dut", 2, 2, 0, 2, "-"],
            ["npu_dut", 2, 2, 0, 1, "Pass"],
        ):
            with self.subTest(bad_row=bad_row):
                self.assertFalse(
                    result_checker._strict_accuracy_summary_valid(
                        strict_header, [bad_row]
                    )
                )
        with mock.patch.object(
            result_checker, "_find_accuracy_reports", return_value=[Path("report.xlsx")]
        ), mock.patch.object(
            result_checker,
            "_parse_summary",
            return_value=(strict_header, strict_rows),
        ):
            strict_result = result_checker.check_accuracy(
                "unused", "chunk_kda_fwd_prepare", require_accuracy_pass=True
            )
        self.assertTrue(strict_result["all_pass"])
        no_status_header = header[:-1]
        with mock.patch.object(
            result_checker, "_find_accuracy_reports", return_value=[Path("report.xlsx")]
        ), mock.patch.object(
            result_checker,
            "_parse_summary",
            return_value=(no_status_header, [rows[0][:-1]]),
        ):
            strict_result = result_checker.check_accuracy(
                "unused", "chunk_kda_fwd_prepare", require_accuracy_pass=True
            )
        self.assertFalse(strict_result["all_pass"])

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"), "需要 openpyxl 生成测试报告"
    )
    def test_performance_checker_requires_successful_statistic_rows(self):
        import openpyxl

        module_spec = importlib.util.spec_from_file_location(
            "chunk_kda_fwd_prepare_performance_checker_test",
            ROOT / "tests/atk/common/check_atk_result.py",
        )
        self.assertIsNotNone(module_spec)
        self.assertIsNotNone(module_spec.loader)
        result_checker = importlib.util.module_from_spec(module_spec)
        module_spec.loader.exec_module(result_checker)
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            expected_cases = root / "performance_cases.json"
            expected_cases.write_text(
                json.dumps([{"id": 0}, {"id": 1}]), encoding="utf-8"
            )
            report = (
                root
                / "perf/atk_output/atk_chunk_kda_fwd_prepare_test"
                / "report/result.xlsx"
            )
            report.parent.mkdir(parents=True)
            workbook = openpyxl.Workbook()
            statistic = workbook.active
            statistic.title = "statistic"
            statistic.append(
                [
                    "编号",
                    "npu_dut_Device性能（us）",
                    "npu_dut_性能波动校验结果",
                    "运行结果",
                ]
            )
            statistic.append([0, 12.5, "Pass", "SUCCESS"])
            statistic.append([1, 13.5, "Pass", "SUCCESS"])
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                require_performance_stability=True,
            )
            self.assertTrue(result["all_pass"])
            self.assertEqual(result["total"], 2)

            # 公共 checker 默认兼容 ATK 未开启波动校验时的 None 列。
            workbook = openpyxl.load_workbook(report)
            workbook["statistic"]["C2"] = None
            workbook["statistic"]["C3"] = None
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
            )
            self.assertTrue(result["all_pass"])
            workbook = openpyxl.load_workbook(report)
            workbook["statistic"]["C2"] = "Pass"
            workbook["statistic"]["C3"] = "Pass"
            workbook.save(report)
            workbook.close()

            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_start=0,
                expected_case_end=1,
            )
            self.assertFalse(result["all_pass"])

            self.assertIn("--expected-case-file", result["detail"])
            workbook = openpyxl.load_workbook(report)
            workbook["statistic"]["D2"] = "FAILED"
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
            )
            self.assertFalse(result["all_pass"])

            # 分段 smoke 只应校验请求的半开区间，而不是强制等待整套性能矩阵。
            workbook = openpyxl.load_workbook(report)
            statistic = workbook["statistic"]
            statistic.delete_rows(2, statistic.max_row)
            statistic.append([0, 12.5, "Pass", "SUCCESS"])
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                expected_case_start=0,
                expected_case_end=1,
                require_performance_stability=True,
            )
            self.assertTrue(result["all_pass"])
            self.assertEqual(result["total"], 1)

            # 稳定性列存在时，空值和 UNKNOWN 都不能被当作通过。
            workbook = openpyxl.load_workbook(report)
            statistic = workbook["statistic"]
            statistic["C2"] = "UNKNOWN"
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                expected_case_start=0,
                expected_case_end=1,
                require_performance_stability=True,
            )
            self.assertFalse(result["all_pass"])

            workbook = openpyxl.load_workbook(report)
            workbook["statistic"]["C2"] = None
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                expected_case_start=0,
                expected_case_end=1,
                require_performance_stability=True,
            )
            self.assertFalse(result["all_pass"])
            self.assertEqual(result["fail"], 1)

            workbook = openpyxl.load_workbook(report)
            statistic = workbook["statistic"]
            statistic["D2"] = "SUCCESS"
            statistic.delete_rows(3)
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                require_performance_stability=True,
            )
            self.assertFalse(result["all_pass"])
            self.assertIn("缺少=[1]", result["detail"])

            workbook = openpyxl.load_workbook(report)
            statistic = workbook["statistic"]
            statistic.append([None, 14.5, "Pass", "SUCCESS"])
            workbook.save(report)
            workbook.close()
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
            )
            self.assertFalse(result["all_pass"])

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"), "需要 openpyxl 生成测试报告"
    )
    def test_performance_checker_enforces_targets_and_report_freshness(self):
        """声明目标的 case 必须有新报告中的设备耗时且不超目标。"""
        import openpyxl

        module_spec = importlib.util.spec_from_file_location(
            "chunk_kda_fwd_prepare_performance_target_checker_test",
            ROOT / "tests/atk/common/check_atk_result.py",
        )
        self.assertIsNotNone(module_spec)
        self.assertIsNotNone(module_spec.loader)
        result_checker = importlib.util.module_from_spec(module_spec)
        module_spec.loader.exec_module(result_checker)

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            expected_cases = root / "performance_cases.json"
            expected_cases.write_text(
                json.dumps(
                    [
                        {
                            "id": 0,
                            "inputs": [
                                {
                                    "name": "case_spec",
                                    "range_values": json.dumps(
                                        {"performance_target_us": 10}
                                    ),
                                }
                            ],
                        },
                        {"id": 1},
                    ]
                ),
                encoding="utf-8",
            )
            report = (
                root
                / "perf/atk_output/atk_chunk_kda_fwd_prepare_test"
                / "report/result.xlsx"
            )
            report.parent.mkdir(parents=True)

            def write_report(device_header, value0):
                workbook = openpyxl.Workbook()
                statistic = workbook.active
                statistic.title = "statistic"
                statistic.append(["编号", device_header, "运行结果"])
                statistic.append([0, value0, "SUCCESS"])
                statistic.append([1, 999, "SUCCESS"])
                workbook.save(report)
                workbook.close()

            write_report("npu_dut_Device性能（us）", 9)
            fresh_ns = report.stat().st_mtime_ns
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                report_min_mtime_ns=fresh_ns,
            )
            self.assertTrue(result["all_pass"])
            self.assertEqual(result["target_case_count"], 1)
            self.assertEqual(result["target_failed_ids"], [])

            write_report("npu_dut_Device性能（us）", 11)
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
            )
            self.assertFalse(result["all_pass"])
            self.assertEqual(result["target_failed_ids"], [0])
            self.assertIn("目标未达标=[0]", result["detail"])

            write_report("npu_dut_运行结果", 9)
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
            )
            self.assertFalse(result["all_pass"])
            self.assertEqual(result["target_failed_ids"], [0])

            # 早于本次运行标记的历史报告不能被选中。
            write_report("npu_dut_Device性能（us）", 9)
            old_mtime_ns = report.stat().st_mtime_ns
            result = result_checker.check_performance(
                str(root),
                "chunk_kda_fwd_prepare",
                expected_case_file=str(expected_cases),
                report_min_mtime_ns=old_mtime_ns + 1,
            )
            self.assertFalse(result["found"])
            self.assertIn("本次运行后", result["detail"])

            # 没有目标的公共算子不要求设备耗时列。
            generic_cases = root / "generic_cases.json"
            generic_cases.write_text(json.dumps([{"id": 0}]), encoding="utf-8")
            generic_report = (
                root
                / "perf/atk_output/atk_generic_op_test"
                / "report/result.xlsx"
            )
            generic_report.parent.mkdir(parents=True)
            workbook = openpyxl.Workbook()
            statistic = workbook.active
            statistic.title = "statistic"
            statistic.append(["编号", "运行结果"])
            statistic.append([0, "SUCCESS"])
            workbook.save(generic_report)
            workbook.close()
            result = result_checker.check_performance(
                str(root), "generic_op", expected_case_file=str(generic_cases)
            )
            self.assertTrue(result["all_pass"])
            self.assertEqual(
                result_checker._load_expected_performance_targets(
                    str(generic_cases)
                )[1],
                {},
            )

    def test_mssanitizer_uses_one_log_for_outer_and_atk_for_prepare(self):
        source = (ROOT / "tests/atk/run_test_cpu.sh").read_text(
            encoding="utf-8"
        )
        mssanitizer = source.split("if should_run mssanitizer; then", 1)[1]
        self.assertIn(
            'mssanitizer --tool="$MSS_TOOL" --log-file "$MSS_LOG_PATH" --',
            mssanitizer,
        )
        self.assertIn('-msl "$MSS_LOG_PATH"', mssanitizer)
        self.assertIn('MSS_SANITIZER_LOG_PATH', source)
        self.assertIn(': > "$MSS_LOG_PATH"', mssanitizer)
        self.assertIn('[[ ! -L "$MSS_LOG_PATH" ]]', mssanitizer)
        self.assertIn('realpath -m -- "$MSS_LOG_PATH"', mssanitizer)
        self.assertIn('realpath -m -- "$MSS_SANITIZER_LOG_PATH"', mssanitizer)
        self.assertIn(
            '[[ "$MSS_SANITIZER_LOG_PATH" == "$MSS_LOG_PATH" ]]',
            mssanitizer,
        )
        self.assertIn('sanitizer-log', mssanitizer)
        self.assertIn('"${SINGLE_PROCESS_ARGS[@]}"', mssanitizer)
        self.assertIn('"${MSS_TIMEOUT_ARGS[@]}"', mssanitizer)

    def test_sanitizer_log_requires_clean_finish_for_each_kernel(self):
        kernels = [
            "ChunkKdaFwdPrepare_171e9a19df4c0d44ca00a177d8fe300c_2570",
            "ChunkKdaFwdPrepare_171e9a19df4c0d44ca00a177d8fe300c_2571",
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            outer = root / "outer.log"
            atk = root / "atk.log"
            outer.write_text(
                "\n".join(
                    [
                        *(f"Start memcheck sanitizer on kernel {kernel}" for kernel in kernels),
                        *(
                            f"[mssanitizer] Sanitizer finished on kernel {kernel}. "
                            "No error detected."
                            for kernel in kernels
                        ),
                    ]
                )
                + "\n",
                encoding="utf-8",
            )
            atk.write_text("ATK completed\n", encoding="utf-8")
            result = VERIFIER._sanitizer_log_summary(
                (outer, atk), "memcheck", expected_finish_count=2
            )
            self.assertTrue(result["passed"])
            self.assertEqual(result["clean_finish_count"], 2)

            outer.write_text(
                outer.read_text(encoding="utf-8").replace(
                    "No error detected.", "See all detected errors above.", 1
                ),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "完成状态失败"):
                VERIFIER._sanitizer_log_summary(
                    (outer, atk), "memcheck", expected_finish_count=2
                )

            outer.write_text(
                "Start memcheck sanitizer on kernel " + kernels[0] + "\n"
                "[mssanitizer] Sanitizer finished on kernel "
                + kernels[0]
                + ". No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "数量不匹配"):
                VERIFIER._sanitizer_log_summary(
                    (outer, atk), "memcheck", expected_finish_count=2
                )

            outer.write_text(
                "Start memcheck sanitizer on kernel " + kernels[0] + "\n"
                "See all detected errors above.\n"
                "[mssanitizer] Sanitizer finished on kernel "
                + kernels[0]
                + ". No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "检测到 sanitizer 异常"):
                VERIFIER._sanitizer_log_summary(
                    (outer, atk), "memcheck", expected_finish_count=1
                )

    def test_verifier_rejects_symlink_evidence_before_reading(self):
        path = Path("evidence")
        with mock.patch.object(Path, "is_symlink", return_value=True):
            with self.assertRaisesRegex(ValueError, "符号链接"):
                VERIFIER._read_sanitizer_logs((path,))
            with self.assertRaisesRegex(ValueError, "符号链接"):
                VERIFIER._read_summary(path)
            with self.assertRaisesRegex(ValueError, "符号链接"):
                VERIFIER._report_case_ids(path, "accuracy")
            with self.assertRaisesRegex(ValueError, "符号链接"):
                VERIFIER._load_runtime_manifest(path, path, False, False)
            with self.assertRaisesRegex(ValueError, "符号链接"):
                VERIFIER.verify_runtime(SimpleNamespace(output=path))

    def test_runtime_artifact_symlink_chain_is_rejected_before_resolve(self):
        """运行时对象的精确文件和祖先目录都不能通过符号链接替换。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            artifact = root / "vendor" / "kernel" / "prepare.o"
            with mock.patch.object(
                Path,
                "is_symlink",
                return_value=True,
            ), mock.patch.object(
                Path,
                "resolve",
                side_effect=AssertionError("符号链接检查必须先于 resolve"),
            ):
                with self.assertRaisesRegex(ValueError, "符号链接"):
                    VERIFIER._reject_symlink_chain(artifact, "kernel 对象")

    def test_evidence_paths_reject_ancestor_symlink_before_resolve(self):
        """各类外部证据路径都必须检查祖先目录，而不只检查文件本身。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            evidence = root / "linked" / "evidence"
            linked_parent = evidence.parent

            checks = (
                ("runtime manifest", lambda: VERIFIER._load_runtime_manifest(
                    evidence, evidence, False, False
                )),
                ("ATK report", lambda: VERIFIER._report_case_ids(
                    evidence, "accuracy"
                )),
                ("summary", lambda: VERIFIER._read_summary(evidence)),
                ("sanitizer log", lambda: VERIFIER._read_sanitizer_logs(
                    (evidence,)
                )),
                ("console log", lambda: VERIFIER._logged_keys(evidence)),
                ("matrix root", lambda: VERIFIER._matrix_shards(evidence, 1)),
                ("matrix summary", lambda: VERIFIER._verify_or_write_matrix_summary(
                    evidence, {}
                )),
                ("shard summary", lambda: VERIFIER._verify_or_write_summary(
                    SimpleNamespace(summary=evidence), {}
                )),
                ("runtime output", lambda: VERIFIER.verify_runtime(
                    SimpleNamespace(output=evidence)
                )),
            )
            with mock.patch.object(
                Path,
                "is_symlink",
                new=lambda current: current == linked_parent,
            ), mock.patch.object(
                Path,
                "resolve",
                side_effect=AssertionError(
                    "证据祖先符号链接必须在 resolve 前被拒绝"
                ),
            ):
                for name, check in checks:
                    with self.subTest(name=name):
                        with self.assertRaisesRegex(ValueError, "符号链接"):
                            check()

    def test_candidate_opp_roots_reject_raw_paths_before_resolve(self):
        """所有 OPP 搜索根都要在 resolve 前检查原始路径及其祖先。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            package_spec = SimpleNamespace(
                submodule_search_locations=[str(root / "python/fla_npu")]
            )
            sources = (
                ("FLA_NPU_OPP_PATH", root / "opp", None),
                ("ASCEND_CUSTOM_OPP_PATH", root / "custom_opp", None),
                ("package_root/opp", None, package_spec),
            )
            for source, value, spec in sources:
                with self.subTest(source=source):
                    environment = {
                        "FLA_NPU_OP_API_LIB": "",
                        "FLA_NPU_OPP_PATH": "",
                        "ASCEND_CUSTOM_OPP_PATH": "",
                    }
                    if value is not None:
                        environment[source] = str(value)
                    with mock.patch.object(
                        VERIFIER.importlib.util, "find_spec", return_value=spec
                    ), mock.patch.dict(
                        "os.environ", environment, clear=False
                    ), mock.patch.object(
                        Path, "is_symlink", return_value=True
                    ), mock.patch.object(
                        Path,
                        "resolve",
                        side_effect=AssertionError(
                            "OPP 根目录必须先完成符号链接检查"
                        ),
                    ):
                        with self.assertRaisesRegex(ValueError, "符号链接"):
                            VERIFIER._candidate_opp_roots()

    def test_runtime_manifest_checks_artifact_symlink_boundaries(self):
        source = (ATK_DIR / "scripts/verify_matrix.py").read_text(
            encoding="utf-8"
        )
        for fragment in (
            '_reject_symlink_chain(path, "kernel metadata", kernel_dir)',
            '_reject_symlink_chain(object_path, "kernel 对象", kernel_dir)',
            '_reject_symlink_chain(path, "测试程序")',
            '_reject_symlink_chain(path, "Python 包装器", package_root)',
            '_reject_symlink_chain(candidate, "op_api 动态库", root)',
        ):
            with self.subTest(fragment=fragment):
                self.assertIn(fragment, source)

    def test_sanitizer_suite_rejects_aggregate_symlink_before_resolve(self):
        """聚合汇总路径必须先做符号链接检查，再进行规范化。"""
        cases = [
            {
                "id": case_id,
                "inputs": [
                    {
                        "name": "case_spec",
                        "range_values": json.dumps(
                            {"expected_tiling_key": 2570 + case_id}
                        ),
                    }
                ],
            }
            for case_id in range(432)
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            case_file = root / "cases.json"
            case_file.write_text(json.dumps(cases), encoding="utf-8")
            aggregates = []
            for tool in ("memcheck", "racecheck", "initcheck", "synccheck"):
                aggregate = root / tool / "aggregate_summary.json"
                aggregate.parent.mkdir()
                aggregate.write_text("{}\n", encoding="utf-8")
                aggregates.append(aggregate)
            args = SimpleNamespace(
                case_file=case_file,
                soc="ascend950",
                test_artifact=[root / "runner.py"],
                aggregate=aggregates,
                output=root / "suite.json",
            )
            linked_parent = aggregates[0].parent
            with mock.patch.object(
                Path,
                "is_symlink",
                new=lambda current: current == linked_parent,
            ), mock.patch.object(
                Path,
                "resolve",
                side_effect=AssertionError("resolve must follow symlink check"),
            ):
                with self.assertRaisesRegex(ValueError, "符号链接"):
                    VERIFIER.verify_sanitizer_suite(args)

    def test_matrix_script_rejects_root_symlink_before_mkdir_or_cd(self):
        source = (ATK_DIR / "scripts/run_matrix.sh").read_text(
            encoding="utf-8"
        )
        guard = source.index('if [[ -L "$matrix_root" ]]')
        lexical_realpath = source.index(
            'matrix_root_nosymlink=$(realpath -m -s -- "$matrix_root")', guard
        )
        physical_realpath = source.index(
            'matrix_root_realpath=$(realpath -m -- "$matrix_root")', guard
        )
        comparison = source.index(
            'if [[ "$matrix_root_nosymlink" != "$matrix_root_realpath" ]]',
            guard,
        )
        mkdir = source.index('mkdir -p "$matrix_root_realpath"', guard)
        self.assertLess(guard, mkdir)
        self.assertLess(guard, lexical_realpath)
        self.assertLess(lexical_realpath, physical_realpath)
        self.assertLess(physical_realpath, comparison)
        self.assertLess(comparison, mkdir)

    def test_sanitizer_parser_covers_official_diagnostic_variants(self):
        kernel_name = (
            "ChunkKdaFwdPrepare_171e9a19df4c0d44ca00a177d8fe300c_2570"
        )
        runtime_manifest = _runtime_manifest_fixture({2570})
        variants = {
            "memcheck": (
                "====== ERROR: LeakCheck: detected memory leaks",
                "====== WARNING: Unused memory of 128 bytes",
            ),
            "synccheck": (
                "====== WARNING: Redundant set_flag instructions detected",
                "====== ERROR: Sync error detected kernel locked up at",
                "====== ERROR: Sync error detected. Divergent thread(s) in vec_add",
            ),
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            console = root / "console.log"
            sanitizer = root / "sanitizer.log"
            for tool, diagnostics in variants.items():
                for diagnostic in diagnostics:
                    console.write_text(
                        f"Start {tool} sanitizer on kernel {kernel_name}\n"
                        f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                        "No error detected.\n",
                        encoding="utf-8",
                    )
                    sanitizer.write_text(diagnostic + "\n", encoding="utf-8")
                    with self.assertRaisesRegex(ValueError, "检测到 sanitizer 异常"):
                        VERIFIER._sanitizer_evidence(
                            console, sanitizer, tool, {2570}, runtime_manifest
                        )

            sanitizer.write_text(
                "[mssanitizer]Warning:Register FFTS_BASE_ADDR was not reset "
                f"to default in block aic(1) on kernel {kernel_name}. "
                "Expected default value is (0), but current value is (1)\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "寄存器状态异常"):
                VERIFIER._sanitizer_evidence(
                    console, sanitizer, "memcheck", {2570}, runtime_manifest
                )

    def test_readme_binds_sanitizer_build_to_matrix_soc(self):
        source = (ATK_DIR / "README.md").read_text(encoding="utf-8")
        self.assertIn("export KDA_PREPARE_ATK_SOC=ascend950", source)
        self.assertIn('FLA_NPU_SOC="$KDA_PREPARE_ATK_SOC"', source)
        self.assertIn(
            "A2 使用 ascend910b，A3 使用 ascend910_93，A5 使用 ascend950",
            source,
        )

    def test_matrix_runtime_manifest_binds_runner_and_verifier(self):
        source = (ATK_DIR / "scripts/run_matrix.sh").read_text(
            encoding="utf-8"
        )
        for artifact in (
            '"$op_dir/executor_chunk_kda_fwd_prepare.py"',
            '"$op_dir/chunk_kda_fwd_prepare.yaml"',
            '"$op_dir/gen_chunk_kda_fwd_prepare.py"',
            '"$repo_root/tests/op_cases/chunk_kda_fwd_prepare.json"',
            '"$repo_root/tests/atk/common/_ascendc_common_executor.py"',
            '"$repo_root/tests/atk/common/check_atk_result.py"',
            '"$runner"',
            '"$coverage_checker"',
            '"$script_dir/run_matrix.sh"',
            '"$verifier"',
        ):
            self.assertIn(artifact, source)
        self.assertIn(
            'runtime_args+=(--test-artifact "$artifact")', source
        )
        self.assertIn(
            'aggregate_runtime_args+=(--test-artifact "$artifact")', source
        )
        self.assertIn('aggregate_runtime_args=(--soc "$soc")', source)
        self.assertIn("正式矩阵必须显式设置 KDA_PREPARE_ATK_SOC", source)
        self.assertLess(
            source.index('python3 "$coverage_checker"'),
            source.index('python3 "$verifier" "${runtime_args[@]}"'),
        )

    def test_matrix_resume_revalidates_prefix_at_shard_boundaries(self):
        source = (ATK_DIR / "scripts/run_matrix.sh").read_text(
            encoding="utf-8"
        )
        prefix = source.split(
            "# 续跑前先从原始报告和日志重验完整前缀", 1
        )[1].split("for ((start = matrix_start;", 1)[0]
        self.assertIn("start < matrix_start", prefix)
        self.assertIn('python3 "$verifier" shard', prefix)
        self.assertIn("续跑前缀缺少完整分片", prefix)
        self.assertIn("matrix_start % shard_size", source)
        self.assertIn('[[ "$shard_size" == "1" ]]', source)
        self.assertIn("ATK_SINGLE_PROCESS=off", source)
        self.assertIn("--single-process-mode off", source)
        for name in (
            "ACCURACY_START",
            "ACCURACY_END",
            "DETERMINISM_START",
            "DETERMINISM_END",
            "MSS_START",
            "MSS_END",
        ):
            expected = "end" if name.endswith("_END") else "start"
            self.assertIn(
                f'common_env+=({name}="${expected}")',
                source,
            )
        self.assertIn(
            'common_env+=(MSS_SANITIZER_LOG_PATH="$shard_root/${tool}.log")',
            source,
        )

    def test_runner_rejects_requested_and_physical_soc_mismatch(self):
        source = (ROOT / "tests/atk/run_test_cpu.sh").read_text(
            encoding="utf-8"
        )
        self.assertIn(
            'detected_soc="$(detect_soc_from_npu "$NPU_DEVICE_ID")"', source
        )
        self.assertIn("resolve_physical_device_id()", source)
        self.assertIn("ASCEND_RT_VISIBLE_DEVICES", source)
        self.assertIn('awk -F\'|\' -v target="$physical_id"', source)
        self.assertIn('elif [[ "$SOC" != "$detected_soc" ]]', source)
        self.assertIn("目标 SoC 与实际 NPU 不一致", source)
        self.assertIn("实际 NPU SoC", source)

    def test_runner_detects_selected_a2_a3_a5_device(self):
        bash = shutil.which("bash")
        if os.name == "nt":
            git_bash = Path(os.environ.get("ProgramFiles", "C:/Program Files"))
            git_bash /= "Git/bin/bash.exe"
            if git_bash.is_file():
                bash = str(git_bash)
        if bash is None:
            self.skipTest("需要 Bash 执行 SoC 解析 fixture")

        script = r"""
set -euo pipefail
source <(sed -n '/^resolve_physical_device_id()/,/^}/p' "$1")
source <(sed -n '/^detect_soc_from_npu()/,/^}/p' "$1")
MOCK_NPU_SMI='| 4     910_9391 | OK | data |
| 5 | Ascend950PR | OK | data |
| 6     910B3 | OK | data |
| 7     9579 | OK | data |'
npu-smi() { printf '%s\n' "$MOCK_NPU_SMI"; }
export ASCEND_RT_VISIBLE_DEVICES=4,5
[[ "$(detect_soc_from_npu 0)" == ascend910_93 ]]
[[ "$(detect_soc_from_npu 1)" == ascend950 ]]
[[ -z "$(detect_soc_from_npu 2)" ]]
unset ASCEND_RT_VISIBLE_DEVICES
[[ "$(detect_soc_from_npu 6)" == ascend910b ]]
[[ "$(detect_soc_from_npu 7)" == ascend950 ]]
"""
        result = subprocess.run(
            [bash, "-c", script, "soc-fixture", "tests/atk/run_test_cpu.sh"],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_matrix_verifier_requires_matching_host_and_launch_keys(self):
        key = 8391178
        log = (
            "ChunkKdaFwdPrepare tiling: outputMode=1, "
            f"tilingKey={key}, workspace=1024\n"
            "OpName:[aclnnChunkKdaFwdPrepare_0_ChunkKdaFwdPrepare] "
            f"Tiling Key: {key}, len: 48\n"
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "console.log"
            path.write_text(log, encoding="utf-8")
            self.assertEqual(VERIFIER._logged_keys(path), ({key}, {key}))

            path.write_text(
                log.replace("outputMode=1", "outputMode=0"),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "outputMode"):
                VERIFIER._logged_keys(path)

    def test_matrix_shards_require_canonical_single_case_cover(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            (root / "shard_0_1").mkdir()
            (root / "shard_1_2").mkdir()
            self.assertEqual(
                [(start, end) for _, start, end in VERIFIER._matrix_shards(root, 2)],
                [(0, 1), (1, 2)],
            )

        invalid_layouts = (
            ("shard_00_1",),
            ("shard_0_2",),
            ("shard_1_2",),
            ("shard_0_1", "shard_2_3"),
        )
        for names in invalid_layouts:
            with self.subTest(names=names), tempfile.TemporaryDirectory() as temp_dir:
                root = Path(temp_dir)
                for name in names:
                    (root / name).mkdir()
                with self.assertRaisesRegex(
                    ValueError, "非法|单 case|不连续|未覆盖"
                ):
                    VERIFIER._matrix_shards(root, 3)

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"), "需要 openpyxl 生成测试报告"
    )
    def test_existing_shard_is_rebuilt_from_raw_evidence(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            cases = GENERATOR._payloads(GENERATOR.build_accuracy_specs())
            case_file = root / "cases.json"
            case_file.write_text(json.dumps(cases), encoding="utf-8")
            runtime_manifest = root / "runtime_manifest.json"
            runtime_manifest.write_text("{}\n", encoding="utf-8")
            console_log = root / "console.log"
            console_log.write_text("ATK finished\n", encoding="utf-8")
            report = (
                root
                / "accuracy/atk_output/atk_chunk_kda_fwd_prepare_test"
                / "report/result.xlsx"
            )
            report.parent.mkdir(parents=True)
            workbook = openpyxl.Workbook()
            statistic = workbook.active
            statistic.title = "statistic"
            statistic.append(["编号", "npu_npu_dut_精度通过", "运行结果"])
            statistic.append([0, True, "SUCCESS"])
            summary_sheet = workbook.create_sheet("summary")
            summary_sheet.append(
                [
                    "名称",
                    "总用例数",
                    "执行成功用例个数",
                    "执行失败用例个数",
                    "通过用例个数",
                    "精度是否达标",
                ]
            )
            summary_sheet.append(["cpu_cpu_golden", 1, 1, 0, 1, "Pass"])
            workbook.save(report)
            workbook.close()

            args = SimpleNamespace(
                case_file=case_file,
                scope="accuracy",
                tool="",
                runtime_manifest=runtime_manifest,
                timeout=60,
                loop_nums=1,
                gm_init_mode="off",
                single_process_mode="off",
                require_tiling_log=False,
                shard_root=root,
                console_log=console_log,
                summary=root / "summary.json",
                start=0,
                end=1,
            )
            with self.assertRaisesRegex(ValueError, "schema 不是 v2"):
                VERIFIER.summarize_shard(args)
            runtime_manifest.write_text(
                json.dumps(
                    _runtime_manifest_fixture(
                        {
                            spec["expected_tiling_key"]
                            for spec in GENERATOR.build_accuracy_specs()
                        },
                        binary_count=4,
                    )
                ),
                encoding="utf-8",
            )
            # accuracy 正式报告不能用 ATK 的“不适用”占位符伪装成通过。
            workbook = openpyxl.load_workbook(report)
            workbook["summary"]["F2"] = "-"
            workbook.save(report)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "精度达标列失败"):
                VERIFIER._report_case_ids(report, "accuracy")
            workbook = openpyxl.load_workbook(report)
            workbook["summary"]["F2"] = "Pass"
            workbook.save(report)
            workbook.close()
            self.assertEqual(VERIFIER.summarize_shard(args), 0)
            console_log.write_text("ATK evidence changed\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "原始证据不一致"):
                VERIFIER.summarize_shard(args)

            console_log.write_text("ATK finished\n", encoding="utf-8")
            report.unlink()
            with self.assertRaisesRegex(ValueError, "没有找到"):
                VERIFIER.summarize_shard(args)

    def test_aggregate_rebuilds_every_shard_before_trusting_summary(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            case_file = root / "cases.json"
            cases = GENERATOR._payloads(GENERATOR.build_accuracy_specs())
            case_file.write_text(json.dumps(cases), encoding="utf-8")
            runtime_manifest = root / "runtime_manifest.json"
            runtime_manifest.write_text("{}\n", encoding="utf-8")
            runtime_hash = hashlib.sha256(runtime_manifest.read_bytes()).hexdigest()
            case_hash = hashlib.sha256(case_file.read_bytes()).hexdigest()

            canonical = {}
            for case_id in range(200):
                shard = root / f"shard_{case_id}_{case_id + 1}"
                shard.mkdir()
                value = {
                    "schema": "kda-prepare-atk-shard/v2",
                    "scope": "accuracy",
                    "tool": "",
                    "timeout_seconds": 60,
                    "loop_nums": 1,
                    "gm_init_mode": "off",
                    "single_process_mode": "off",
                    "require_tiling_log": False,
                    "runtime_manifest_sha256": runtime_hash,
                    "start": case_id,
                    "end": case_id + 1,
                    "expected_cases": 1,
                    "case_ids": [case_id],
                    "case_json_sha256": case_hash,
                    "expected_tiling_keys": [],
                    "host_tiling_keys": [],
                    "launch_tiling_keys": [],
                    "key_pair_sha256": hashlib.sha256(b"").hexdigest(),
                    "passed": True,
                }
                canonical[case_id] = value
                (shard / "summary.json").write_text(
                    json.dumps(value), encoding="utf-8"
                )

            args = SimpleNamespace(
                case_file=case_file,
                scope="accuracy",
                tool="",
                runtime_manifest=runtime_manifest,
                timeout=60,
                loop_nums=1,
                gm_init_mode="off",
                single_process_mode="off",
                require_tiling_log=False,
                matrix_root=root,
                soc="ascend950",
                test_artifact=[root / "runner.py"],
                output=root / "aggregate_summary.json",
            )

            def rebuild(shard_args):
                return canonical[shard_args.start]

            with mock.patch.object(
                VERIFIER, "_verify_current_runtime", return_value={}
            ), mock.patch.object(
                VERIFIER, "_build_shard_summary", side_effect=rebuild
            ) as builder:
                self.assertEqual(VERIFIER.aggregate_matrix(args), 0)
                self.assertEqual(builder.call_count, 200)

            tampered = dict(canonical[0])
            tampered["passed"] = False
            (root / "shard_0_1/summary.json").write_text(
                json.dumps(tampered), encoding="utf-8"
            )
            with mock.patch.object(
                VERIFIER, "_verify_current_runtime", return_value={}
            ), mock.patch.object(
                VERIFIER, "_build_shard_summary", side_effect=rebuild
            ):
                with self.assertRaisesRegex(ValueError, "原始证据不一致"):
                    VERIFIER.aggregate_matrix(args)

    def test_matrix_verifier_requires_active_sanitizer_kernel(self):
        kernel_name = (
            "ChunkKdaFwdPrepare_171e9a19df4c0d44ca00a177d8fe300c_2570"
        )
        runtime_manifest = _runtime_manifest_fixture({2570})
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            console = root / "console.log"
            sanitizer = root / "memcheck.log"
            console.write_text(
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            sanitizer.write_text("memcheck completed\n", encoding="utf-8")
            count, keys, names, console_hash, sanitizer_hash = (
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    runtime_manifest,
                )
            )
            self.assertEqual(count, 1)
            self.assertEqual(keys, {2570})
            self.assertEqual(names, {kernel_name})
            self.assertEqual(len(console_hash), 64)
            self.assertEqual(len(sanitizer_hash), 64)

            # 同一 key 可以有多个安装包候选，但一次 case 只能命中一个。
            multi_manifest = _runtime_manifest_fixture(
                {2570}, sanitizer=True, binary_count=4
            )
            expected_case = {
                "id": 0,
                "inputs": [
                    {
                        "name": "case_spec",
                        "range_values": {
                            "expected_tiling_key": 2570,
                            "gate_dtype": "bf16",
                            "beta_dtype": "bf16",
                        },
                    }
                ],
            }
            candidate_names = sorted(
                VERIFIER._manifest_key_to_kernels(multi_manifest)[2570]
            )
            selected = VERIFIER._manifest_dispatch_map(multi_manifest)[
                (2570, "bf16", "bf16")
            ]
            wrong_dtype_candidate = next(
                name for name in candidate_names if name != selected
            )
            console.write_text("case completed\n", encoding="utf-8")
            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {selected}\n"
                f"[mssanitizer] Sanitizer finished on kernel {selected}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            count, keys, names, _, _ = VERIFIER._sanitizer_evidence(
                console,
                sanitizer,
                "memcheck",
                {2570},
                multi_manifest,
                [expected_case],
            )
            self.assertEqual((count, keys, names), (1, {2570}, {selected}))

            console_mirrors = (
                f"Start memcheck sanitizer on kernel {selected}\n",
                "[mssanitizer] Sanitizer finished on kernel "
                f"{selected}. No error detected.\n",
                f"Start memcheck sanitizer on kernel {selected}\n"
                "[mssanitizer] Sanitizer finished on kernel "
                f"{selected}. No error detected.\n",
            )
            for mirror in console_mirrors:
                console.write_text(mirror, encoding="utf-8")
                count, keys, names, _, _ = VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )
                self.assertEqual(
                    (count, keys, names), (1, {2570}, {selected})
                )
            console.write_text("case completed\n", encoding="utf-8")

            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {wrong_dtype_candidate}\n"
                "[mssanitizer] Sanitizer finished on kernel "
                f"{wrong_dtype_candidate}. No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "manifest 之外"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )

            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {selected}\n"
                f"[mssanitizer] Sanitizer finished on kernel {selected}. "
                "No error detected.\n",
                encoding="utf-8",
            )

            console.write_text(
                f"Start memcheck sanitizer on kernel {wrong_dtype_candidate}\n"
                "[mssanitizer] Sanitizer finished on kernel "
                f"{wrong_dtype_candidate}. No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "记录不一致"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )
            console.write_text("case completed\n", encoding="utf-8")

            sanitizer.write_text(
                "\n".join(
                    [
                        f"Start memcheck sanitizer on kernel {candidate_names[0]}",
                        f"Start memcheck sanitizer on kernel {candidate_names[1]}",
                        (
                            "[mssanitizer] Sanitizer finished on kernel "
                            f"{candidate_names[0]}. No error detected."
                        ),
                        (
                            "[mssanitizer] Sanitizer finished on kernel "
                            f"{candidate_names[1]}. No error detected."
                        ),
                    ]
                )
                + "\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "manifest 之外"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )

            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {selected}\n"
                "[mssanitizer] Sanitizer finished on kernel "
                f"{wrong_dtype_candidate}. No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "非目标 kernel|缺少目标 kernel"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )

            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {selected}\n"
                f"Start memcheck sanitizer on kernel {selected}\n"
                f"[mssanitizer] Sanitizer finished on kernel {selected}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "Start 次数不为 1"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )

            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {selected}\n"
                f"[mssanitizer] Sanitizer finished on kernel {selected}. "
                "No error detected.\n"
                f"[mssanitizer] Sanitizer finished on kernel {selected}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "Finish 次数不为 1"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    multi_manifest,
                    [expected_case],
                )

            # 后续异常分支继续使用单候选 fixture。
            console.write_text(
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            sanitizer.write_text("memcheck completed\n", encoding="utf-8")

            stale_kernel = (
                "ChunkKdaFwdPrepare_aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa_2570"
            )
            console.write_text(
                f"Start memcheck sanitizer on kernel {stale_kernel}\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "manifest 之外"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    runtime_manifest,
                )

            sanitizer.write_text(
                "No active sanitizer tool on kernel "
                "ChunkKdaFwdPrepare_171e9a19df4c0d44ca00a177d8fe300c_2570\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "未加载 sanitizer"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    runtime_manifest,
                )

            diagnostics = {
                "memcheck": "========= WARNING: out of bounds of size 2048",
                "racecheck": "========= ERROR: Potential RAW hazard detected",
                "initcheck": "========= ERROR: uninitialized read of size 224",
                "synccheck": "========= WARNING: Unpaired set_flag instructions detected",
            }
            for tool, diagnostic in diagnostics.items():
                for destination in ("console", "sanitizer"):
                    with self.subTest(tool=tool, destination=destination):
                        start = (
                            f"Start {tool} sanitizer on kernel {kernel_name}\n"
                            f"[mssanitizer] Sanitizer finished on kernel "
                            f"{kernel_name}. No error detected.\n"
                        )
                        console.write_text(
                            start + (diagnostic + "\n" if destination == "console" else ""),
                            encoding="utf-8",
                        )
                        sanitizer.write_text(
                            diagnostic + "\n" if destination == "sanitizer" else "completed\n",
                            encoding="utf-8",
                        )
                        with self.assertRaisesRegex(
                            ValueError, "检测到 sanitizer 异常"
                        ):
                            VERIFIER._sanitizer_evidence(
                                console,
                                sanitizer,
                                tool,
                                {2570},
                                runtime_manifest,
                            )

            for diagnostic in (
                "ERROR SUMMARY: 1",
            ):
                with self.subTest(diagnostic=diagnostic):
                    console.write_text(
                        f"Start memcheck sanitizer on kernel {kernel_name}\n"
                        f"[mssanitizer] Sanitizer finished on kernel "
                        f"{kernel_name}. No error detected.\n",
                        encoding="utf-8",
                    )
                    sanitizer.write_text(diagnostic + "\n", encoding="utf-8")
                    with self.assertRaisesRegex(
                        ValueError, "检测到 sanitizer 异常"
                    ):
                        VERIFIER._sanitizer_evidence(
                            console,
                            sanitizer,
                            "memcheck",
                            {2570},
                            runtime_manifest,
                        )

            for tool in sorted(VERIFIER.SANITIZER_TOOLS):
                with self.subTest(tool=tool, finish="failed"):
                    console.write_text("case completed\n", encoding="utf-8")
                    sanitizer.write_text(
                        f"Start {tool} sanitizer on kernel {kernel_name}\n"
                        f"[mssanitizer] Sanitizer finished on kernel "
                        f"{kernel_name}. See all detected errors above.\n",
                        encoding="utf-8",
                    )
                    with self.assertRaisesRegex(ValueError, "完成状态失败"):
                        VERIFIER._sanitizer_evidence(
                            console,
                            sanitizer,
                            tool,
                            {2570},
                            runtime_manifest,
                        )

                with self.subTest(tool=tool, finish="missing"):
                    console.write_text(
                        f"Start {tool} sanitizer on kernel {kernel_name}\n",
                        encoding="utf-8",
                    )
                    sanitizer.write_text("completed\n", encoding="utf-8")
                    with self.assertRaisesRegex(ValueError, "缺少目标 kernel"):
                        VERIFIER._sanitizer_evidence(
                            console,
                            sanitizer,
                            tool,
                            {2570},
                            runtime_manifest,
                        )

            console.write_text(
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            sanitizer.write_text(
                "[mssanitizer] Warning:Register FFTS_BASE_ADDR was not reset "
                f"to default in block aic(1) on kernel {kernel_name}. "
                "Expected default value is (0), but current value is (1)\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "寄存器状态异常"):
                VERIFIER._sanitizer_evidence(
                    console,
                    sanitizer,
                    "memcheck",
                    {2570},
                    runtime_manifest,
                )

            sanitizer.write_text("ERROR SUMMARY: 0\n", encoding="utf-8")
            count, keys, names, _, _ = VERIFIER._sanitizer_evidence(
                console,
                sanitizer,
                "memcheck",
                {2570},
                runtime_manifest,
            )
            self.assertEqual((count, keys, names), (1, {2570}, {kernel_name}))

            for tool in sorted(VERIFIER.SANITIZER_TOOLS):
                with self.subTest(tool=tool, finish="success"):
                    console.write_text("case completed\n", encoding="utf-8")
                    sanitizer.write_text(
                        f"Start {tool} sanitizer on kernel {kernel_name}\n"
                        f"[mssanitizer] Sanitizer finished on kernel "
                        f"{kernel_name}. No error detected.\n",
                        encoding="utf-8",
                    )
                    count, keys, names, _, _ = VERIFIER._sanitizer_evidence(
                        console,
                        sanitizer,
                        tool,
                        {2570},
                        runtime_manifest,
                    )
                    self.assertEqual(
                        (count, keys, names), (1, {2570}, {kernel_name})
                    )

            console.write_text(
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                "No error detected.\n",
                encoding="utf-8",
            )
            sanitizer.write_text(
                "ERROR: model output mismatch\n========= INFO: completed\n",
                encoding="utf-8",
            )
            count, keys, names, _, _ = VERIFIER._sanitizer_evidence(
                console,
                sanitizer,
                "memcheck",
                {2570},
                runtime_manifest,
            )
            self.assertEqual((count, keys, names), (1, {2570}, {kernel_name}))

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"), "需要 openpyxl 生成测试报告"
    )
    def test_mssanitizer_raw_error_overrides_passing_atk_report(self):
        import openpyxl

        keys = set(range(2570, 2570 + 432))
        kernel_name = (
            "ChunkKdaFwdPrepare_171e9a19df4c0d44ca00a177d8fe300c_2570"
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            cases = [
                {
                    "id": case_id,
                    "inputs": [
                        {
                            "name": "case_spec",
                            "range_values": json.dumps(
                                {
                                    "expected_tiling_key": 2570 + case_id,
                                    "gate_dtype": "bf16",
                                    "beta_dtype": "bf16",
                                }
                            ),
                        }
                    ],
                }
                for case_id in range(432)
            ]
            case_file = root / "cases.json"
            case_file.write_text(json.dumps(cases), encoding="utf-8")
            runtime_manifest = root / "runtime_manifest.json"
            runtime_manifest.write_text(
                json.dumps(
                    _runtime_manifest_fixture(
                        keys, sanitizer=True, complete=True
                    )
                ),
                encoding="utf-8",
            )

            report = (
                root
                / "mssanitizer/atk_output/atk_chunk_kda_fwd_prepare_test"
                / "report/result.xlsx"
            )
            report.parent.mkdir(parents=True)
            workbook = openpyxl.Workbook()
            statistic = workbook.active
            statistic.title = "statistic"
            statistic.append(["编号", "npu_dut_内存检测通过", "运行结果"])
            statistic.append([0, True, "SUCCESS"])
            summary = workbook.create_sheet("summary")
            summary.append(
                ["名称", "总用例数", "内存检测通过率", "内存检测是否达标"]
            )
            summary.append(["npu_dut", 1, "100%", "Pass"])
            workbook.save(report)
            workbook.close()

            console = root / "console.log"
            console.write_text(
                "ChunkKdaFwdPrepare tiling: outputMode=0, tilingKey=2570\n"
                "OpName:[ChunkKdaFwdPrepare] Tiling Key: 2570\n",
                encoding="utf-8",
            )
            sanitizer = root / "memcheck.log"
            sanitizer.write_text(
                "====== WARNING: out of bounds of size 2048\n"
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                "See all detected errors above.\n",
                encoding="utf-8",
            )
            args = SimpleNamespace(
                case_file=case_file,
                scope="mssanitizer",
                tool="memcheck",
                runtime_manifest=runtime_manifest,
                timeout=60,
                loop_nums=1,
                gm_init_mode="not_applicable",
                single_process_mode="off",
                require_tiling_log=True,
                shard_root=root,
                console_log=console,
                summary=root / "summary.json",
                start=0,
                end=1,
                sanitizer_log=sanitizer,
            )
            with self.assertRaisesRegex(ValueError, "sanitizer 异常"):
                VERIFIER._build_shard_summary(args)

            clean_finish = (
                f"[mssanitizer] Sanitizer finished on kernel {kernel_name}. "
                "No error detected.\n"
            )
            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                + clean_finish,
                encoding="utf-8",
            )
            summary_value = VERIFIER._build_shard_summary(args)
            self.assertEqual(
                summary_value["sanitizer_outer_log_sha256"],
                hashlib.sha256(sanitizer.read_bytes()).hexdigest(),
            )
            sanitizer.write_text(
                f"Start memcheck sanitizer on kernel {kernel_name}\n"
                + clean_finish.replace(kernel_name, kernel_name[:-1] + "1"),
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "kernel 集合不一致"):
                VERIFIER._build_shard_summary(args)
            sanitizer.write_text("changed\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "sanitizer 完成状态|缺少 sanitizer"):
                VERIFIER._build_shard_summary(args)

    def test_matrix_verifier_rejects_noncanonical_scope_contracts(self):
        valid = {
            "accuracy": SimpleNamespace(
                scope="accuracy",
                timeout=60,
                loop_nums=1,
                gm_init_mode="off",
                single_process_mode="off",
                tool="",
                require_tiling_log=False,
            ),
            "determinism": SimpleNamespace(
                scope="determinism",
                timeout=60,
                loop_nums=50,
                gm_init_mode="not_applicable",
                single_process_mode="off",
                tool="",
                require_tiling_log=True,
            ),
            "mssanitizer": SimpleNamespace(
                scope="mssanitizer",
                timeout=60,
                loop_nums=1,
                gm_init_mode="not_applicable",
                single_process_mode="off",
                tool="memcheck",
                require_tiling_log=True,
            ),
        }
        case_counts = {"accuracy": 200, "determinism": 432, "mssanitizer": 432}
        for scope, args in valid.items():
            with self.subTest(scope=scope):
                VERIFIER._validate_scope_contract(args, case_counts[scope])

        invalid = (
            ("timeout", valid["accuracy"], 200, {"timeout": 0}),
            ("timeout", valid["accuracy"], 200, {"timeout": 61}),
            ("正式矩阵", valid["accuracy"], 199, {}),
            ("loop_nums", valid["accuracy"], 200, {"loop_nums": 2}),
            ("gm_init_mode", valid["accuracy"], 200, {"gm_init_mode": "on"}),
            (
                "single_process_mode",
                valid["accuracy"],
                200,
                {"single_process_mode": "on"},
            ),
            ("tool", valid["accuracy"], 200, {"tool": "memcheck"}),
            (
                "require_tiling_log",
                valid["accuracy"],
                200,
                {"require_tiling_log": True},
            ),
            (
                "require_tiling_log",
                valid["determinism"],
                432,
                {"require_tiling_log": False},
            ),
            ("tool", valid["mssanitizer"], 432, {"tool": "invalid"}),
        )
        for message, base_args, case_count, override in invalid:
            values = vars(base_args).copy()
            values.update(override)
            with self.subTest(message=message, values=values):
                with self.assertRaisesRegex(ValueError, message):
                    VERIFIER._validate_scope_contract(
                        SimpleNamespace(**values), case_count
                    )

    @mock.patch.object(
        VERIFIER,
        "_runtime_toolchain",
        return_value=_toolchain_fixture(False),
    )
    def test_runtime_manifest_binds_compiled_tiling_keys(self, _mock_toolchain):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            kernel_dir = (
                root
                / "vendors/fla_npu_transformer/op_impl/ai_core/tbe/kernel"
                / "ascend950/chunk_kda_fwd_prepare"
            )
            kernel_dir.mkdir(parents=True)
            compile_hash = "171e9a19df4c0d44ca00a177d8fe300c"
            binary_name = f"ChunkKdaFwdPrepare_{compile_hash}"
            object_path = kernel_dir / f"{binary_name}.o"
            object_path.write_bytes(b"kernel")
            object_sha256 = hashlib.sha256(object_path.read_bytes()).hexdigest()
            (kernel_dir / f"{binary_name}.json").write_text(
                json.dumps(
                    {
                        "binFileName": binary_name,
                        "binFileSuffix": ".o",
                        "kernelName": binary_name,
                        "sha256": object_sha256,
                        "supportInfo": _support_info_fixture("bf16", "bf16"),
                        "kernelList": [
                            {
                                "tilingKey": 2570,
                                "kernelName": f"{binary_name}_2570",
                            },
                            {
                                "tilingKey": 8391178,
                                "kernelName": f"{binary_name}_8391178",
                            },
                        ],
                    }
                ),
                encoding="utf-8",
            )
            op_api_lib = (
                root
                / "vendors/fla_npu_transformer/op_api/lib/libcust_opapi.so"
            )
            op_api_lib.parent.mkdir(parents=True)
            op_api_lib.write_bytes(b"op_api")
            package_root = root / "python/fla_npu"
            for relative in (
                "__init__.py",
                "ops/ascendc/__init__.py",
                "ops/ascendc/_aclnn_ctypes.py",
                "ops/ascendc/_kda_policy.py",
                "ops/ascendc/_runtime.py",
            ):
                path = package_root / relative
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(f"# {relative}\n", encoding="utf-8")
            package_spec = SimpleNamespace(
                submodule_search_locations=[str(package_root)]
            )
            case_file = root / "cases.json"
            case_file.write_text(
                json.dumps(
                    [
                        {
                            "id": 0,
                            "inputs": [
                                {
                                    "name": "case_spec",
                                    "range_values": json.dumps(
                                        {
                                            "expected_tiling_key": 2570,
                                            "gate_dtype": "bf16",
                                            "beta_dtype": "bf16",
                                        }
                                    ),
                                }
                            ],
                        }
                    ]
                ),
                encoding="utf-8",
            )
            with mock.patch.object(
                VERIFIER.importlib.util, "find_spec", return_value=package_spec
            ), mock.patch.dict(
                "os.environ",
                {
                    "FLA_NPU_OPP_PATH": str(root),
                    "FLA_NPU_OP_API_LIB": str(op_api_lib),
                    "ASCEND_CUSTOM_OPP_PATH": "",
                },
                clear=False,
            ):
                manifest = VERIFIER._runtime_manifest(
                    case_file, "ascend950", False, False
                )
            self.assertEqual(manifest["platform"], "ascend950")
            self.assertEqual(manifest["object_count"], 1)
            self.assertEqual(manifest["compiled_tiling_key_count"], 2)
            self.assertEqual(manifest["schema"], "kda-prepare-runtime/v2")
            self.assertEqual(manifest["kernel_name_count"], 2)
            self.assertEqual(
                manifest["kernel_binaries"][0]["compile_hash"], compile_hash
            )
            self.assertEqual(
                manifest["kernel_binaries"][0]["bin_file_sha256"],
                object_sha256,
            )
            with mock.patch.object(
                VERIFIER.importlib.util, "find_spec", return_value=package_spec
            ), mock.patch.dict(
                "os.environ",
                {
                    "FLA_NPU_OPP_PATH": str(root),
                    "FLA_NPU_OP_API_LIB": str(op_api_lib),
                    "ASCEND_CUSTOM_OPP_PATH": "",
                },
                clear=False,
            ):
                with self.assertRaisesRegex(ValueError, "编译 key"):
                    VERIFIER._runtime_manifest(
                        case_file, "ascend950", False, True
                    )

            # 不同 hashed 编译对象可以为同一个 TilingKey 提供多个候选。
            second_hash = "22222222222222222222222222222222"
            second_name = f"ChunkKdaFwdPrepare_{second_hash}"
            second_object = kernel_dir / f"{second_name}.o"
            second_object.write_bytes(b"kernel-2")
            second_sha256 = hashlib.sha256(second_object.read_bytes()).hexdigest()
            (kernel_dir / f"{second_name}.json").write_text(
                json.dumps(
                    {
                        "binFileName": second_name,
                        "binFileSuffix": ".o",
                        "kernelName": second_name,
                        "sha256": second_sha256,
                        "supportInfo": _support_info_fixture("fp32", "bf16"),
                        "kernelList": [
                            {
                                "tilingKey": 2570,
                                "kernelName": f"{second_name}_2570",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            with mock.patch.object(
                VERIFIER.importlib.util, "find_spec", return_value=package_spec
            ), mock.patch.dict(
                "os.environ",
                {
                    "FLA_NPU_OPP_PATH": str(root),
                    "FLA_NPU_OP_API_LIB": str(op_api_lib),
                    "ASCEND_CUSTOM_OPP_PATH": "",
                },
                clear=False,
            ):
                manifest = VERIFIER._runtime_manifest(
                    case_file, "ascend950", False, False
                )
            self.assertEqual(manifest["object_count"], 2)
            self.assertEqual(manifest["compiled_tiling_key_count"], 2)
            self.assertEqual(manifest["kernel_name_count"], 3)
            _, candidates = VERIFIER._manifest_kernel_maps(manifest)
            self.assertEqual(len(candidates[2570]), 2)
            self.assertEqual(len(candidates[8391178]), 1)

    def test_runtime_manifest_reader_accepts_multiple_candidates_per_key(self):
        keys = {
            int(spec["expected_tiling_key"])
            for spec in GENERATOR.build_mss_specs()
        }
        manifest = _runtime_manifest_fixture(keys, binary_count=4)
        kernel_map, candidates = VERIFIER._manifest_kernel_maps(manifest)
        self.assertEqual(len(keys), 432)
        self.assertEqual(len(kernel_map), 4 * 432)
        self.assertEqual(set(kernel_map.values()), keys)
        self.assertEqual(len(candidates[2570]), 4)
        self.assertEqual(
            len(VERIFIER._manifest_dispatch_map(manifest)), 4 * 432
        )
        mss_cases = json.loads(
            (ATK_DIR / "atk_chunk_kda_fwd_prepare_mss.json").read_text(
                encoding="utf-8"
            )
        )
        bindings = VERIFIER._case_kernel_bindings(manifest, mss_cases)
        self.assertEqual(len(bindings), 432)
        self.assertEqual(len({item["kernel_name"] for item in bindings}), 432)
        for binding, case in zip(bindings, mss_cases):
            spec = VERIFIER._case_spec(case)
            self.assertEqual(binding["tiling_key"], spec["expected_tiling_key"])
            self.assertEqual(binding["gate_dtype"], spec["gate_dtype"])
            self.assertEqual(binding["beta_dtype"], spec["beta_dtype"])

        duplicate = json.loads(json.dumps(manifest))
        duplicate["kernel_binaries"].append(
            json.loads(json.dumps(duplicate["kernel_binaries"][0]))
        )
        duplicate["object_count"] += 1
        duplicate["metadata_count"] += 1
        with self.assertRaisesRegex(ValueError, "kernelName 重复"):
            VERIFIER._manifest_kernel_map(duplicate)

        duplicate_dispatch = json.loads(json.dumps(manifest))
        duplicate_dispatch["kernel_binaries"][1]["dispatch_signature"] = (
            duplicate_dispatch["kernel_binaries"][0]["dispatch_signature"]
        )
        with self.assertRaisesRegex(ValueError, "调度签名重复"):
            VERIFIER._manifest_kernel_map(duplicate_dispatch)

        only_bf16 = _runtime_manifest_fixture({2570})
        fp32_case = {
            "id": 0,
            "inputs": [
                {
                    "name": "case_spec",
                    "range_values": {
                        "expected_tiling_key": 2570,
                        "gate_dtype": "fp32",
                        "beta_dtype": "bf16",
                    },
                }
            ],
        }
        with self.assertRaisesRegex(ValueError, "没有匹配输入 dtype"):
            VERIFIER._case_kernel_bindings(only_bf16, [fp32_case])

    def test_metadata_dispatch_contract_rejects_ambiguous_inputs(self):
        missing_g = _support_info_fixture("bf16", "bf16")
        missing_g["inputs"] = [
            item for item in missing_g["inputs"] if item["name"] != "g"
        ]
        with self.assertRaisesRegex(ValueError, "缺少 g"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": missing_g}, "missing-g.json"
            )

        duplicate_name = _support_info_fixture("bf16", "bf16")
        duplicate_name["inputs"].append(dict(duplicate_name["inputs"][3]))
        with self.assertRaisesRegex(ValueError, "name 缺失或重复"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": duplicate_name}, "duplicate-name.json"
            )

        duplicate_index = _support_info_fixture("bf16", "bf16")
        duplicate_index["inputs"][4]["index"] = 3
        with self.assertRaisesRegex(ValueError, "index 非法或重复"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": duplicate_index}, "duplicate-index.json"
            )

        wrong_mode = _support_info_fixture("bf16", "bf16")
        wrong_mode["simplifiedKeyMode"] = 1
        with self.assertRaisesRegex(ValueError, "simplifiedKeyMode"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": wrong_mode}, "wrong-mode.json"
            )

        wrong_g_token = _support_info_fixture("bf16", "bf16")
        fields = wrong_g_token["simplifiedKey"][0].split("/")
        fields[5] = "0,2"
        wrong_g_token["simplifiedKey"][0] = "/".join(fields)
        with self.assertRaisesRegex(ValueError, "g 不一致"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": wrong_g_token}, "wrong-g-token.json"
            )

        wrong_beta_token = _support_info_fixture("bf16", "bf16")
        fields = wrong_beta_token["simplifiedKey"][0].split("/")
        fields[6] = "0,2"
        wrong_beta_token["simplifiedKey"][0] = "/".join(fields)
        with self.assertRaisesRegex(ValueError, "beta 不一致"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": wrong_beta_token}, "wrong-beta-token.json"
            )

        wrong_field_count = _support_info_fixture("bf16", "bf16")
        fields = wrong_field_count["simplifiedKey"][0].split("/")
        wrong_field_count["simplifiedKey"][0] = "/".join(fields[:-1])
        with self.assertRaisesRegex(ValueError, "结构非法"):
            VERIFIER._metadata_dispatch_contract(
                {"supportInfo": wrong_field_count}, "wrong-fields.json"
            )

    def test_all_scopes_reject_forged_runtime_manifest(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            case_file = root / "cases.json"
            case_file.write_text(
                json.dumps(
                    [
                        {
                            "id": 0,
                            "inputs": [
                                {
                                    "name": "case_spec",
                                    "range_values": json.dumps(
                                        {
                                            "expected_tiling_key": 2570,
                                            "gate_dtype": "bf16",
                                            "beta_dtype": "bf16",
                                        }
                                    ),
                                }
                            ],
                        }
                    ]
                ),
                encoding="utf-8",
            )
            manifest_path = root / "runtime_manifest.json"
            manifest_path.write_text("{}\n", encoding="utf-8")
            for scope in ("accuracy", "determinism", "mssanitizer"):
                with self.subTest(scope=scope), self.assertRaisesRegex(
                    ValueError, "schema 不是 v2"
                ):
                    VERIFIER._load_runtime_manifest(
                        manifest_path,
                        case_file,
                        scope == "mssanitizer",
                        scope != "accuracy",
                    )

    def test_aggregate_rejects_stale_runtime_manifest(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            case_file = root / "cases.json"
            case_file.write_text(
                json.dumps(
                    [
                        {
                            "id": 0,
                            "inputs": [
                                {
                                    "name": "case_spec",
                                    "range_values": json.dumps(
                                        {
                                            "expected_tiling_key": 2570,
                                            "gate_dtype": "bf16",
                                            "beta_dtype": "bf16",
                                        }
                                    ),
                                }
                            ],
                        }
                    ]
                ),
                encoding="utf-8",
            )
            manifest = _runtime_manifest_fixture({2570})
            manifest_path = root / "runtime_manifest.json"
            manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            args = SimpleNamespace(
                case_file=case_file,
                runtime_manifest=manifest_path,
                scope="accuracy",
                soc="ascend950",
                test_artifact=[root / "runner.py"],
            )
            with mock.patch.object(
                VERIFIER, "_runtime_manifest", return_value=manifest
            ):
                self.assertEqual(VERIFIER._verify_current_runtime(args), manifest)

            changed = dict(manifest)
            changed["runtime_sha256"] = "b" * 64
            with mock.patch.object(
                VERIFIER, "_runtime_manifest", return_value=changed
            ):
                with self.assertRaisesRegex(ValueError, "当前实际"):
                    VERIFIER._verify_current_runtime(args)

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"), "需要 openpyxl 生成测试报告"
    )
    def test_matrix_verifier_parses_mssanitizer_columns_by_name(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            report = Path(temp_dir) / "mss.xlsx"
            workbook = openpyxl.Workbook()
            statistic = workbook.active
            statistic.title = "statistic"
            statistic.append(
                ["编号", "npu_npu_dut_内存检测通过", "运行结果"]
            )
            statistic.append([0, True, "SUCCESS"])
            statistic.append([1, True, "SUCCESS"])
            summary = workbook.create_sheet("summary")
            summary.append(
                ["名称", "总用例数", "内存检测通过率", "内存检测是否达标"]
            )
            summary.append(["npu_npu_dut", 2, 100.0, "Pass"])
            workbook.save(report)
            workbook.close()

            self.assertEqual(
                VERIFIER._report_case_ids(report, "mssanitizer"), ([0, 1], 2)
            )

    @unittest.skipUnless(
        importlib.util.find_spec("openpyxl"), "需要 openpyxl 生成测试报告"
    )
    def test_sanitizer_suite_integrates_dtype_bound_shards(self):
        import openpyxl

        source_cases = json.loads(
            (ATK_DIR / "atk_chunk_kda_fwd_prepare_mss.json").read_text(
                encoding="utf-8"
            )
        )
        selected_cases = []
        for gate_dtype, beta_dtype in (
            ("bf16", "bf16"),
            ("fp32", "bf16"),
        ):
            source = next(
                case
                for case in source_cases
                if (
                    VERIFIER._case_spec(case)["gate_dtype"],
                    VERIFIER._case_spec(case)["beta_dtype"],
                )
                == (gate_dtype, beta_dtype)
            )
            case = json.loads(json.dumps(source))
            case_id = len(selected_cases)
            case["id"] = case_id
            for item in case["inputs"]:
                if item.get("name") == "case_spec":
                    spec = json.loads(item["range_values"])
                    spec["case_id"] = case_id
                    item["range_values"] = json.dumps(spec)
                    break
            selected_cases.append(case)

        keys = {
            int(VERIFIER._case_spec(case)["expected_tiling_key"])
            for case in selected_cases
        }
        manifest = _runtime_manifest_fixture(
            keys,
            sanitizer=True,
            complete=True,
            binary_count=4,
        )
        bindings = {
            int(item["case_id"]): item
            for item in VERIFIER._case_kernel_bindings(
                manifest, selected_cases
            )
        }

        with tempfile.TemporaryDirectory() as temp_dir, mock.patch.dict(
            VERIFIER.SCOPE_CONTRACTS["mssanitizer"], {"case_count": 2}
        ), mock.patch.object(
            VERIFIER, "_verify_current_runtime", return_value=manifest
        ):
            root = Path(temp_dir)
            case_file = root / "cases.json"
            case_file.write_text(
                json.dumps(selected_cases), encoding="utf-8"
            )
            aggregates = []
            for tool in ("memcheck", "racecheck", "initcheck", "synccheck"):
                matrix_root = root / tool
                matrix_root.mkdir()
                manifest_path = matrix_root / "runtime_manifest.json"
                manifest_path.write_text(
                    json.dumps(manifest), encoding="utf-8"
                )
                for case_id, case in enumerate(selected_cases):
                    shard_root = matrix_root / f"shard_{case_id}_{case_id + 1}"
                    report = (
                        shard_root
                        / "mssanitizer/atk_output"
                        / "atk_chunk_kda_fwd_prepare_fixture"
                        / "report/result.xlsx"
                    )
                    report.parent.mkdir(parents=True)
                    workbook = openpyxl.Workbook()
                    statistic = workbook.active
                    statistic.title = "statistic"
                    statistic.append(
                        ["编号", "npu_dut_内存检测通过", "运行结果"]
                    )
                    statistic.append([case_id, True, "SUCCESS"])
                    summary_sheet = workbook.create_sheet("summary")
                    summary_sheet.append(
                        ["名称", "总用例数", "内存检测通过率", "内存检测是否达标"]
                    )
                    summary_sheet.append(["npu_dut", 1, "100%", "Pass"])
                    workbook.save(report)
                    workbook.close()

                    spec = VERIFIER._case_spec(case)
                    tiling_key = int(spec["expected_tiling_key"])
                    console = shard_root / "console.log"
                    console.write_text(
                        "ChunkKdaFwdPrepare tiling: "
                        f"outputMode=0, tilingKey={tiling_key}\n"
                        "OpName:[ChunkKdaFwdPrepare] "
                        f"Tiling Key: {tiling_key}\n",
                        encoding="utf-8",
                    )
                    kernel_name = str(bindings[case_id]["kernel_name"])
                    sanitizer_log = shard_root / f"{tool}.log"
                    sanitizer_log.write_text(
                        f"Start {tool} sanitizer on kernel {kernel_name}\n"
                        "[mssanitizer] Sanitizer finished on kernel "
                        f"{kernel_name}. No error detected.\n",
                        encoding="utf-8",
                    )
                    shard_args = SimpleNamespace(
                        case_file=case_file,
                        scope="mssanitizer",
                        tool=tool,
                        runtime_manifest=manifest_path,
                        timeout=60,
                        loop_nums=1,
                        gm_init_mode="not_applicable",
                        single_process_mode="off",
                        require_tiling_log=True,
                        shard_root=shard_root,
                        console_log=console,
                        summary=shard_root / "summary.json",
                        start=case_id,
                        end=case_id + 1,
                        sanitizer_log=sanitizer_log,
                    )
                    shard_summary = VERIFIER._build_shard_summary(shard_args)
                    shard_args.summary.write_text(
                        json.dumps(shard_summary), encoding="utf-8"
                    )

                matrix_args = SimpleNamespace(
                    case_file=case_file,
                    scope="mssanitizer",
                    tool=tool,
                    runtime_manifest=manifest_path,
                    timeout=60,
                    loop_nums=1,
                    gm_init_mode="not_applicable",
                    single_process_mode="off",
                    require_tiling_log=True,
                    matrix_root=matrix_root,
                    soc="ascend950",
                    test_artifact=[],
                )
                aggregate = VERIFIER._build_matrix_summary(matrix_args)
                aggregate_path = matrix_root / "aggregate_summary.json"
                aggregate_path.write_text(
                    json.dumps(aggregate), encoding="utf-8"
                )
                aggregates.append(aggregate_path)

            suite_args = SimpleNamespace(
                case_file=case_file,
                aggregate=aggregates,
                soc="ascend950",
                test_artifact=[],
                output=root / "suite.json",
            )
            self.assertEqual(VERIFIER.verify_sanitizer_suite(suite_args), 0)
            suite = json.loads(suite_args.output.read_text(encoding="utf-8"))
            self.assertTrue(suite["passed"])

            expected_name = str(bindings[0]["kernel_name"])
            wrong_name = next(
                name
                for name in VERIFIER._manifest_key_to_kernels(manifest)[
                    int(bindings[0]["tiling_key"])
                ]
                if name != expected_name
            )
            bad_log = root / "memcheck/shard_0_1/memcheck.log"
            bad_log.write_text(
                f"Start memcheck sanitizer on kernel {wrong_name}\n"
                "[mssanitizer] Sanitizer finished on kernel "
                f"{wrong_name}. No error detected.\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "manifest 之外"):
                VERIFIER.verify_sanitizer_suite(suite_args)

    def test_sanitizer_suite_requires_all_four_tools(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            case_file = root / "cases.json"
            cases = [
                {
                    "id": case_id,
                    "inputs": [
                        {
                            "name": "case_spec",
                            "range_values": json.dumps(
                                {
                                    "expected_tiling_key": 1000 + case_id,
                                    "gate_dtype": "bf16",
                                    "beta_dtype": "bf16",
                                }
                            ),
                        }
                    ],
                }
                for case_id in range(432)
            ]
            case_file.write_text(
                json.dumps(cases),
                encoding="utf-8",
            )
            pair_digest = hashlib.sha256(
                "".join(
                    f"{case_id},{1000 + case_id}\n" for case_id in range(432)
                ).encode("ascii")
            ).hexdigest()
            aggregates = []
            canonical = {}
            for tool in ("memcheck", "racecheck", "initcheck", "synccheck"):
                matrix_root = root / tool
                matrix_root.mkdir()
                path = matrix_root / "aggregate_summary.json"
                summary = {
                    "schema": "kda-prepare-atk-matrix/v2",
                    "scope": "mssanitizer",
                    "tool": tool,
                    "timeout_seconds": 60,
                    "loop_nums": 1,
                    "gm_init_mode": "not_applicable",
                    "single_process_mode": "off",
                    "require_tiling_log": True,
                    "runtime_manifest_sha256": "runtime",
                    "expected_cases": 432,
                    "observed_cases": 432,
                    "case_json_sha256": hashlib.sha256(
                        case_file.read_bytes()
                    ).hexdigest(),
                    "expected_tiling_key_count": 432,
                    "host_tiling_key_count": 432,
                    "launch_tiling_key_count": 432,
                    "sanitizer_started_tiling_key_count": 432,
                    "sanitizer_started_kernel_name_count": 432,
                    "sanitizer_started_kernel_name_sha256": "kernel-names",
                    "sanitizer_started_kernel_binding_count": 432,
                    "sanitizer_started_kernel_binding_sha256": "c" * 64,
                    "sanitizer_outer_log_count": 432,
                    "sanitizer_outer_log_sha256": "b" * 64,
                    "key_pair_sha256": pair_digest,
                    "complete": True,
                    "passed": True,
                }
                canonical[tool] = summary
                path.write_text(json.dumps(summary), encoding="utf-8")
                aggregates.append(path)

            args = SimpleNamespace(
                case_file=case_file,
                aggregate=aggregates,
                soc="ascend950",
                test_artifact=[root / "runner.py"],
                output=root / "suite.json",
            )
            with mock.patch.object(
                VERIFIER,
                "_build_matrix_summary",
                side_effect=lambda matrix_args: canonical[matrix_args.tool],
            ) as builder:
                self.assertEqual(VERIFIER.verify_sanitizer_suite(args), 0)
                self.assertEqual(builder.call_count, 4)
            suite = json.loads(args.output.read_text(encoding="utf-8"))
            self.assertTrue(suite["passed"])

            mismatched_binding = dict(canonical["racecheck"])
            mismatched_binding[
                "sanitizer_started_kernel_binding_sha256"
            ] = "d" * 64
            canonical["racecheck"] = mismatched_binding
            aggregates[1].write_text(
                json.dumps(mismatched_binding), encoding="utf-8"
            )
            with mock.patch.object(
                VERIFIER,
                "_build_matrix_summary",
                side_effect=lambda matrix_args: canonical[matrix_args.tool],
            ):
                with self.assertRaisesRegex(ValueError, "调度绑定不一致"):
                    VERIFIER.verify_sanitizer_suite(args)
            canonical["racecheck"] = dict(mismatched_binding)
            canonical["racecheck"][
                "sanitizer_started_kernel_binding_sha256"
            ] = "c" * 64
            aggregates[1].write_text(
                json.dumps(canonical["racecheck"]), encoding="utf-8"
            )

            forged = dict(canonical["memcheck"])
            forged["passed"] = False
            aggregates[0].write_text(json.dumps(forged), encoding="utf-8")
            with mock.patch.object(
                VERIFIER,
                "_build_matrix_summary",
                side_effect=lambda matrix_args: canonical[matrix_args.tool],
            ):
                with self.assertRaisesRegex(ValueError, "原始证据不一致"):
                    VERIFIER.verify_sanitizer_suite(args)

            aggregates[0].write_text(
                json.dumps(canonical["memcheck"]), encoding="utf-8"
            )
            args.aggregate = aggregates[:-1]
            with self.assertRaisesRegex(ValueError, "四种 sanitizer"):
                VERIFIER.verify_sanitizer_suite(args)


if __name__ == "__main__":
    unittest.main()
